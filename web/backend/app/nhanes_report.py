"""Human-readable NHANES results alongside the complete TAME analysis record."""
from datetime import datetime, timezone
from hashlib import sha256
from html import escape
import json
from pathlib import Path
import shutil
import threading
import zipfile

import pandas as pd
from tametools import __version__
from tametools.io import write_tame, write_xlsx
from tametools.planned_analysis import write_analysis_report
from .nhanes_data import AGE_GROUPS, BY_ID, CAUTIONS, WEIGHT_GUIDE
from .tametools_bridge import json_safe

REPORT_LOCK = threading.Lock()
GROUP_LABELS={'SELECTED':'선택 대상 전체','MALE':'남성','FEMALE':'여성'}


def format_number(value):
    if value is None or pd.isna(value): return '산출 불가'
    return f'{value:,.3f}'.rstrip('0').rstrip('.')


def result_payload(output, settings, info):
    survey=output.tables['survey_means']
    sample=output.tables['sample_summary']
    rows=[]
    for _,s in sample.loc[sample.group.isin(GROUP_LABELS)].iterrows():
        w=survey.loc[survey.result_id.eq(s.result_id)&survey.group.eq(s.group)].iloc[0]
        rows.append(dict(result_id=s.result_id, label=BY_ID[s.result_id]['name'], unit=s.unit,
            group=s.group, group_label=GROUP_LABELS[s.group], n=int(s.analysis_n), domain_n=int(s.domain_n),
            missing_n=int(s.missing_n), below_limit_n=int(s.censored_n) if any(v['id']==s.result_id and v.get('flag') for v in info['variables']) else None,
            policy_excluded_n=int(s.policy_excluded_n),
            weighted_mean=w['mean'], ci_low=w.ci95_low, ci_high=w.ci95_high, se=w.se,
            median=s['median'], q25=s.q25, q75=s.q75,
            status=w.status, note='유효 표본의 설계 자유도가 부족해 신뢰구간을 산출하지 않았습니다.' if w.status=='insufficient_domain_df' else ''))
    histograms=output.tables['histograms'].to_dict('records')
    selected=[r for r in rows if r['group']=='SELECTED']
    cautions=list(CAUTIONS)
    cautions.append('2017–2018년 ALT는 공개 대체값과 검출한계 미만 표시를 함께 보존합니다.' if info['censoring_available'] else '2021–2023년의 선택 가능한 6개 검사에는 검출한계 표시 열이 제공되지 않습니다. 공개값을 그대로 사용하며 검출한계 미만 인원은 산출하지 않습니다.')
    for r in selected:
        if r['domain_n'] and r['missing_n']/r['domain_n']>.1:
            cautions.insert(0,f"{r['label']}: 선택 대상 중 검사값 결측 비율이 {r['missing_n']/r['domain_n']:.1%}입니다. 관측값이 있는 참여자만의 추정임을 확인하세요.")
        if r['n']<30:
            cautions.insert(0,f"{r['label']}: 유효 관측이 {r['n']}명으로 적어 요약과 신뢰구간을 주의해서 해석해야 합니다.")
    age=AGE_GROUPS[settings['age_group']]['label']
    first=selected[0]
    intro=(f"{info['label']} 자료에서 {age}의 {len(selected)}개 검사항목을 분석했습니다. "
           f"{first['label']}의 유효 관측은 {first['n']:,}명이고, "
           f"조사 가중 평균은 {format_number(first['weighted_mean'])} {first['unit']}입니다.")
    return json_safe(dict(settings=settings, cycle=info['cycle'], cycle_label=info['label'], age_label=age,
        rows=rows, histograms=histograms, interpretation=intro, cautions=cautions,
        weight_variable=info['weight_variable'], weight_label=info['weight_label'],
        total_people=info['people'], positive_weight_people=info['positive_weight_people'],
        sources=info['sources'], age_top_code='80 = 80세 이상',
        sampling_note='가중 평균·95% 신뢰구간: 조사설계 적용 / 중앙값·분포: 비가중 관측 요약',
        policy_label='공개값 유지' if settings['policy']=='RELEASED' else '검출한계 미만 ALT 제외'))


def display_table(result):
    return pd.DataFrame([{
        '검사항목':r['label'], '단위':r['unit'], '그룹':r['group_label'], '유효 인원':r['n'],
        '선택 대상 인원':r['domain_n'], '결측 인원':r['missing_n'],
        '가중 평균':r['weighted_mean'], '95% CI 하한':r['ci_low'], '95% CI 상한':r['ci_high'],
        '표본 중앙값':r['median'], '표본 Q1':r['q25'], '표본 Q3':r['q75'],
        '검출한계 미만 인원':r['below_limit_n'], '정책에 따른 제외 인원':r['policy_excluded_n'],
    } for r in result['rows']])


def write_report(output, result, info, directory, source_directory):
    from docx import Document
    from docx.shared import Cm, Pt
    directory=Path(directory)
    directory.mkdir(exist_ok=False)
    # Matplotlib's global renderer is used by the existing report writer.
    with REPORT_LOCK:
        write_analysis_report(output,directory/'audit',docx=False)
    write_tame(directory/'analysis.tame',output.dataset)
    write_xlsx(directory/'data.xlsx',output.dataset)
    summary=display_table(result)
    summary.to_csv(directory/'results.csv',index=False,encoding='utf-8-sig')
    with pd.ExcelWriter(directory/'results.xlsx',engine='openpyxl') as writer:
        summary.to_excel(writer,sheet_name='결과',index=False)
        pd.DataFrame({'해석상 주의사항':result['cautions']}).to_excel(writer,sheet_name='해석 안내',index=False)
        pd.DataFrame([dict(항목=k,설정=str(v)) for k,v in result['settings'].items()]).to_excel(writer,sheet_name='분석 조건',index=False)
        pd.DataFrame(info['source_files']).to_excel(writer,sheet_name='자료 출처',index=False)
        for sheet in writer.book:
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(70,max(15,max(len(str(c.value or '')) for c in col)+2))
    headers=['검사항목 / 단위','그룹','유효 인원','가중 평균','95% 신뢰구간','표본 중앙값','결측 인원']
    values=[[f"{r['label']} ({r['unit']})",r['group_label'],str(r['n']),format_number(r['weighted_mean']),
             f"{format_number(r['ci_low'])} – {format_number(r['ci_high'])}",format_number(r['median']),str(r['missing_n'])] for r in result['rows']]
    doc=Document()
    section=doc.sections[0];section.page_width=Cm(29.7);section.page_height=Cm(21)
    section.left_margin=section.right_margin=Cm(1.8)
    style=doc.styles['Normal'];style.font.name='Malgun Gothic';style.font.size=Pt(9)
    doc.add_heading('NHANES 임상화학 분석 보고서',0)
    doc.add_paragraph(result['interpretation'])
    doc.add_paragraph(f"분석 대상: {result['age_label']} | {result['weight_variable']} ({result['weight_label']}) | {result['policy_label']}")
    doc.add_paragraph(result['sampling_note'])
    table=doc.add_table(rows=1,cols=len(headers));table.style='Light Shading Accent 1'
    for c,label in zip(table.rows[0].cells,headers):c.text=label
    for row in values:
        for c,value in zip(table.add_row().cells,row):c.text=value
    doc.add_heading('결과 해석과 주의사항',1)
    for text in result['cautions']:doc.add_paragraph(text,style='List Bullet')
    doc.add_heading('출처와 재현',1)
    for s in info['sources']:doc.add_paragraph(s['name']+': '+s['url'])
    doc.add_paragraph(f'tametools {__version__}. 같은 조건의 재실행: tametools analyze analysis.tame --output-dir rerun')
    doc.add_paragraph('analysis.tame에 전체 참여자·가중치·자료 출처·분석 설정을 저장했습니다. audit 폴더에는 전체 표본설계와 전체 검사 대상(ALL)을 포함한 원시 분석표 및 코드 해시가 있습니다.')
    doc.save(directory/'report_ko.docx')
    html=['<!doctype html><html lang="ko"><meta charset="utf-8"><title>NHANES 분석 보고서</title>',
          '<style>body{max-width:1120px;margin:48px auto;font:15px/1.8 system-ui,sans-serif;color:#163c3c;padding:0 24px}table{border-collapse:collapse;width:100%;font-size:13px}td,th{padding:10px;border-bottom:1px solid #d8e5e2;text-align:left}th{background:#edf4f1}li{margin:10px 0}h1,h2{line-height:1.3}a{color:#126c5c}</style>',
          '<h1>NHANES 임상화학 분석 보고서</h1><p>'+escape(result['interpretation'])+'</p>',
          '<p>'+escape(result['sampling_note'])+'</p><table><thead><tr>'+''.join('<th>'+escape(v)+'</th>' for v in headers)+'</tr></thead><tbody>']
    html+=['<tr>'+''.join('<td>'+escape(v)+'</td>' for v in row)+'</tr>' for row in values]
    html+=['</tbody></table><h2>분석 조건</h2><p>'+escape(result['age_label']+' / '+result['weight_variable']+' / '+result['policy_label'])+'</p>',
           '<h2>결과 해석과 주의사항</h2><ul>'+''.join('<li>'+escape(v)+'</li>' for v in result['cautions'])+'</ul>',
           '<h2>자료 출처</h2><ul>'+''.join('<li><a href="'+escape(s['url'],quote=True)+'">'+escape(s['name'])+'</a></li>' for s in info['sources'])+'</ul>',
           '<p>같은 조건으로 다시 실행: <code>tametools analyze analysis.tame --output-dir rerun</code></p></html>']
    (directory/'report_ko.html').write_text('\n'.join(html),encoding='utf-8')
    (directory/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    shutil.copy2(Path(source_directory)/'source_manifest.json',directory/'source_manifest.json')
    files={p.relative_to(directory).as_posix():sha256(p.read_bytes()).hexdigest() for p in directory.rglob('*') if p.is_file()}
    (directory/'bundle_manifest.json').write_text(json.dumps(dict(version=__version__,created_utc=datetime.now(timezone.utc).isoformat(),files=files),indent=2),encoding='utf-8')
    archive=directory.parent/'nhanes_results.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(directory.rglob('*')):
            if p.is_file():z.write(p,'results/'+p.relative_to(directory).as_posix())
    return dict(zip=archive,docx=directory/'report_ko.docx',html=directory/'report_ko.html',
                csv=directory/'results.csv',xlsx=directory/'results.xlsx',tame=directory/'analysis.tame',data_xlsx=directory/'data.xlsx')
