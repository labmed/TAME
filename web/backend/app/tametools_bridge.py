from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
import importlib
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from typing import Any

import openpyxl
import pandas as pd

from tametools.actions import action_payloads, action_pipeline_payloads, execute_action, execute_action_pipeline
from tametools.age import AGE_CANONICAL_UNITS, AGE_DEFAULT_UNIT, AGE_UCUM_SYSTEM, age_value_profile, standardize_age_dataset
from tametools.analysis import describe_dataset, exploratory_data_analysis, validate_dataset
from tametools.cellstate import parse_serialized_cell, serialize_cell
from tametools.config import ci_get
from tametools.io import (
    CONTROL_SHEETS,
    _augment_multisheet_meta,
    _bundle_from_raw_sections,
    _serialize_control_sections,
    _sheet_dataset,
    _sheet_text,
    _source_column_name,
    read_tame,
    read_xlsx,
    write_tame,
    write_xlsx,
)
from tametools.merge import merge_datasets
from tametools.models import ColumnSpec, TameDataset, merged_column_specs
from tametools.plugin_base import configured_plugin_modules, list_plugins, run_plugin
from tametools.plugin_base.base import normalize_plugin_name
from tametools.plugin_base.roles import role_contract_payload
from tametools.presets import apply_preset, preset_payloads
from tametools.models import OperationOutput
from tametools.provenance import append_log_entry, inherit_logs
from tametools.provenance_audit import history_payload
from .provenance_ui import web_operation
from tametools.reporting import with_visualizations
from tametools.review_profiles import category_value_profile, datetime_value_profile
from tametools.sex import SEX_CANONICAL_VALUES, sex_value_profile, standardize_sex_dataset
from tametools.tag_placement import tags_to_header, tags_to_meta
from tametools.tag_catalog import flat_tag_names, tag_catalog_payload
from tametools.tags import build_header, merge_tags
from tametools.toml_compat import dumps as dumps_toml
from tametools.toml_compat import loads as loads_toml
from tametools.transforms import anonymize_dataset
from tametools.transforms import fix_num_comparator_values


MAX_GRID_ROWS = 5000


@dataclass(frozen=True)
class StoredUpload:
    upload_id: str
    filename: str
    path: Path


def inspect_xlsx(path: str | Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        return [_inspect_sheet(workbook[sheet_name]) for sheet_name in workbook.sheetnames]
    finally:
        workbook.close()


def load_tame_payload(
    path: str | Path,
    *,
    filename: str = "",
    plugin_workspace: str | Path | None = None,
    allow_web_plugins: bool = False,
) -> dict[str, Any]:
    dataset = read_tame(path)
    if plugin_workspace is not None and allow_web_plugins:
        dataset = materialize_web_plugins(dataset, plugin_workspace)
    return dataset_payload(dataset, filename=filename or Path(path).name)


def is_tame_workbook(path: str | Path) -> bool:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        sheet_names = {name.upper() for name in workbook.sheetnames}
        return "DATA" in sheet_names and "META" in sheet_names
    finally:
        workbook.close()


def load_tame_workbook_payload(
    path: str | Path,
    *,
    filename: str = "",
    plugin_workspace: str | Path | None = None,
    allow_web_plugins: bool = False,
) -> dict[str, Any]:
    dataset = read_xlsx(path)
    if plugin_workspace is not None and allow_web_plugins:
        dataset = materialize_web_plugins(dataset, plugin_workspace)
    return dataset_payload(
        dataset,
        filename=filename or Path(path).name,
        warnings=["Loaded workbook DATA and META sheets directly as a TAME dataset."],
    )


def convert_xlsx_selection(
    path: str | Path,
    *,
    sheet_names: list[str],
    mode: str,
) -> tuple[TameDataset, list[str]]:
    selected = [str(name) for name in sheet_names if str(name).strip()]
    if not selected:
        raise ValueError("Select at least one worksheet.")

    normalized_mode = str(mode or "single").lower()
    if normalized_mode not in {"single", "merge"}:
        raise ValueError("mode must be one of: single, merge")
    if normalized_mode == "single":
        selected = selected[:1]

    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"Unknown worksheet(s): {', '.join(missing)}")

        controls = _workbook_controls(workbook)
        warnings = selection_warnings(path, selected, mode=normalized_mode)
        datasets = [
            _sheet_dataset(
                workbook,
                sheet_name,
                meta=deepcopy(controls.meta),
                schema=deepcopy(controls.schema),
                job=deepcopy(controls.job),
                source_path=str(path),
            )
            for sheet_name in selected
        ]

        from hashlib import sha256
        source_sha = sha256(Path(path).read_bytes()).hexdigest()
        datasets = [append_log_entry(ds, action='IMPORT_XLSX', parent=Path(path).name,
            message=f'Excel 시트 {sheet_name}: {len(ds.df)}행을 가져왔습니다.',
            counts={'INPUT_ROWS':len(ds.df), 'SOURCE_FILES':1},
            parameters={'source_file_sha256': source_sha, 'sheet': sheet_name, 'mode': normalized_mode,
                        'formula_policy': 'cached values (data_only=True)'})
            for sheet_name, ds in zip(selected, datasets)]
        if normalized_mode == "single":
            dataset = datasets[0].replace(raw_sections=_serialize_control_sections(datasets[0].meta, datasets[0].schema, datasets[0].job))
            return dataset, warnings

        sheet_source_column = _source_column_name(controls.meta, None)
        merged = merge_datasets(
            datasets,
            source_labels=selected,
            add_source_column=True,
            source_column_name=sheet_source_column,
            source_column_tags=("SHEET", "STR"),
            numeric_conflict="promote",
            prefer_tag_names=True,
        )
        meta = _augment_multisheet_meta(merged.dataset.meta, selected, source_column_name=sheet_source_column)
        dataset = merged.dataset.replace(
            meta=meta,
            schema=controls.schema,
            job=controls.job,
            raw_sections=_serialize_control_sections(meta, controls.schema, controls.job),
            source_path=str(path),
        )
        dataset = append_log_entry(dataset, action='XLSX_MERGE_BINDING', input_dataset=merged.dataset,
            message='병합한 시트의 출처 열과 시트 목록을 연결했습니다.',
            parameters={'sheets': selected, 'source_column': sheet_source_column})
        return dataset, warnings + merged.warnings
    finally:
        workbook.close()


def selection_warnings(path: str | Path, selected: list[str], *, mode: str) -> list[str]:
    sheets = {sheet["name"]: sheet for sheet in inspect_xlsx(path)}
    warnings: list[str] = []
    for name in selected:
        warnings.extend(f"{name}: {warning}" for warning in sheets.get(name, {}).get("warnings", []))
    if mode == "merge" and len(selected) > 1:
        warnings.append("Selected worksheets are merged with a [[SHEET::STR]]시트명 source column.")
    if mode == "single" and len(selected) > 1:
        warnings.append("Single-sheet conversion uses the first selected worksheet.")
    return warnings


def dataset_payload(
    dataset: TameDataset,
    *,
    filename: str = "",
    warnings: list[str] | None = None,
    tag_storage: str = "preserve",
    max_rows: int = MAX_GRID_ROWS,
) -> dict[str, Any]:
    from .workbench_analysis import data_context, data_role
    plugin_list, plugin_warnings = plugin_payloads(dataset.meta)
    all_warnings = [*(warnings or []), *plugin_warnings]
    settings = dataset.settings()
    row_limit = min(len(dataset.df), max_rows)
    data_rows = serialized_data_rows(dataset, settings)
    rows: list[dict[str, Any]] = []
    for row_index in range(row_limit):
        record: dict[str, Any] = {"__rowid": row_index + 1}
        for column_index, _column in enumerate(dataset.columns):
            record[_column_field(column_index)] = data_rows[row_index][column_index]
        rows.append(record)

    validation = validate_dataset(dataset)
    provenance = history_payload(dataset)
    return {
        "provenance": provenance,
        "auditBaseline": provenance['verification']['current'],
        "kind": "dataset",
        "dataRole": data_role(dataset),
        "dataContext": data_context(dataset),
        "referenceSettings": json_safe(ci_get(dataset.meta, "RI_EP28", {})),
        "sexNormalization": json_safe(ci_get(dataset.meta, "SEX_NORMALIZATION", {})),
        "filename": filename,
        "rowCount": len(dataset.df),
        "visibleRowCount": row_limit,
        "truncated": len(dataset.df) > row_limit,
        "columns": [
            {
                "field": _column_field(index),
                "index": index,
                "name": column.name,
                "tags": list(column.tags),
                "taggedHeader": column.tagged_header,
                "originalHeader": column.original_header,
                "metadata": column_metadata_payload(dataset.column_metadata(column), column.tags),
            }
            for index, column in enumerate(dataset.columns)
        ],
        "rows": rows,
        "dataRows": data_rows,
        "metaText": dumps_toml(dataset.meta),
        "sections": {
            "SCHEMA": dataset.raw_sections.get("SCHEMA") or dumps_toml(dataset.schema),
            "JOB": dataset.raw_sections.get("JOB") or dumps_toml(dataset.job),
        },
        "issues": issues_payload(validation.issues),
        "warnings": all_warnings,
        "actions": action_payloads(dataset.meta),
        "actionPipelines": action_pipeline_payloads(dataset.meta),
        "analyses": analysis_payloads(dataset.meta),
        "charts": chart_payloads(dataset),
        "presets": preset_payloads(),
        "plugins": plugin_list,
        "tagCatalog": tag_catalog_payload(dataset.meta),
        "tagDefinitions": tag_definition_payloads(dataset.meta),
        "commonTags": flat_tag_names(dataset.meta),
        "roleSummary": role_summary_payload(dataset),
        "tagStorage": tag_storage,
    }


def dataset_from_payload(payload: dict[str, Any], *, standardize: bool = True) -> TameDataset:
    columns_payload = sorted(payload.get("columns", []), key=lambda item: int(item.get("index", 0)))
    rows_payload = payload.get("dataRows", payload.get("rows", []))
    if not isinstance(rows_payload, list):
        rows_payload = []
    meta_text = str(payload.get("metaText", ""))
    sections = payload.get("sections", {}) if isinstance(payload.get("sections", {}), dict) else {}
    schema_text = str(sections.get("SCHEMA", ""))
    job_text = str(sections.get("JOB", ""))

    meta = loads_toml(meta_text)
    meta = merge_payload_extensions(meta, columns_payload, payload.get("tagDefinitions", []))
    schema = loads_toml(schema_text)
    job = loads_toml(job_text)
    settings = meta.get("SETTINGS", {}) if isinstance(meta.get("SETTINGS", {}), dict) else {}

    columns: list[ColumnSpec] = []
    for index, column in enumerate(columns_payload):
        name = str(column.get("name") or f"column_{index + 1}")
        tags = merge_tags(column.get("tags", []))
        columns.append(
            ColumnSpec(
                original_header=build_header(name, tags),
                name=name,
                tags=tags,
            )
        )

    matrix: list[list[Any]] = []
    for row in rows_payload:
        matrix.append(serialized_payload_row(row, len(columns), settings))

    frame = pd.DataFrame(matrix, columns=[column.name for column in columns], dtype=object)
    raw_sections = {}
    if meta_text.strip():
        raw_sections["META"] = meta_text
    if schema_text.strip():
        raw_sections["SCHEMA"] = schema_text
    if job_text.strip():
        raw_sections["JOB"] = job_text
    dataset = TameDataset(
        df=frame,
        columns=columns,
        meta=meta,
        schema=schema,
        job=job,
        raw_sections=raw_sections,
        source_path=str(payload.get("filename") or ""),
    )
    if standardize:
        return standardize_sex_dataset(dataset, binary_map=payload.get("sexBinaryMap"))
    return dataset


def serialized_data_rows(dataset: TameDataset, settings: dict[str, Any]) -> list[list[Any]]:
    return [[json_safe(serialize_cell(value, settings, for_excel=False)) for value in row]
            for row in dataset.df.itertuples(index=False, name=None)]


def serialized_payload_row(row: Any, column_count: int, settings: dict[str, Any]) -> list[Any]:
    if isinstance(row, dict):
        return [parse_serialized_cell(row.get(_column_field(index), None), settings) for index in range(column_count)]
    if isinstance(row, list):
        values = row[:column_count] + [None] * max(0, column_count - len(row))
        return [parse_serialized_cell(value, settings) for value in values]
    return [None for _index in range(column_count)]


def merge_payload_extensions(meta: dict[str, Any], columns_payload: list[dict[str, Any]], tag_definitions_payload: Any) -> dict[str, Any]:
    updated = deepcopy(meta) if isinstance(meta, dict) else {}
    updated = merge_column_metadata_payload(updated, columns_payload)
    updated = merge_tag_definitions_payload(updated, tag_definitions_payload)
    return updated


def merge_column_metadata_payload(meta: dict[str, Any], columns_payload: list[dict[str, Any]]) -> dict[str, Any]:
    columns_section = ci_get(meta, "COLUMN", {})
    columns_section = deepcopy(columns_section) if isinstance(columns_section, dict) else {}
    changed = False

    for column in columns_payload:
        name = str(column.get("name") or "").strip()
        if not name:
            continue
        tags = merge_tags(column.get("tags", []))
        metadata = column.get("metadata", {})
        entry = deepcopy(metadata) if isinstance(metadata, dict) else {}
        entry = clean_metadata_value(entry)
        if tags:
            entry["TAGS"] = list(tags)
        if entry:
            columns_section[name] = entry
            changed = True

    if changed:
        meta = deepcopy(meta)
        meta["COLUMN"] = columns_section
    return meta


def merge_tag_definitions_payload(meta: dict[str, Any], tag_definitions_payload: Any) -> dict[str, Any]:
    incoming: dict[str, dict[str, Any]] = {}
    if isinstance(tag_definitions_payload, dict):
        for name, config in tag_definitions_payload.items():
            if isinstance(config, dict):
                incoming[str(name)] = clean_metadata_value(deepcopy(config))
    elif isinstance(tag_definitions_payload, list):
        for item in tag_definitions_payload:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or "").strip()
            if not name:
                continue
            properties = item.get("properties", {})
            config = deepcopy(properties) if isinstance(properties, dict) else {}
            label = str(item.get("label") or "").strip()
            description = str(item.get("description") or "").strip()
            inherits = item.get("inherits", [])
            age_bin_width = item.get("ageBinWidth", "")
            if label:
                config["LABEL"] = label
            if description:
                config["DESCRIPTION"] = description
            if isinstance(inherits, list):
                config["INHERITS"] = [str(tag).strip() for tag in inherits if str(tag).strip()]
            elif str(inherits).strip():
                config["INHERITS"] = [tag.strip() for tag in str(inherits).split(",") if tag.strip()]
            if str(age_bin_width).strip():
                try:
                    config["AGE_BIN_WIDTH"] = int(age_bin_width)
                except ValueError:
                    config["AGE_BIN_WIDTH"] = str(age_bin_width).strip()
            config = clean_metadata_value(config)
            if config:
                incoming[name] = config

    if not incoming:
        return meta

    section = ci_get(meta, "TAG_DEFINITIONS", {})
    section = deepcopy(section) if isinstance(section, dict) else {}
    section.update(incoming)
    meta = deepcopy(meta)
    meta["TAG_DEFINITIONS"] = section
    return meta


def clean_metadata_value(value: Any) -> Any:
    if isinstance(value, dict):
        result: dict[str, Any] = {}
        for key, item in value.items():
            if str(key).upper() in {"TAGS", "EFFECTIVETAGS"}:
                continue
            cleaned = clean_metadata_value(item)
            if cleaned in ("", None, [], {}):
                continue
            result[str(key)] = cleaned
        return result
    if isinstance(value, list):
        return [clean_metadata_value(item) for item in value if clean_metadata_value(item) not in ("", None, [], {})]
    return value


def column_metadata_payload(metadata: dict[str, Any], tags: tuple[str, ...] | list[str]) -> dict[str, Any]:
    result = deepcopy(metadata) if isinstance(metadata, dict) else {}
    result.pop("TAGS", None)
    result["effectiveTags"] = list(tags)
    return json_safe(result)


def tag_definition_payloads(meta: dict[str, Any] | None) -> list[dict[str, Any]]:
    section = ci_get(meta, "TAG_DEFINITIONS", {}) if isinstance(meta, dict) else {}
    if not isinstance(section, dict):
        return []
    payloads: list[dict[str, Any]] = []
    for name, config in section.items():
        if not isinstance(config, dict):
            continue
        properties = deepcopy(config)
        label = str(ci_get(properties, "LABEL", name))
        description = str(ci_get(properties, "DESCRIPTION", ""))
        inherits = ci_get(properties, "INHERITS", [])
        age_bin_width = ci_get(properties, "AGE_BIN_WIDTH", "")
        for key in ("LABEL", "DESCRIPTION", "INHERITS", "AGE_BIN_WIDTH"):
            properties.pop(key, None)
        payloads.append(
            {
                "name": str(name),
                "label": label,
                "description": description,
                "inherits": [str(item) for item in inherits] if isinstance(inherits, list) else [],
                "ageBinWidth": age_bin_width,
                "properties": json_safe(properties),
            }
        )
    return payloads


def role_summary_payload(dataset: TameDataset) -> dict[str, Any]:
    result_columns = [
        {
            "name": column.name,
            "field": _column_field(index),
            "tags": list(column.tags),
            "pivotContext": json_safe(ci_get(dataset.column_metadata(column), "PIVOT_CONTEXT", {})),
        }
        for index, column in enumerate(dataset.columns)
        if dataset.column_has_tag(column, "RESULT")
    ]
    qualified_ids = [
        {
            "name": column.name,
            "field": _column_field(index),
            "tags": [tag for tag in column.tags if str(tag).startswith("ID(") or str(tag).endswith("_ID") or str(tag) == "ID"],
        }
        for index, column in enumerate(dataset.columns)
        if dataset.column_has_tag(column, "ID")
    ]
    age_columns = [
        {
            "name": column.name,
            "field": _column_field(index),
            "tags": list(column.tags),
            "ageBinWidth": _age_bin_width_from_profile(dataset, column.name),
        }
        for index, column in enumerate(dataset.columns)
        if dataset.column_has_tag(column, "AGE")
    ]
    return {
        "resultColumns": result_columns,
        "multiResult": len(result_columns) > 1,
        "pivotContextColumns": [item for item in result_columns if item.get("pivotContext")],
        "qualifiedIds": qualified_ids,
        "ageColumns": age_columns,
        "customTagCount": len(tag_definition_payloads(dataset.meta)),
    }


def _age_bin_width_from_profile(dataset: TameDataset, column_name: str) -> int | None:
    frame = age_value_profile(dataset)
    if frame.empty:
        return None
    values = frame.loc[frame["column"] == column_name, "preferred_band_width"].dropna().unique()
    if len(values) == 0:
        return None
    try:
        return int(values[0])
    except (TypeError, ValueError):
        return None


def validation_review_payload(payload: dict[str, Any]) -> dict[str, Any]:
    dataset = dataset_from_payload(payload, standardize=False)
    validation = validate_dataset(dataset)
    return {
        "issues": issues_payload(validation.issues),
        "profiles": profile_payloads(dataset, binary_map=payload.get("sexBinaryMap")),
    }


@web_operation('MANUAL_EDIT')
def refresh_payload(payload: dict[str, Any]) -> dict[str, Any]:
    dataset = dataset_from_payload(payload, standardize=False)
    return dataset_payload(
        dataset,
        filename=payload_filename(payload),
        tag_storage=str(payload.get("tagStorage") or "preserve"),
    )


@web_operation('FIX')
def apply_validation_fix_payload(payload: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    dataset = dataset_from_payload(payload, standardize=False)
    action = str(options.get("action", "")).strip().lower()
    log_parameters: dict[str, Any] = {"action": action}
    if action == "normalize-sex-by-source":
        from tametools.sex_normalization import normalize_sex_by_source
        dataset = normalize_sex_by_source(dataset, options.get('sexNormalization'))
        log_parameters.update(rows_preserved=True, raw_values_preserved=True)
    elif action in {"standardize-sex", "standardize_sex", "sex"}:
        dataset = standardize_sex_dataset(dataset, binary_map=payload.get("sexBinaryMap"))
    elif action in {"standardize-age", "standardize_age", "age"}:
        dataset = standardize_age_dataset(dataset)
    elif action == "comparator-policy":
        from .workbench_analysis import comparator_policy
        policy = str(options.get("numComparatorHandling") or "")
        dataset, affected = comparator_policy(dataset, policy)
        log_parameters.update(policy=policy, affected_cells=affected, rows_preserved=True, raw_values_preserved=True)
    elif action in {"fix-num-comparator", "fix_num_comparator", "num-comparator", "num"}:
        handling = str(options.get("handling") or options.get("numComparatorHandling") or "delete").strip().lower()
        dataset, _table = fix_num_comparator_values(dataset, handling=handling)
        log_parameters["handling"] = handling
    else:
        raise ValueError(f"Unsupported validation fix action: {options.get('action')}")
    suffix = {"normalize-sex-by-source":"sex_std", "standardize-sex":"sex_std", "standardize-age":"age_std", "comparator-policy":"num_policy"}.get(action,"fixed")
    filename = derived_filename(payload, suffix)
    dataset = append_log_entry(
        dataset,
        action="fix",
        message="Applied validation fix.",
        parent=payload_filename(payload),
        output=filename,
        parameters=log_parameters,
    )
    result = dataset_payload(dataset, filename=filename)
    if action == 'normalize-sex-by-source':
        result['operationSummary'] = f"출처별 성별 정규화를 완료했습니다. 원 코드와 {len(dataset.df):,}행을 보존하고 표준 성별 열을 작업 입력에 연결했습니다."
    if action == "comparator-policy":
        result['operationSummary'] = f"부등호 {affected}셀에 분석 정책을 저장했습니다. 원문과 {len(dataset.df):,}행은 보존했습니다."
    return result


def create_web_plugin_payload(payload: dict[str, Any], options: dict[str, Any], workspace: str | Path) -> dict[str, Any]:
    plugin_name = normalize_plugin_name(str(options.get("name") or "").strip())
    if not plugin_name:
        raise ValueError("Plugin name is required.")
    description = str(options.get("description") or "").strip()
    source = str(options.get("source") or "").strip()
    if not source:
        raise ValueError("Plugin source is required.")
    compile(source, f"<web plugin {plugin_name}>", "exec")

    dataset = dataset_from_payload(payload, standardize=False)
    module_name = _web_plugin_module_name(plugin_name)
    meta = _meta_with_web_plugin(dataset.meta, plugin_name, module_name, description, source)
    updated = dataset.replace(meta=meta, raw_sections=_raw_sections_with_meta(dataset, meta))
    output_name = derived_filename(payload, f"plugin_{_plugin_slug(plugin_name)}")
    updated = append_log_entry(
        updated,
        action="plugin-create",
        message=f"Created web plugin {plugin_name}.",
        parent=payload_filename(payload),
        output=output_name,
        parameters={"plugin": plugin_name, "module": module_name},
    )
    return dataset_payload(updated, filename=output_name)


@web_operation('PLUGIN_RUN')
def run_web_plugin_payload(payload: dict[str, Any], options: dict[str, Any], workspace: str | Path) -> dict[str, Any]:
    plugin_name = str(options.get("plugin") or options.get("name") or "").strip()
    if not plugin_name:
        raise ValueError("Plugin name is required.")
    plugin_options = options.get("pluginOptions", {})
    if not isinstance(plugin_options, dict):
        raise ValueError("pluginOptions must be an object.")

    if plugin_name in {"RI_EP28", "REFERENCE_INTERVAL_EP28"}:
        from .workbench_analysis import reference_analysis
        return reference_analysis(payload, plugin_options, workspace)

    dataset = dataset_from_payload(payload, standardize=False)
    output = run_plugin(dataset, plugin_name, dataset.meta, plugin_name, plugin_options)
    if output is None:
        if not plugins_explicitly_allowed(options) and configured_external_plugin(dataset.meta, plugin_name):
            raise PermissionError(
                "External plugin execution is disabled. Trust this TAME file and retry with allowPlugins=true."
            )
        if plugins_explicitly_allowed(options):
            dataset = materialize_web_plugins(dataset, workspace)
            output = run_plugin(
                dataset,
                plugin_name,
                dataset.meta,
                plugin_name,
                plugin_options,
                allow_external=True,
            )
    if output is None:
        raise ValueError(f"Unknown plugin: {plugin_name}")

    result_dataset = _dataset_from_plugin_output(output, dataset, plugin_name)
    if not result_dataset.columns_with_tag('RESULT'):
        result_meta = deepcopy(result_dataset.meta)
        result_meta['WEB_WORKBENCH'] = dict(KIND='result', ANALYSIS=plugin_name, SOURCE=payload_filename(payload))
        result_dataset = result_dataset.replace(meta=result_meta)
    filename = derived_filename(payload, f"plugin_{_plugin_slug(plugin_name)}_run")
    result_dataset = append_log_entry(
        result_dataset,
        action="plugin-run",
        message=output.message or f"Ran plugin {plugin_name}.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"plugin": plugin_name, "options": plugin_options},
        warnings=list(output.warnings or []), operation_output=output,
    )
    asset_warnings = [f"file: {path}" for path in output.files or []]
    return dataset_payload(result_dataset, filename=filename, warnings=[*(output.warnings or []), *asset_warnings])


@web_operation('META_ACTION')
def apply_meta_action_payload(payload: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    action_name = str(options.get("action") or options.get("name") or "").strip()
    if not action_name:
        raise ValueError("Action name is required.")

    dataset = dataset_from_payload(payload, standardize=False)
    output = execute_action(dataset, action_name)
    if output.dataset is None:
        raise ValueError(f"ACTION did not return a dataset: {action_name}")

    filename = derived_filename(payload, _slug_suffix(action_name))
    result_dataset = append_log_entry(
        output.dataset,
        action="meta-action",
        message=output.message or f"Executed ACTION {action_name}.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"action": action_name, "config": ci_get(ci_get(dataset.meta,"ACTIONS",{}),action_name,{})},
        warnings=list(output.warnings or []), operation_output=output,
    )
    return dataset_payload(result_dataset, filename=filename, warnings=list(output.warnings or []))


@web_operation('ACTION_PIPELINE')
def apply_action_pipeline_payload(payload: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    pipeline_name = str(options.get("pipeline") or options.get("name") or "DEFAULT").strip() or "DEFAULT"

    dataset = dataset_from_payload(payload, standardize=False)
    result = execute_action_pipeline(dataset, pipeline_name)

    warnings = [
        warning
        for output in result.outputs
        for warning in (output.warnings or [])
    ]
    messages = [
        f"{output.name}: {output.message}"
        for output in result.outputs
        if output.message
    ]
    filename = derived_filename(payload, f"pipeline_{_slug_suffix(pipeline_name)}")
    result_dataset = append_log_entry(
        result.final_dataset,
        action="action-pipeline",
        message="; ".join(messages) or f"Executed ACTION pipeline {pipeline_name}.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"pipeline": pipeline_name, "actions": [output.name for output in result.outputs]},
        warnings=warnings,
    )
    return dataset_payload(result_dataset, filename=filename, warnings=warnings)


@web_operation('ANALYSIS')
def run_meta_analysis_payload(
    payload: dict[str, Any],
    options: dict[str, Any],
    workspace: str | Path | None = None,
) -> dict[str, Any]:
    analysis_name = str(options.get("analysis") or options.get("name") or "").strip()
    if not analysis_name:
        raise ValueError("Analysis name is required.")

    dataset = dataset_from_payload(payload, standardize=False)
    name, config = _resolve_analysis_config(dataset.meta, analysis_name)
    plugin_name = str(ci_get(config, "PLUGIN", "CHEMISTRY_ANALYSIS")).strip() or "CHEMISTRY_ANALYSIS"
    plugin_options = ci_get(config, "OPTIONS", {})
    if not isinstance(plugin_options, dict):
        plugin_options = {}
    inline_options = {key: value for key, value in config.items() if key.upper() not in {"LABEL", "DESCRIPTION", "PLUGIN", "OPTIONS"}}
    merged_options = {**inline_options, **plugin_options}
    output = run_plugin(dataset, plugin_name, dataset.meta, name, merged_options)
    if output is None:
        if not plugins_explicitly_allowed(options) and configured_external_plugin(dataset.meta, plugin_name):
            raise PermissionError(
                "External analysis plugin execution is disabled. Trust this TAME file and retry with allowPlugins=true."
            )
        if plugins_explicitly_allowed(options):
            dataset = materialize_web_plugins(dataset, workspace or tempfile.gettempdir())
            output = run_plugin(
                dataset,
                plugin_name,
                dataset.meta,
                name,
                merged_options,
                allow_external=True,
            )
    if output is None or output.dataset is None:
        raise ValueError(f"Analysis plugin did not return a dataset: {plugin_name}")

    filename = derived_filename(payload, f"analysis_{_slug_suffix(name)}")
    result_dataset = append_log_entry(
        output.dataset,
        action="analysis",
        message=output.message or f"Ran analysis {name}.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"analysis": name, "plugin": plugin_name, "options": merged_options},
        warnings=list(output.warnings or []), operation_output=output,
    )
    return dataset_payload(result_dataset, filename=filename, warnings=list(output.warnings or []))


@web_operation('PRESET')
def apply_preset_payload(payload: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    preset_name = str(options.get("preset") or options.get("name") or "CLINICAL_CHEMISTRY").strip() or "CLINICAL_CHEMISTRY"
    dataset = dataset_from_payload(payload, standardize=False)
    output = apply_preset(dataset, preset_name)
    if output.dataset is None:
        raise ValueError(f"Preset did not return a dataset: {preset_name}")

    filename = derived_filename(payload, f"preset_{_slug_suffix(preset_name)}")
    result_dataset = append_log_entry(
        output.dataset,
        action="preset",
        message=output.message or f"Applied preset {preset_name}.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"preset": preset_name},
        warnings=list(output.warnings or []), operation_output=output,
    )
    return dataset_payload(result_dataset, filename=filename, warnings=list(output.warnings or []), tag_storage="meta")


@web_operation('TAG_PLACEMENT')
def apply_tag_placement_payload(payload: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    mode = str(options.get("mode") or options.get("storage") or "").strip().lower().replace("-", "_")
    dataset = dataset_from_payload(payload, standardize=False)
    if mode in {"meta", "metadata"}:
        output = tags_to_meta(dataset)
        tag_storage = "meta"
        suffix = "tags_meta"
    elif mode in {"header", "headers", "data"}:
        output = tags_to_header(dataset)
        tag_storage = "header"
        suffix = "tags_header"
    else:
        raise ValueError("Tag placement mode must be one of: meta, header.")

    filename = derived_filename(payload, suffix)
    result_dataset = append_log_entry(
        output.dataset or dataset,
        action="tag-placement",
        message=output.message or f"Reorganized tags into {tag_storage}.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"mode": tag_storage},
        warnings=list(output.warnings or []), operation_output=output,
    )
    return dataset_payload(
        result_dataset,
        filename=filename,
        warnings=list(output.warnings or []),
        tag_storage=tag_storage,
    )


def dataset_to_tame_bytes(dataset: TameDataset, workspace: str | Path, *, tag_storage: str = "preserve") -> bytes:
    path = _temporary_output_path(Path(workspace), ".tame")
    try:
        write_tame(path, dataset, tag_storage=tag_storage, standardize_sex_values=False)
        return path.read_bytes()
    finally:
        _safe_unlink(path)


def dataset_to_xlsx_bytes(dataset: TameDataset, workspace: str | Path, *, tag_storage: str = "preserve") -> bytes:
    path = _temporary_output_path(Path(workspace), ".xlsx")
    try:
        write_xlsx(path, dataset, tag_storage=tag_storage, standardize_sex_values=False)
        return path.read_bytes()
    finally:
        _safe_unlink(path)


@web_operation('EDA')
def eda_payload(payload: dict[str, Any]) -> dict[str, Any]:
    dataset = dataset_from_payload(payload)
    from .workbench_analysis import data_context
    validation = validate_dataset(dataset)
    report = exploratory_data_analysis(dataset)
    tables = {
        "validation_issues": pd.DataFrame(issues_payload(validation.issues)),
        "describe": describe_dataset(dataset),
        "summary": report.summary,
        "category_distribution": report.category_distribution,
        "numeric_percentiles": report.numeric_percentiles,
        "numeric_percentile_bands": report.numeric_percentile_bands,
        "result_by_summary": report.result_by_summary,
        "comparator_profile": report.comparator_profile,
        "comparator_policy_impact": report.comparator_policy_impact,
        "harmonization_preview": report.harmonization_preview,
    }
    report.warnings.extend(w for w in data_context(dataset)['warnings'] if w not in report.warnings)
    filename = derived_filename(payload, "eda")
    result_dataset = eda_result_dataset(dataset, tables=tables, warnings=report.warnings)
    result_dataset = append_log_entry(
        result_dataset,
        action="eda",
        message="Generated exploratory data analysis TAME dataset.",
        parent=payload_filename(payload),
        output=filename,
        parameters={"tables": list(tables.keys()), "SETTINGS": dataset.settings()},
        operation_output=OperationOutput(name="EDA", tables=tables),
        warnings=list(report.warnings),
    )
    return {
        "tables": {name: frame_records(table) for name, table in tables.items()},
        "warnings": list(report.warnings),
        "dataset": dataset_payload(result_dataset, filename=filename, warnings=list(report.warnings)),
    }


def reference_interval_payload(payload: dict[str, Any], options: dict[str, Any] | None = None) -> dict[str, Any]:
    from .workbench_analysis import reference_analysis
    return reference_analysis(payload, options)


@web_operation('ANONYMIZE')
def anonymize_payload(payload: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    dataset = dataset_from_payload(payload)
    default_hash_tags = "ID,ID(patient),ID(hospital),PATIENT_ID,HOSPITAL_ID"
    transformed, mappings = anonymize_dataset(
        dataset,
        hash_tags=split_option(options.get("hashTags", default_hash_tags)),
        drop_tags=split_option(options.get("dropTags", "NAME")),
        hash_columns=split_option(options.get("hashColumns", "")),
        drop_columns=split_option(options.get("dropColumns", "")),
        salt=str(options.get("salt", "")),
    )
    filename = derived_filename(payload, "anonymized")
    transformed = append_log_entry(
        transformed,
        action="anonymize",
        message="Generated anonymized TAME dataset.",
        parent=payload_filename(payload),
        output=filename,
        parameters={
            "hashTags": split_option(options.get("hashTags", default_hash_tags)),
            "dropTags": split_option(options.get("dropTags", "NAME")),
            "hashColumns": split_option(options.get("hashColumns", "")),
            "dropColumns": split_option(options.get("dropColumns", "")),
        },
    )
    result = dataset_payload(transformed, filename=filename)
    result["mappingTables"] = {name: frame_records(table) for name, table in mappings.items()}
    return result


def eda_result_dataset(source_dataset: TameDataset, *, tables: dict[str, pd.DataFrame], warnings: list[str]) -> TameDataset:
    rows: list[dict[str, Any]] = []
    for table_name, table in tables.items():
        if table is None or table.empty:
            continue
        for row_index, record in enumerate(table.to_dict(orient="records"), start=1):
            for field, value in record.items():
                rows.append(
                    {
                        "table": table_name,
                        "row": row_index,
                        "field": str(field),
                        "value": "" if value is None or (isinstance(value, float) and math.isnan(value)) else str(value),
                    }
                )

    headers = ["[[EDA::TABLE::STR]]table", "[[EDA::ROW_INDEX::NUM]]row", "[[EDA::FIELD::STR]]field", "[[EDA::VALUE::STR]]value"]
    columns = merged_column_specs(headers)
    df = pd.DataFrame(rows, columns=[column.name for column in columns]) if rows else pd.DataFrame(columns=[column.name for column in columns])
    meta = {
        "INFO": {"DESCRIPTION": "Exploratory data analysis output dataset"},
        "LOG": deepcopy(source_dataset.meta.get("LOG", [])),
        "SETTINGS": {"VALIDATE_ERROR": "REPORT"},
        "WORKS": {"DEFAULT": ["DESCRIBE"]},
        "EDA": {
            "SOURCE_PATH": source_dataset.source_path or "",
            "TABLES": list(tables.keys()),
            "WARNINGS": list(warnings),
        },
    }
    return TameDataset(
        df=df,
        columns=columns,
        meta=meta,
        schema={},
        job={},
        raw_sections={},
        source_path=source_dataset.source_path,
    )


def plugin_payloads(meta: dict | None) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        plugins = list_plugins(meta)
    except Exception as exc:
        return [], [f"Plugin load failed: {exc}"]
    warnings = external_plugin_warnings(meta)
    payloads = [
        {
            "name": plugin.name,
            "description": plugin.description,
            "roles": role_contract_payload(plugin.roles),
            "external": False,
            "requiresTrust": False,
        }
        for plugin in plugins
    ]
    seen = {item["name"] for item in payloads}
    for plugin_name, config in web_plugin_configs(meta).items():
        normalized = normalize_plugin_name(str(plugin_name))
        if normalized in seen:
            continue
        description = str(ci_get(config, "DESCRIPTION", "")) if isinstance(config, dict) else ""
        payloads.append(
            {
                "name": normalized,
                "description": description,
                "roles": [],
                "external": True,
                "requiresTrust": True,
            }
        )
        seen.add(normalized)
    return payloads, warnings


def analysis_payloads(meta: dict | None) -> list[dict[str, Any]]:
    section = ci_get(meta, "ANALYSES", {})
    if not isinstance(section, dict):
        return []
    builtin_plugins = {plugin.name for plugin in list_plugins(None)}
    external_modules = configured_plugin_modules(meta)
    web_plugin_names = {normalize_plugin_name(name) for name in web_plugin_configs(meta)}
    analyses: list[dict[str, Any]] = []
    for name, config in section.items():
        if not isinstance(config, dict):
            continue
        plugin_name = str(ci_get(config, "PLUGIN", "CHEMISTRY_ANALYSIS")).strip() or "CHEMISTRY_ANALYSIS"
        normalized_plugin = normalize_plugin_name(plugin_name)
        analyses.append(
            {
                "name": str(name),
                "label": str(ci_get(config, "LABEL", name)),
                "description": str(ci_get(config, "DESCRIPTION", "")),
                "plugin": plugin_name,
                "options": ci_get(config, "OPTIONS", {}) if isinstance(ci_get(config, "OPTIONS", {}), dict) else {},
                "mutates": True,
                "requiresTrust": normalized_plugin in web_plugin_names
                or (normalized_plugin not in builtin_plugins and bool(external_modules)),
            }
        )
    return analyses


def web_plugin_configs(meta: dict | None) -> dict[str, Any]:
    web_plugins = ci_get(meta, "WEB_PLUGINS", {}) if isinstance(meta, dict) else {}
    return web_plugins if isinstance(web_plugins, dict) else {}


def external_plugin_warnings(meta: dict | None) -> list[str]:
    modules = configured_plugin_modules(meta)
    if not modules:
        return []
    return [
        "External plugin modules are disabled until explicitly trusted: "
        + ", ".join(modules)
    ]


def configured_external_plugin(meta: dict | None, plugin_name: str) -> bool:
    normalized = normalize_plugin_name(plugin_name)
    if normalized in {normalize_plugin_name(name) for name in web_plugin_configs(meta)}:
        return True
    if not configured_plugin_modules(meta):
        return False
    return normalized not in {plugin.name for plugin in list_plugins(None)}


def plugins_explicitly_allowed(options: dict[str, Any]) -> bool:
    return bool(
        options.get("allowPlugins")
        or options.get("allowExternalPlugins")
        or options.get("trusted")
        or options.get("trustPlugins")
    )


def _resolve_analysis_config(meta: dict | None, analysis_name: str) -> tuple[str, dict[str, Any]]:
    section = ci_get(meta, "ANALYSES", {})
    if not isinstance(section, dict):
        raise ValueError("No META[ANALYSES] section is defined.")
    for name, config in section.items():
        if str(name).lower() == str(analysis_name).lower():
            if not isinstance(config, dict):
                raise ValueError(f"ANALYSIS {name} must be a table.")
            return str(name), config
    available = ", ".join(str(name) for name in section) or "(none)"
    raise ValueError(f"Undefined ANALYSIS: {analysis_name}. Available analyses: {available}")


def chart_payloads(dataset: TameDataset) -> list[dict[str, Any]]:
    section = ci_get(dataset.meta, "VISUALIZATIONS", {})
    if not isinstance(section, dict):
        return []
    charts: list[dict[str, Any]] = []
    for name, config in section.items():
        if not isinstance(config, dict):
            continue
        x_col = str(ci_get(config, "X", "")).strip()
        y_col = str(ci_get(config, "Y", "")).strip()
        series_col = str(ci_get(config, "SERIES", ci_get(config, "COLOR", ""))).strip()
        embedded = ci_get(config, 'ROWS', None)
        frame = pd.DataFrame(embedded) if isinstance(embedded, list) else dataset.df
        if not x_col or not y_col or x_col not in frame.columns or y_col not in frame.columns:
            continue
        max_points = int(ci_get(config, "MAX_POINTS", 80))
        filters = ci_get(config, 'FILTER', {})
        for column, value in filters.items():
            if column not in frame.columns:
                frame = frame.iloc[:0]
                break
            matches = frame[column].map(lambda v: str(v).strip().lower() in {'true','1'}) if value is True else frame[column].map(lambda v: str(v).strip().lower() in {'false','0'}) if value is False else frame[column].eq(value)
            frame = frame.loc[matches]
        frame = frame.head(max(0, max_points))
        rows: list[dict[str, Any]] = []
        for record in frame.to_dict(orient="records"):
            y_value = _chart_number(record.get(y_col))
            if y_value is None:
                continue
            rows.append(
                {
                    "x": str(json_safe(record.get(x_col))),
                    "y": y_value,
                    "series": str(json_safe(record.get(series_col))) if series_col and series_col in record else "",
                    **{target: _chart_number(record.get(ci_get(config, key, ''))) for target, key in
                       [('low','Y_LOW'), ('high','Y_HIGH'), ('lowCiLow','LOW_CI_LOW'), ('lowCiHigh','LOW_CI_HIGH'),
                        ('highCiLow','HIGH_CI_LOW'), ('highCiHigh','HIGH_CI_HIGH')] if ci_get(config, key, '')},
                }
            )
        charts.append(
            {
                "name": str(name),
                "type": str(ci_get(config, "TYPE", "BAR")).strip().lower() or "bar",
                "title": str(ci_get(config, "TITLE", name)),
                "x": x_col,
                "y": y_col,
                "series": series_col,
                "rows": rows,
            }
        )
    return charts


def _chart_number(value: Any) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def materialize_web_plugins(dataset: TameDataset, workspace: str | Path) -> TameDataset:
    web_plugins = ci_get(dataset.meta, "WEB_PLUGINS", {})
    if not isinstance(web_plugins, dict):
        return dataset

    meta = deepcopy(dataset.meta)
    changed = False
    for plugin_name, config in web_plugins.items():
        if not isinstance(config, dict):
            continue
        source = str(ci_get(config, "SOURCE", "")).strip()
        if not source:
            continue
        description = str(ci_get(config, "DESCRIPTION", ""))
        module_name = str(ci_get(config, "MODULE", "")).strip()
        if not module_name:
            module_name = _web_plugin_module_name(str(plugin_name))
        write_web_plugin_module(workspace, str(plugin_name), description, source, module_name=module_name)
        meta = _meta_with_plugin_module(meta, module_name)
        changed = True

    if not changed:
        return dataset
    return dataset.replace(meta=meta, raw_sections=_raw_sections_with_meta(dataset, meta))


def write_web_plugin_module(
    workspace: str | Path,
    plugin_name: str,
    description: str,
    source: str,
    *,
    module_name: str | None = None,
) -> str:
    workspace_path = Path(workspace)
    package_name = "tametools_web_plugins"
    package_dir = workspace_path / package_name
    package_dir.mkdir(parents=True, exist_ok=True)
    (package_dir / "__init__.py").write_text("", encoding="utf-8")
    if str(workspace_path) not in sys.path:
        sys.path.insert(0, str(workspace_path))

    normalized_name = normalize_plugin_name(plugin_name)
    resolved_module = module_name or _web_plugin_module_name(normalized_name)
    module_path = package_dir / f"{resolved_module.rsplit('.', 1)[-1]}.py"
    module_path.write_text(_web_plugin_module_source(normalized_name, description, source), encoding="utf-8")

    sys.modules.pop(resolved_module, None)
    importlib.invalidate_caches()
    importlib.import_module(resolved_module)
    return resolved_module


def _web_plugin_module_source(plugin_name: str, description: str, source: str) -> str:
    return f'''from __future__ import annotations

import inspect

import pandas as pd

from tametools.models import OperationOutput, TameDataset, merged_column_specs
from tametools.plugin_base.base import register_plugin

{source}


def _invoke_user_plugin(dataset, meta, step_name, options):
    if "run" not in globals() or not callable(globals()["run"]):
        raise RuntimeError("Plugin source must define callable run(...).")
    handler = globals()["run"]
    values = {{"dataset": dataset, "meta": meta, "step_name": step_name, "options": options}}
    signature = inspect.signature(handler)
    if any(parameter.kind == inspect.Parameter.VAR_KEYWORD for parameter in signature.parameters.values()):
        return handler(**values)
    accepted = {{name: value for name, value in values.items() if name in signature.parameters}}
    if accepted:
        return handler(**accepted)
    return handler(dataset, meta, step_name, options)


def _coerce_plugin_result(result, dataset, step_name):
    if isinstance(result, OperationOutput):
        return result
    if isinstance(result, TameDataset):
        return OperationOutput(name=step_name, dataset=result, table=result.df, message=f"plugin {{step_name}} rows={{len(result.df)}}")
    if isinstance(result, pd.DataFrame):
        frame = result.copy()
        frame.columns = [str(column) for column in frame.columns]
        columns = merged_column_specs(list(frame.columns))
        frame.columns = [column.name for column in columns]
        output_dataset = TameDataset(
            df=frame,
            columns=columns,
            meta={{"INFO": {{"DESCRIPTION": f"Plugin output: {{step_name}}"}}}},
            schema={{}},
            job={{}},
            raw_sections={{}},
            source_path=dataset.source_path,
        )
        return OperationOutput(name=step_name, dataset=output_dataset, table=frame, message=f"plugin {{step_name}} rows={{len(frame)}}")
    frame = pd.DataFrame([{{"value": "" if result is None else str(result)}}])
    columns = merged_column_specs(list(frame.columns))
    return OperationOutput(name=step_name, table=frame, message=f"plugin {{step_name}} value returned")


@register_plugin({plugin_name!r}, description={description!r})
def _registered_web_plugin(dataset, meta, step_name, options):
    return _coerce_plugin_result(_invoke_user_plugin(dataset, meta, step_name, options), dataset, step_name)
'''


def _meta_with_web_plugin(meta: dict[str, Any], plugin_name: str, module_name: str, description: str, source: str) -> dict[str, Any]:
    updated = _meta_with_plugin_module(deepcopy(meta), module_name)
    web_plugins = ci_get(updated, "WEB_PLUGINS", {})
    if not isinstance(web_plugins, dict):
        web_plugins = {}
    web_plugins[normalize_plugin_name(plugin_name)] = {
        "DESCRIPTION": description,
        "MODULE": module_name,
        "SOURCE": source,
    }
    updated["WEB_PLUGINS"] = web_plugins
    return updated


def _meta_with_plugin_module(meta: dict[str, Any], module_name: str) -> dict[str, Any]:
    plugins = ci_get(meta, "PLUGINS", {})
    if not isinstance(plugins, dict):
        plugins = {}
    modules = ci_get(plugins, "MODULES", [])
    if not isinstance(modules, list):
        modules = []
    module_items = [str(item) for item in modules]
    if module_name not in module_items:
        module_items.append(module_name)
    plugins["MODULES"] = module_items
    meta["PLUGINS"] = plugins
    return meta


def _dataset_from_plugin_output(output, source_dataset: TameDataset, plugin_name: str) -> TameDataset:
    if output.dataset is not None:
        return with_visualizations(output.dataset, output.charts)
    if output.table is not None:
        frame = output.table.copy()
    else:
        rows: list[dict[str, Any]] = []
        if output.message:
            rows.append({"field": "message", "value": output.message})
        for warning in output.warnings or []:
            rows.append({"field": "warning", "value": warning})
        frame = pd.DataFrame(rows or [{"field": "plugin", "value": plugin_name}])

    frame.columns = [str(column) for column in frame.columns]
    columns = merged_column_specs(list(frame.columns))
    frame.columns = [column.name for column in columns]
    result = TameDataset(
        df=frame,
        columns=columns,
        meta={"INFO": {"DESCRIPTION": f"Plugin output: {plugin_name}"}},
        schema={},
        job={},
        raw_sections={},
        source_path=source_dataset.source_path,
    )
    return with_visualizations(result, output.charts)


def _raw_sections_with_meta(dataset: TameDataset, meta: dict[str, Any]) -> dict[str, str]:
    raw_sections = dict(dataset.raw_sections)
    raw_sections["META"] = dumps_toml(meta)
    return raw_sections


def payload_filename(payload: dict[str, Any]) -> str:
    return str(payload.get("filename") or "dataset.tame")


def derived_filename(payload: dict[str, Any], suffix: str) -> str:
    source = Path(payload_filename(payload))
    name = source.name
    if name.endswith(".data.tame"):
        stem = name[: -len(".data.tame")]
    elif name.endswith(".meta.tame"):
        stem = name[: -len(".meta.tame")]
    elif source.suffix:
        stem = source.stem
    else:
        stem = name or "dataset"
    return f"{stem}.{suffix}.tame"


def _slug_suffix(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_").lower()
    return slug or "action"


def _plugin_slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_]+", "_", str(value).strip()).strip("_").lower()
    if not slug:
        return "plugin"
    if slug[0].isdigit():
        return f"p_{slug}"
    return slug


def _web_plugin_module_name(plugin_name: str) -> str:
    return f"tametools_web_plugins.{_plugin_slug(plugin_name)}"


def frame_records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    if frame is None or frame.empty:
        return []
    return [json_safe(record) for record in frame.to_dict(orient="records")]


def issues_payload(issues) -> list[dict[str, Any]]:
    return [
        {
            "row": issue.row_number,
            "column": issue.column,
            "tag": issue.tag,
            "value": str(issue.value),
            "message": issue.message,
            "severity": getattr(issue, "severity", "error"),
        }
        for issue in issues[:200]
    ]


def profile_payloads(dataset: TameDataset, *, binary_map: dict[str, str] | None = None) -> list[dict[str, Any]]:
    profiles = [*sex_profile_payloads(dataset, binary_map=binary_map), *age_profile_payloads(dataset),
                *category_profile_payloads(dataset), *datetime_profile_payloads(dataset)]
    specific = {p['field'] for p in profiles if p['tag'] != 'CATEGORY'}
    return [p for p in profiles if p['tag'] != 'CATEGORY' or p['field'] not in specific]


def sex_profile_payloads(dataset: TameDataset, *, binary_map: dict[str, str] | None = None) -> list[dict[str, Any]]:
    frame = sex_value_profile(dataset, binary_map=binary_map)
    if frame.empty:
        return []

    fields = {column.name: _column_field(index) for index, column in enumerate(dataset.columns)}
    profiles: list[dict[str, Any]] = []
    for column_name, column_frame in frame.groupby("column", sort=False):
        categories: list[dict[str, Any]] = []
        for canonical in SEX_CANONICAL_VALUES:
            group = column_frame.loc[column_frame["canonical_value"] == canonical].copy()
            if group.empty:
                continue
            raw_values = [
                {
                    "raw": str(row.raw_value),
                    "count": int(row.raw_count),
                    "willChange": bool(row.will_change),
                }
                for row in group.sort_values(["will_change", "raw_value"], ascending=[False, True]).itertuples(index=False)
            ]
            total = int(group["raw_count"].sum())
            categories.append(
                {
                    "canonical": canonical,
                    "total": total,
                    "rawValues": raw_values,
                    "summary": _profile_summary(canonical, total, raw_values),
                    "hasSuggestedChanges": any(item["willChange"] for item in raw_values),
                }
            )

        unrecognized = [
            {"raw": str(row.raw_value), "count": int(row.raw_count)}
            for row in column_frame.loc[~column_frame["recognized"]].itertuples(index=False)
        ]
        profiles.append(
            {
                "tag": "SEX",
                "column": str(column_name),
                "field": fields.get(str(column_name), ""),
                "categories": categories,
                "unrecognized": unrecognized,
                "hasSuggestedChanges": any(category["hasSuggestedChanges"] for category in categories),
            }
        )
    return profiles


def age_profile_payloads(dataset: TameDataset) -> list[dict[str, Any]]:
    frame = age_value_profile(dataset)
    if frame.empty:
        return []

    fields = {column.name: _column_field(index) for index, column in enumerate(dataset.columns)}
    profiles: list[dict[str, Any]] = []
    for column_name, column_frame in frame.groupby("column", sort=False):
        categories: list[dict[str, Any]] = []
        for unit_code, unit in AGE_CANONICAL_UNITS.items():
            group = column_frame.loc[column_frame["unit_code"] == unit_code].copy()
            if group.empty:
                continue
            raw_values = [
                {
                    "raw": str(row.raw_value),
                    "count": int(row.raw_count),
                    "willChange": bool(row.will_change),
                    "canonical": str(row.canonical_value),
                    "years": json_safe(row.years),
                }
                for row in group.sort_values(["years", "raw_value"], ascending=[True, True]).itertuples(index=False)
            ]
            total = int(group["raw_count"].sum())
            canonical = f"{unit['label']} ({unit_code})"
            categories.append(
                {
                    "canonical": canonical,
                    "total": total,
                    "rawValues": raw_values,
                    "summary": _profile_summary(canonical, total, raw_values),
                    "hasSuggestedChanges": any(item["willChange"] for item in raw_values),
                }
            )

        unrecognized = [
            {"raw": str(row.raw_value), "count": int(row.raw_count)}
            for row in column_frame.loc[~column_frame["recognized"]].itertuples(index=False)
        ]
        profiles.append(
            {
                "tag": "AGE",
                "column": str(column_name),
                "field": fields.get(str(column_name), ""),
                "categories": categories,
                "unrecognized": unrecognized,
                "hasSuggestedChanges": any(category["hasSuggestedChanges"] for category in categories),
                "distributions": [
                    {
                        "label": _preferred_age_distribution_label(column_frame),
                        "width": _preferred_age_distribution_width(column_frame),
                        "bands": _age_distribution_payload(column_frame, "preferred_band"),
                    }
                ],
                "standard": {
                    "system": AGE_UCUM_SYSTEM,
                    "defaultUnit": AGE_DEFAULT_UNIT,
                    "storage": "years without suffix; months with mo; days with d",
                    "allowedUnits": [
                        {"code": code, "label": str(unit["label"])}
                        for code, unit in AGE_CANONICAL_UNITS.items()
                    ],
                },
            }
        )
    return profiles


def _preferred_age_distribution_width(frame: pd.DataFrame) -> int:
    if "preferred_band_width" not in frame.columns:
        return 10
    values = frame.loc[frame["recognized"] & frame["preferred_band_width"].notna(), "preferred_band_width"].drop_duplicates()
    if values.empty:
        return 10
    try:
        width = int(values.iloc[0])
    except (TypeError, ValueError):
        return 10
    return width if width > 0 else 10


def _preferred_age_distribution_label(frame: pd.DataFrame) -> str:
    return f"{_preferred_age_distribution_width(frame)}-year distribution"


def _age_distribution_payload(frame: pd.DataFrame, column_name: str) -> list[dict[str, Any]]:
    recognized = frame.loc[frame["recognized"] & frame[column_name].notna()].copy()
    if recognized.empty:
        return []
    grouped = recognized.groupby(column_name, sort=False)["raw_count"].sum()
    return [{"band": str(band), "count": int(count)} for band, count in grouped.items()]


def category_profile_payloads(dataset: TameDataset) -> list[dict[str, Any]]:
    frame = category_value_profile(dataset)
    if frame.empty:
        return []

    fields = {column.name: _column_field(index) for index, column in enumerate(dataset.columns)}
    profiles: list[dict[str, Any]] = []
    for column_name, column_frame in frame.groupby("column", sort=False):
        categories = [
            {
                "canonical": str(row.raw_value),
                "total": int(row.raw_count),
                "rawValues": [{"raw": str(row.raw_value), "count": int(row.raw_count), "willChange": False}],
                "summary": f"{row.raw_value}: {int(row.raw_count)}",
                "hasSuggestedChanges": False,
            }
            for row in column_frame.sort_values(["raw_count", "raw_value"], ascending=[False, True]).itertuples(index=False)
        ]
        profiles.append(
            {
                "tag": "CATEGORY",
                "column": str(column_name),
                "field": fields.get(str(column_name), ""),
                "categories": categories,
                "unrecognized": [],
                "hasSuggestedChanges": False,
            }
        )
    return profiles


def datetime_profile_payloads(dataset: TameDataset) -> list[dict[str, Any]]:
    frame = datetime_value_profile(dataset)
    if frame.empty:
        return []

    fields = {column.name: _column_field(index) for index, column in enumerate(dataset.columns)}
    columns_by_name = {column.name: column for column in dataset.columns}
    profiles: list[dict[str, Any]] = []
    for column_name, column_frame in frame.groupby("column", sort=False):
        column = columns_by_name.get(str(column_name))
        tag = "DATETIME" if column is not None and dataset.column_has_tag(column, "DATETIME") else "DATE"
        recognized = column_frame.loc[column_frame["recognized"]].copy()
        unrecognized_frame = column_frame.loc[~column_frame["recognized"]].copy()
        parseable_count = int(recognized["raw_count"].sum()) if not recognized.empty else 0
        unparseable_count = int(unrecognized_frame["raw_count"].sum()) if not unrecognized_frame.empty else 0
        raw_values = [
            {
                "raw": str(row.raw_value),
                "count": int(row.raw_count),
                "willChange": False,
                "parsed": str(row.parsed_value),
            }
            for row in recognized.sort_values(["raw_count", "raw_value"], ascending=[False, True]).head(20).itertuples(index=False)
        ]
        unrecognized = [
            {"raw": str(row.raw_value), "count": int(row.raw_count)}
            for row in unrecognized_frame.sort_values(["raw_count", "raw_value"], ascending=[False, True]).itertuples(index=False)
        ]
        distributions = [
            {
                "label": "Datetime conversion",
                "bands": [
                    {"band": "parseable", "count": parseable_count},
                    {"band": "unparseable", "count": unparseable_count},
                ],
            }
        ]
        categories = [
            {
                "canonical": "parseable",
                "total": parseable_count,
                "rawValues": raw_values,
                "summary": f"parseable: {parseable_count}",
                "hasSuggestedChanges": False,
            }
        ]
        profiles.append(
            {
                "tag": tag,
                "column": str(column_name),
                "field": fields.get(str(column_name), ""),
                "categories": categories,
                "unrecognized": unrecognized,
                "hasSuggestedChanges": False,
                "distributions": distributions,
            }
        )
    return profiles


def _profile_summary(canonical: str, total: int, raw_values: list[dict[str, Any]]) -> str:
    raw_summary = ", ".join(f"{item['raw']}: {item['count']}" for item in raw_values)
    return f"{canonical}: {total} ({raw_summary})"


def json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, list | tuple):
        return [json_safe(item) for item in value]
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if hasattr(value, "item") and callable(value.item):
        try:
            item = value.item()
        except Exception:
            item = value
        if item is not value:
            return json_safe(item)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def split_option(value: Any) -> tuple[str, ...]:
    if isinstance(value, list):
        return tuple(str(item).strip() for item in value if str(item).strip())
    return tuple(part.strip() for part in str(value or "").split(",") if part.strip())


def _inspect_sheet(sheet) -> dict[str, Any]:
    warnings: list[str] = []
    rows = int(sheet.max_row or 0)
    columns = int(sheet.max_column or 0)
    header_values = ["" if value is None else str(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    blank_headers = sum(1 for value in header_values if not str(value).strip())
    duplicates = sorted({value for value in header_values if value and header_values.count(value) > 1})
    formula_count = _formula_count(sheet)

    if sheet.title.upper() in CONTROL_SHEETS:
        warnings.append("Control sheet; not used as DATA unless explicitly selected.")
    if rows == 0 or columns == 0:
        warnings.append("Empty worksheet.")
    if rows <= 1:
        warnings.append("No data rows below the header row.")
    if blank_headers:
        warnings.append(f"{blank_headers} blank header cell(s) will be named column_n.")
    if duplicates:
        warnings.append(f"Duplicate header(s): {', '.join(duplicates[:5])}.")
    if formula_count:
        warnings.append(f"{formula_count} formula cell(s) found; conversion uses cached workbook values.")

    return {
        "name": sheet.title,
        "rows": rows,
        "dataRows": max(rows - 1, 0),
        "columns": columns,
        "headers": header_values[:30],
        "isControl": sheet.title.upper() in CONTROL_SHEETS,
        "warnings": warnings,
    }


def _formula_count(sheet) -> int:
    count = 0
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                count += 1
    return count


def _workbook_controls(workbook: openpyxl.Workbook):
    return _bundle_from_raw_sections(
        {
            key: value
            for key, value in {
                "META": _sheet_text(workbook, "META"),
                "SCHEMA": _sheet_text(workbook, "SCHEMA"),
                "JOB": _sheet_text(workbook, "JOB"),
            }.items()
            if value.strip()
        }
    )


def _temporary_output_path(workspace: Path, suffix: str) -> Path:
    workspace.mkdir(parents=True, exist_ok=True)
    handle, path = tempfile.mkstemp(prefix="export-", suffix=suffix, dir=workspace)
    os.close(handle)
    return Path(path)


def _safe_unlink(path: Path) -> None:
    try:
        path.unlink()
    except FileNotFoundError:
        pass


def _column_field(index: int) -> str:
    return f"c{index}"
