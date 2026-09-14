from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import textwrap
import unittest

import openpyxl
from pydantic import TypeAdapter

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "tametools" / "src"))
sys.path.insert(0, str(ROOT / "web" / "backend"))

from app.tametools_bridge import (
    anonymize_payload,
    apply_action_pipeline_payload,
    apply_meta_action_payload,
    apply_preset_payload,
    apply_tag_placement_payload,
    apply_validation_fix_payload,
    convert_xlsx_selection,
    create_web_plugin_payload,
    dataset_from_payload,
    dataset_payload,
    dataset_to_tame_bytes,
    eda_payload,
    inspect_xlsx,
    is_tame_workbook,
    load_tame_payload,
    load_tame_workbook_payload,
    reference_interval_payload,
    run_meta_analysis_payload,
    run_web_plugin_payload,
    validation_review_payload,
)


class WebBridgeTest(unittest.TestCase):
    def test_data_meta_xlsx_loads_directly_as_tame_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "tame_workbook.xlsx"
            workbook = openpyxl.Workbook()
            data_sheet = workbook.active
            data_sheet.title = "DATA"
            data_sheet.append(["등록번호", "보고값"])
            data_sheet.append(["P001", "10"])
            meta_sheet = workbook.create_sheet("META")
            meta_sheet.append(["[TAGS]"])
            meta_sheet.append(['"등록번호" = ["ID"]'])
            meta_sheet.append(['"보고값" = ["RESULT", "NUM"]'])
            workbook.save(path)
            workbook.close()

            detected = is_tame_workbook(path)
            payload = load_tame_workbook_payload(path)

        self.assertTrue(detected)
        self.assertEqual(payload["kind"], "dataset")
        self.assertIn("[TAGS]", payload["metaText"])
        self.assertEqual(payload["columns"][0]["tags"], ["ID"])
        self.assertEqual(payload["columns"][1]["tags"], ["RESULT", "NUM"])
        self.assertEqual(payload["rows"][0]["c0"], "P001")

    def test_inspect_and_convert_xlsx_single_and_merge(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.xlsx"
            workbook = openpyxl.Workbook()
            chemistry = workbook.active
            chemistry.title = "Chemistry"
            chemistry.append(["[[ID]]등록번호", "[[TESTNAME]]검사항목명", "[[RESULT::<NUM>]]보고값"])
            chemistry.append(["P001", "AST", "<30"])
            hematology = workbook.create_sheet("Hematology")
            hematology.append(["[[ID]]등록번호", "[[TESTNAME]]검사항목명", "[[RESULT::NUM]]보고값"])
            hematology.append(["P002", "WBC", 5.2])
            workbook.save(path)
            workbook.close()

            sheets = inspect_xlsx(path)
            single, single_warnings = convert_xlsx_selection(path, sheet_names=["Chemistry"], mode="single")
            merged, merge_warnings = convert_xlsx_selection(path, sheet_names=["Chemistry", "Hematology"], mode="merge")

        self.assertEqual([sheet["name"] for sheet in sheets], ["Chemistry", "Hematology"])
        self.assertEqual(list(single.df["등록번호"]), ["P001"])
        self.assertFalse(single_warnings)
        self.assertEqual(len(merged.df), 2)
        self.assertIsNotNone(merged.first_column_with_tag("SHEET"))
        self.assertTrue(any("merged" in warning for warning in merge_warnings))

    def test_payload_roundtrip_and_anonymize_use_edited_header_tags(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["등록번호", "이름", "보고값"])
            sheet.append(["P001", "홍길동", "10"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)

        payload["columns"][0]["tags"] = ["ID"]
        payload["columns"][1]["tags"] = ["NAME"]
        payload["columns"][2]["tags"] = ["RESULT", "NUM"]
        rebuilt = dataset_from_payload(payload)
        anonymized = anonymize_payload(payload, {"hashTags": "ID", "dropTags": "NAME", "salt": "s"})

        self.assertIsNotNone(rebuilt.first_column_with_tag("RESULT"))
        self.assertIn("등록번호", anonymized["mappingTables"])
        self.assertNotIn("이름", [column["name"] for column in anonymized["columns"]])
        self.assertTrue(str(anonymized["rows"][0]["c0"]).startswith("anon_"))

    def test_payload_roundtrip_standardizes_confirmed_binary_sex_codes(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "binary_sex.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[SEX]]성별"])
            sheet.append([1])
            sheet.append([0])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)

        self.assertEqual({issue["tag"] for issue in payload["issues"]}, {"SEX"})
        payload["sexBinaryMap"] = {"1": "female", "0": "male"}
        rebuilt = dataset_from_payload(payload)
        normalized = dataset_payload(rebuilt)

        self.assertEqual(rebuilt.df["성별"].tolist(), ["female", "male"])
        self.assertEqual(normalized["issues"], [])
        self.assertEqual([row["c0"] for row in normalized["rows"]], ["female", "male"])

    def test_validation_review_profiles_raw_sex_values_and_applies_fix(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sex_review.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[SEX]]성별"])
            sheet.append(["M"])
            sheet.append(["M"])
            sheet.append(["male"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)

        self.assertEqual([row["c0"] for row in payload["rows"]], ["M", "M", "male"])

        review = validation_review_payload(payload)
        profile = review["profiles"][0]
        male = profile["categories"][0]
        fixed = apply_validation_fix_payload(payload, {"action": "standardize-sex"})

        self.assertEqual(male["canonical"], "male")
        self.assertEqual(male["total"], 3)
        self.assertEqual({item["raw"]: item["count"] for item in male["rawValues"]}, {"M": 2, "male": 1})
        self.assertEqual([row["c0"] for row in fixed["rows"]], ["male", "male", "male"])
        self.assertEqual(fixed["filename"], "dataset.sex_std.tame")
        self.assertIn("[[LOG]]", fixed["metaText"])
        self.assertIn('OPERATION = "FIX"', fixed["metaText"])

    def test_apply_fix_handles_num_comparator_values_by_policy(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "num_fix.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[RESULT::NUM]]보고값", "[[ID]]등록번호"])
            sheet.append(["<3", "A"])
            sheet.append(["4", "B"])
            sheet.append([">5", "C"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)

        deleted = apply_validation_fix_payload(payload, {"action": "fix-num-comparator", "numComparatorHandling": "delete"})
        valued = apply_validation_fix_payload(payload, {"action": "fix-num-comparator", "numComparatorHandling": "value"})

        self.assertEqual([row["c1"] for row in deleted["rows"]], ["B"])
        self.assertEqual([row["c0"] for row in valued["rows"]], ["3", "4", "5"])
        self.assertIn('handling = "value"', valued["metaText"])

    def test_create_and_run_web_plugin(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "plugin_source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[ID]]등록번호", "[[RESULT::NUM]]보고값"])
            sheet.append(["A", "1"])
            sheet.append(["B", "2"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset, filename="plugin_source.tame")
            created = create_web_plugin_payload(
                payload,
                {
                    "name": "HEAD_ROWS",
                    "description": "Return selected rows.",
                    "source": "def run(dataset, options):\n    return dataset.df.head(int(options.get('rows', 1))).copy()\n",
                },
                tmpdir,
            )
            generated = run_web_plugin_payload(
                created,
                {"plugin": "HEAD_ROWS", "pluginOptions": {"rows": 1}, "allowPlugins": True},
                tmpdir,
            )

        self.assertTrue(any(plugin["name"] == "HEAD_ROWS" for plugin in created["plugins"]))
        self.assertIn("[WEB_PLUGINS.HEAD_ROWS]", created["metaText"])
        self.assertEqual(generated["filename"], "plugin_source.plugin_head_rows.plugin_head_rows_run.tame")
        self.assertEqual(generated["rowCount"], 1)
        self.assertEqual(generated["rows"][0]["c0"], "A")
        self.assertIn('OPERATION = "PLUGIN-RUN"', generated["metaText"])

    def test_payload_roundtrip_preserves_rows_beyond_grid_preview_limit(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "large.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[ID]]id", "[[RESULT::NUM]]value"])
            for index in range(6001):
                sheet.append([f"P{index:04d}", index])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)
            rebuilt = dataset_from_payload(payload)

        self.assertTrue(payload["truncated"])
        self.assertEqual(payload["visibleRowCount"], 5000)
        self.assertEqual(len(payload["rows"]), 5000)
        self.assertEqual(len(payload["dataRows"]), 6001)
        self.assertEqual(len(rebuilt.df), 6001)
        self.assertEqual(rebuilt.df.iloc[-1]["id"], "P6000")

    def test_meta_actions_are_listed_and_executed_as_new_payload(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "action_source.xlsx"
            workbook = openpyxl.Workbook()
            data_sheet = workbook.active
            data_sheet.title = "DATA"
            data_sheet.append(["sex", "age"])
            data_sheet.append(["M", "30"])
            meta_sheet = workbook.create_sheet("META")
            meta_sheet.append(["[ACTIONS.ADD_SEX_AGE_TAGS]"])
            meta_sheet.append(['LABEL = "Add SEX/AGE tags"'])
            meta_sheet.append(['TYPE = "TAG_COLUMNS"'])
            meta_sheet.append(['DESCRIPTION = "Add SEX tag to sex column and AGE tag to age column."'])
            meta_sheet.append([""])
            meta_sheet.append(["[[ACTIONS.ADD_SEX_AGE_TAGS.COLUMNS]]"])
            meta_sheet.append(['COLUMN = "sex"'])
            meta_sheet.append(['TAGS = ["SEX"]'])
            meta_sheet.append([""])
            meta_sheet.append(["[[ACTIONS.ADD_SEX_AGE_TAGS.COLUMNS]]"])
            meta_sheet.append(['COLUMN = "age"'])
            meta_sheet.append(['TAGS = ["AGE"]'])
            workbook.save(path)
            workbook.close()

            payload = load_tame_workbook_payload(path, filename="action_source.xlsx")

        self.assertEqual(payload["actions"][0]["name"], "ADD_SEX_AGE_TAGS")
        generated = apply_meta_action_payload(payload, {"action": "ADD_SEX_AGE_TAGS"})

        self.assertEqual(generated["filename"], "action_source.add_sex_age_tags.tame")
        self.assertEqual(generated["columns"][0]["tags"], ["SEX"])
        self.assertEqual(generated["columns"][1]["tags"], ["AGE"])
        self.assertIn('OPERATION = "META-ACTION"', generated["metaText"])
        self.assertIn('action = "ADD_SEX_AGE_TAGS"', generated["metaText"])

    def test_action_pipelines_are_listed_and_executed_as_new_payload(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline_source.xlsx"
            workbook = openpyxl.Workbook()
            data_sheet = workbook.active
            data_sheet.title = "DATA"
            data_sheet.append(["group", "value"])
            data_sheet.append(["A", "1"])
            data_sheet.append(["B", "2"])
            meta_sheet = workbook.create_sheet("META")
            meta_sheet.append(["[ACTIONS.KEEP_A]"])
            meta_sheet.append(['TYPE = "ROW.INCLUDE"'])
            meta_sheet.append(['COLUMN = "group"'])
            meta_sheet.append(['VALUE = "A"'])
            meta_sheet.append([""])
            meta_sheet.append(["[ACTIONS.ADD_BATCH]"])
            meta_sheet.append(['TYPE = "ADD_COLUMN"'])
            meta_sheet.append(['NAME = "batch"'])
            meta_sheet.append(['VALUE = "one"'])
            meta_sheet.append([""])
            meta_sheet.append(["[ACTION_PIPELINES]"])
            meta_sheet.append(['DEFAULT = ["KEEP_A", "ADD_BATCH"]'])
            workbook.save(path)
            workbook.close()

            payload = load_tame_workbook_payload(path, filename="pipeline_source.xlsx")

        self.assertEqual(payload["actionPipelines"], [{"name": "DEFAULT", "actions": ["KEEP_A", "ADD_BATCH"]}])
        generated = apply_action_pipeline_payload(payload, {"pipeline": "DEFAULT"})

        self.assertEqual(generated["filename"], "pipeline_source.pipeline_default.tame")
        self.assertEqual(generated["rowCount"], 1)
        self.assertEqual([column["name"] for column in generated["columns"]], ["group", "value", "batch"])
        self.assertEqual(generated["rows"][0]["c2"], "one")
        self.assertIn('OPERATION = "ACTION-PIPELINE"', generated["metaText"])
        self.assertIn('pipeline = "DEFAULT"', generated["metaText"])

    def test_meta_analysis_runs_plugin_and_returns_chart_payload(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "chemistry_source.xlsx"
            workbook = openpyxl.Workbook()
            data_sheet = workbook.active
            data_sheet.title = "DATA"
            data_sheet.append(["검사항목명", "보고값", "장비명", "접수일", "검사시간", "등록번호", "sex", "age"])
            data_sheet.append(["AST", 10, "A1", "2024-01-01 08:00", "2024-01-01 08:30", "P1", "M", "30"])
            data_sheet.append(["AST", 20, "A1", "2024-01-01 09:00", "2024-01-01 09:45", "P2", "F", "40"])
            data_sheet.append(["ALT", 30, "A2", "2024-01-01 10:00", "2024-01-01 10:20", "P3", "M", "50"])
            meta_sheet = workbook.create_sheet("META")
            meta_sheet.append(["[TAGS]"])
            meta_sheet.append(['"검사항목명" = ["TESTNAME", "CATEGORY"]'])
            meta_sheet.append(['"보고값" = ["RESULT", "NUM"]'])
            meta_sheet.append(['"장비명" = ["INSTRUMENT", "CATEGORY"]'])
            meta_sheet.append(['"접수일" = ["DATETIME", "RECEIVED_AT"]'])
            meta_sheet.append(['"검사시간" = ["DATETIME", "RESULT_TIME"]'])
            meta_sheet.append(['"등록번호" = ["ID"]'])
            meta_sheet.append(['sex = ["SEX"]'])
            meta_sheet.append(['age = ["AGE"]'])
            meta_sheet.append([""])
            meta_sheet.append(["[ANALYSES.DEFAULT]"])
            meta_sheet.append(['LABEL = "TAT by test"'])
            meta_sheet.append(['PLUGIN = "CHEMISTRY_ANALYSIS"'])
            meta_sheet.append([""])
            meta_sheet.append(["[ANALYSES.DEFAULT.OPTIONS]"])
            meta_sheet.append(['MODE = "TAT_BY_TEST"'])
            workbook.save(path)
            workbook.close()

            payload = load_tame_workbook_payload(path, filename="chemistry_source.xlsx")
            generated = run_meta_analysis_payload(payload, {"analysis": "DEFAULT"})

        self.assertEqual(payload["analyses"][0]["name"], "DEFAULT")
        self.assertEqual(generated["filename"], "chemistry_source.analysis_default.tame")
        self.assertEqual(generated["charts"][0]["type"], "bar")
        self.assertEqual(generated["charts"][0]["y"], "중앙값TAT분")
        self.assertIn('OPERATION = "ANALYSIS"', generated["metaText"])

    def test_apply_preset_payload_adds_reusable_analyses_to_raw_dataset(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "raw_chemistry.xlsx"
            workbook = openpyxl.Workbook()
            data_sheet = workbook.active
            data_sheet.title = "DATA"
            data_sheet.append(["검사항목명", "보고값", "장비명", "접수일", "검사시간", "등록번호", "sex", "age"])
            data_sheet.append(["AST", 10, "A1", "2024-01-01 08:00", "2024-01-01 08:30", "P1", "M", "30"])
            data_sheet.append(["ALT", 20, "A2", "2024-01-01 09:00", "2024-01-01 09:30", "P2", "F", "40"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset, filename="raw_chemistry.xlsx")
            generated = apply_preset_payload(payload, {"preset": "CLINICAL_CHEMISTRY"})

        self.assertEqual(generated["filename"], "raw_chemistry.preset_clinical_chemistry.tame")
        self.assertEqual(len(generated["analyses"]), 13)
        self.assertEqual(generated["columns"][0]["tags"], ["CATEGORY", "TESTNAME", "ITEM"])
        self.assertTrue(any(item["name"] == "CORE_ANALYSIS_VIEW" for item in generated["actionPipelines"]))
        self.assertIn('OPERATION = "PRESET"', generated["metaText"])

    def test_tag_placement_payload_preserves_requested_storage_on_download(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "tag_source.xlsx"
            workbook = openpyxl.Workbook()
            data_sheet = workbook.active
            data_sheet.title = "DATA"
            data_sheet.append(["[[SEX]]sex"])
            data_sheet.append(["M"])
            meta_sheet = workbook.create_sheet("META")
            meta_sheet.append(["[TAGS]"])
            meta_sheet.append(['sex = ["CATEGORY"]'])
            workbook.save(path)
            workbook.close()

            payload = load_tame_workbook_payload(path, filename="tag_source.xlsx")
            meta_payload = apply_tag_placement_payload(payload, {"mode": "meta"})
            header_payload = apply_tag_placement_payload(payload, {"mode": "header"})
            meta_bytes = dataset_to_tame_bytes(
                dataset_from_payload(meta_payload),
                tmpdir,
                tag_storage=meta_payload["tagStorage"],
            )
            header_bytes = dataset_to_tame_bytes(
                dataset_from_payload(header_payload),
                tmpdir,
                tag_storage=header_payload["tagStorage"],
            )

        meta_text = meta_bytes.decode("utf-8")
        header_text = header_bytes.decode("utf-8")
        self.assertEqual(meta_payload["tagStorage"], "meta")
        self.assertEqual(header_payload["tagStorage"], "header")
        self.assertIn("[COLUMN.sex]", meta_text)
        self.assertIn('TAGS = ["CATEGORY", "SEX"]', meta_text)
        self.assertNotIn("[[SEX]]sex", meta_text)
        self.assertNotIn("[TAGS]", header_text)
        self.assertIn("[[CATEGORY::SEX]]sex", header_text)

    def test_validation_review_profiles_age_values(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "age_review.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[AGE]]나이"])
            sheet.append(["10a"])
            sheet.append(["2m"])
            sheet.append(["1d"])
            sheet.append(["0"])
            sheet.append(["bad"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)

        review = validation_review_payload(payload)
        age_profiles = [profile for profile in review["profiles"] if profile["tag"] == "AGE"]
        profile = age_profiles[0]

        self.assertEqual(profile["standard"]["system"], "http://unitsofmeasure.org")
        self.assertEqual(profile["standard"]["defaultUnit"], "a")
        self.assertEqual(profile["standard"]["storage"], "years without suffix; months with mo; days with d")
        self.assertEqual([category["canonical"] for category in profile["categories"]], ["year (a)", "month (mo)", "day (d)"])
        self.assertEqual(profile["distributions"][0]["bands"], [{"band": "<1", "count": 3}, {"band": "10-19", "count": 1}])
        self.assertEqual(profile["categories"][0]["rawValues"][0]["canonical"], "0")
        self.assertEqual(profile["categories"][0]["rawValues"][1]["canonical"], "10")
        self.assertEqual(profile["categories"][1]["rawValues"][0]["canonical"], "2mo")
        self.assertEqual(profile["unrecognized"], [{"raw": "bad", "count": 1}])
        fixed = apply_validation_fix_payload(payload, {"action": "standardize-age"})
        self.assertEqual([row["c0"] for row in fixed["rows"]], ["10", "2mo", "1d", "0", "bad"])

    def test_validation_review_profiles_category_and_datetime_values(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "category_datetime_review.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[CATEGORY]]진단군", "[[DATETIME]]채혈시각"])
            sheet.append(["A", "2024-01-01"])
            sheet.append(["B", "2024-01-02 13:30"])
            sheet.append(["A", "bad-date"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset)

        review = validation_review_payload(payload)
        profiles = {profile["tag"]: profile for profile in review["profiles"]}

        category = profiles["CATEGORY"]
        self.assertEqual({item["canonical"]: item["total"] for item in category["categories"]}, {"A": 2, "B": 1})

        datetimes = profiles["DATETIME"]
        self.assertEqual(datetimes["distributions"][0]["bands"], [{"band": "parseable", "count": 2}, {"band": "unparseable", "count": 1}])
        self.assertEqual(datetimes["unrecognized"], [{"raw": "bad-date", "count": 1}])
        self.assertEqual({issue["tag"] for issue in review["issues"]}, {"DATETIME"})

    def test_eda_payload_returns_generated_tame_dataset_with_log(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "eda_source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[TESTNAME::CATEGORY]]검사항목명", "[[SEX::CATEGORY::BY]]성별", "[[AGE]]나이", "[[RESULT::NUM]]보고값"])
            sheet.append(["AST", "M", "30", "10", "U/L"])
            sheet.append(["AST", "F", "40", "20", "U/L"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset, filename="eda_source.tame")

        result = eda_payload(payload)
        generated = result["dataset"]

        self.assertIn("tagCatalog", payload)
        self.assertIn("commonTags", payload)
        self.assertIn("describe", result["tables"])
        self.assertIn("category_distribution", result["tables"])
        self.assertIn("numeric_percentiles", result["tables"])
        self.assertIn("numeric_percentile_bands", result["tables"])
        self.assertIn("result_by_summary", result["tables"])
        self.assertTrue(any(row["category"] == "AST" and row["count"] == 2 for row in result["tables"]["category_distribution"]))
        self.assertTrue(any(row["by_column"] == "성별" for row in result["tables"]["result_by_summary"]))
        self.assertEqual(generated["filename"], "eda_source.eda.tame")
        self.assertEqual([column["name"] for column in generated["columns"]], ["table", "row", "field", "value"])
        self.assertTrue(any(row["c0"] == "category_distribution" for row in generated["rows"]))
        self.assertIn('OPERATION = "EDA"', generated["metaText"])
        self.assertIn('PARENT = "eda_source.tame"', generated["metaText"])
        TypeAdapter(dict[str, object]).dump_json(result)

    def test_custom_age_tag_definition_flows_to_payload_and_eda(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "age5_source.tame"
            path.write_text(
                textwrap.dedent(
                    """
                    <META>
                    [TAG_DEFINITIONS.AGE5]
                    LABEL = "Age with 5-year groups"
                    INHERITS = ["AGE"]
                    AGE_BIN_WIDTH = 5
                    </META>
                    <DATA>
                    [[TESTNAME::CATEGORY]]검사항목명\t[[SEX::CATEGORY::BY]]성별\t[[AGE5]]나이\t[[RESULT::NUM]]보고값
                    AST\tmale\t1\t10
                    AST\tfemale\t4\t20
                    AST\tmale\t6\t30
                    AST\tfemale\t9\t40
                    </DATA>
                    """
                ).strip(),
                encoding="utf-8",
            )
            payload = load_tame_payload(path, filename="age5_source.tame")

        self.assertIn("AGE5", payload["commonTags"])
        custom_groups = [group for group in payload["tagCatalog"] if group["name"] == "CUSTOM"]
        self.assertTrue(custom_groups)
        result = eda_payload(payload)
        age_rows = [row for row in result["tables"]["result_by_summary"] if row["by_column"] == "나이"]
        self.assertTrue(age_rows)
        self.assertEqual({row["by_grouping"] for row in age_rows}, {"age_band_5"})

    def test_web_payload_column_metadata_and_pivot_context_drive_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "wide_result.tame"
            path.write_text(
                textwrap.dedent(
                    """
                    <META>
                    [COLUMN.AST.PIVOT_CONTEXT]
                    TESTNAME = "AST"
                    UNIT = "U/L"
                    REF_LOW = "0"
                    REF_HIGH = "40"

                    [COLUMN.ALT.PIVOT_CONTEXT]
                    TESTNAME = "ALT"
                    UNIT = "U/L"
                    REF_LOW = "0"
                    REF_HIGH = "41"
                    </META>
                    <DATA>
                    [[ID(patient)::STR]]등록번호\t[[ID(sample)::STR]]검체번호\t[[RESULT::NUM]]AST\t[[RESULT::NUM]]ALT
                    P001\tS001\t28\t35
                    P002\tS002\t55\t62
                    P003\tS003\t22\t25
                    </DATA>
                    """
                ).strip(),
                encoding="utf-8",
            )
            payload = load_tame_payload(path, filename="wide_result.tame")

        ast_column = next(column for column in payload["columns"] if column["name"] == "AST")
        self.assertEqual(ast_column["metadata"]["PIVOT_CONTEXT"]["REF_HIGH"], "40")
        self.assertEqual(len(payload["roleSummary"]["resultColumns"]), 2)
        self.assertEqual(len(payload["roleSummary"]["pivotContextColumns"]), 2)
        self.assertEqual({item["tags"][0] for item in payload["roleSummary"]["qualifiedIds"]}, {"ID(patient)", "ID(sample)"})
        self.assertTrue(any(role["name"] == "result" for plugin in payload["plugins"] for role in plugin["roles"]))

        generated = run_web_plugin_payload(payload, {"plugin": "ABNORMAL_FLAG", "pluginOptions": {"MODE": "FLAG"}}, tempfile.gettempdir())

        self.assertIn("AST_이상플래그", [column["name"] for column in generated["columns"]])
        self.assertIn("ALT_이상플래그", [column["name"] for column in generated["columns"]])
        self.assertEqual(generated["roleSummary"]["multiResult"], True)

    def test_web_payload_edited_tag_definition_controls_review_bins(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "edited_age.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["나이", "보고값"])
            sheet.append(["1", "10"])
            sheet.append(["4", "20"])
            sheet.append(["6", "30"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset, filename="edited_age.xlsx")

        payload["columns"][0]["tags"] = ["AGE5"]
        payload["columns"][1]["tags"] = ["RESULT", "NUM"]
        payload["tagDefinitions"] = [
            {
                "name": "AGE5",
                "label": "5-year age",
                "description": "AGE-compatible web-created tag.",
                "inherits": ["AGE"],
                "ageBinWidth": "5",
            }
        ]

        review = validation_review_payload(payload)
        generated = dataset_payload(dataset_from_payload(payload, standardize=False))
        profile = next(profile for profile in review["profiles"] if profile["tag"] == "AGE")

        self.assertEqual(profile["distributions"][0]["label"], "5-year distribution")
        self.assertEqual(profile["distributions"][0]["bands"], [{"band": "1-4", "count": 2}, {"band": "5-9", "count": 1}])
        self.assertIn("AGE5", generated["commonTags"])
        self.assertEqual(generated["tagDefinitions"][0]["ageBinWidth"], 5)

    def test_reference_interval_payload_returns_generated_tame_dataset_with_log(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "ri_source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[TESTNAME]]검사항목명", "[[SEX]]성별", "[[AGE]]나이", "[[RESULT::NUM]]보고값", "[[UNIT]]단위"])
            sheet.append(["AST", "M", "30", "10", "U/L"])
            sheet.append(["AST", "F", "40", "20", "U/L"])
            sheet.append(["AST", "M", "50", "30", "U/L"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset, filename="ri_source.tame")

        generated = reference_interval_payload(payload, {})

        self.assertEqual(generated["filename"], "ri_source.ri_ep28.tame")
        self.assertIn("status", [column["name"] for column in generated["columns"]])
        self.assertTrue(generated["analysisTables"]["reference_intervals"])
        self.assertIn('OPERATION = "REFERENCE-INTERVAL"', generated["metaText"])

    def test_anonymize_payload_records_log(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "anon_source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["[[ID]]등록번호", "[[NAME]]이름", "[[RESULT::NUM]]보고값"])
            sheet.append(["P001", "홍길동", "10"])
            workbook.save(path)
            workbook.close()

            dataset, _warnings = convert_xlsx_selection(path, sheet_names=["DATA"], mode="single")
            payload = dataset_payload(dataset, filename="anon_source.tame")

        generated = anonymize_payload(payload, {"hashTags": "ID", "dropTags": "NAME", "salt": "s"})

        self.assertEqual(generated["filename"], "anon_source.anonymized.tame")
        self.assertIn('OPERATION = "ANONYMIZE"', generated["metaText"])


if __name__ == "__main__":
    unittest.main()
