from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import os
import tempfile
from typing import Any, Iterable
import warnings

import pandas as pd

from .config import ci_get
from .models import OperationOutput, TameDataset

_FONT_WARNING_EMITTED = False


def chart_spec(
    name: str,
    *,
    type: str = "bar",
    title: str = "",
    x: str,
    y: str,
    series: str = "",
    table: str = "",
    rows: list[dict[str, Any]] | None = None,
    max_points: int = 80,
) -> dict[str, Any]:
    """Return a portable chart declaration usable by web, CLI, and reports."""
    return {
        "name": str(name),
        "TYPE": str(type).upper(),
        "TITLE": str(title or name),
        "X": str(x),
        "Y": str(y),
        "SERIES": str(series or ""),
        "TABLE": str(table or ""),
        "ROWS": list(rows or []),
        "MAX_POINTS": int(max_points),
    }


def visualization_meta(charts: Iterable[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    meta: dict[str, dict[str, Any]] = {}
    for chart in charts:
        name = str(chart.get("name") or chart.get("NAME") or f"CHART_{len(meta) + 1}").strip() or f"CHART_{len(meta) + 1}"
        config = {
            "TYPE": str(chart.get("TYPE") or chart.get("type") or "BAR").upper(),
            "TITLE": str(chart.get("TITLE") or chart.get("title") or name),
            "X": str(chart.get("X") or chart.get("x") or ""),
            "Y": str(chart.get("Y") or chart.get("y") or ""),
            "MAX_POINTS": int(chart.get("MAX_POINTS") or chart.get("max_points") or 80),
        }
        series = str(chart.get("SERIES") or chart.get("series") or "").strip()
        if series:
            config["SERIES"] = series
        table = str(chart.get("TABLE") or chart.get("table") or "").strip()
        if table:
            config["TABLE"] = table
        for key in ("Y_LOW", "Y_HIGH", "X_LABEL", "Y_LABEL", "LOW_CI_LOW", "LOW_CI_HIGH", "HIGH_CI_LOW", "HIGH_CI_HIGH"):
            if key in chart:
                config[key] = str(chart[key])
        if chart.get("ROWS") and not chart.get("FILTER"):
            config["ROWS"] = deepcopy(chart["ROWS"])
        if chart.get("FILTER"):
            config["FILTER"] = deepcopy(chart["FILTER"])
        meta[name] = config
    return meta


def with_visualizations(dataset: TameDataset, charts: Iterable[dict[str, Any]] | None) -> TameDataset:
    chart_list = list(charts or [])
    if not chart_list:
        return dataset
    meta = deepcopy(dataset.meta)
    existing = ci_get(meta, "VISUALIZATIONS", {})
    if not isinstance(existing, dict):
        existing = {}
    merged = dict(existing)
    merged.update(visualization_meta(chart_list))
    meta["VISUALIZATIONS"] = merged
    return dataset.replace(meta=meta)


def output_tables(output: OperationOutput) -> dict[str, pd.DataFrame]:
    tables: dict[str, pd.DataFrame] = {}
    if output.table is not None:
        tables["main"] = output.table
    if output.dataset is not None:
        tables.setdefault("dataset", output.dataset.df)
    if output.tables:
        tables.update(output.tables)
    return tables


def write_docx_report(
    path: str | Path,
    *,
    title: str,
    outputs: Iterable[OperationOutput],
    charts: Iterable[dict[str, Any]] | None = None,
    summary: str = "",
    max_table_rows: int = 30,
) -> Path:
    """Write a docx report containing plugin tables and rendered charts.

    The imports are intentionally lazy so regular tametools use does not require
    report-generation dependencies until a report plugin calls this function.
    """
    try:
        from docx import Document  # noqa: PLC0415
        from docx.shared import Inches  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - depends on optional env
        raise RuntimeError(
            "docx report generation requires the report extra. Install with "
            "python3 -m pip install './tametools[report]' --no-build-isolation "
            "or use './tametools[report,web]' for the web workbench."
        ) from exc

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image_dir = output_path.parent / f"{output_path.stem}_charts"
    image_dir.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_heading(title, 0)
    if summary:
        document.add_paragraph(summary)

    outputs_list = list(outputs)
    for output in outputs_list:
        document.add_heading(output.name, level=1)
        if output.message:
            document.add_paragraph(str(output.message))
        for warning in output.warnings or []:
            document.add_paragraph(f"Warning: {warning}")
        for table_name, table in output_tables(output).items():
            _add_table(document, str(table_name), table, max_rows=max_table_rows)

    all_charts: list[dict[str, Any]] = []
    for output in outputs_list:
        all_charts.extend(output.charts or [])
    all_charts.extend(list(charts or []))

    if all_charts:
        document.add_heading("Charts", level=1)
    combined_tables: dict[str, pd.DataFrame] = {}
    for output in outputs_list:
        combined_tables.update(output_tables(output))
    for index, chart in enumerate(all_charts, start=1):
        image_path = render_chart_png(chart, combined_tables, image_dir / f"chart_{index:02d}.png")
        document.add_paragraph(str(chart.get("TITLE") or chart.get("title") or chart.get("name") or f"Chart {index}"))
        document.add_picture(str(image_path), width=Inches(6.2))

    document.save(output_path)
    return output_path


def render_chart_png(chart: dict[str, Any], tables: dict[str, pd.DataFrame], path: str | Path) -> Path:
    mpl_config_dir = Path(tempfile.gettempdir()) / "tametools-matplotlib"
    mpl_config_dir.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(mpl_config_dir))
    try:
        import matplotlib  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - depends on optional env
        raise RuntimeError(
            "chart rendering requires the report extra. Install with "
            "python3 -m pip install './tametools[report]' --no-build-isolation "
            "or use './tametools[report,web]' for the web workbench."
        ) from exc

    matplotlib.use("Agg")
    from matplotlib import font_manager  # noqa: PLC0415
    import matplotlib.pyplot as plt  # noqa: PLC0415

    _configure_matplotlib_font(plt, font_manager)
    frame = _chart_frame(chart, tables)
    x = str(chart.get("X") or chart.get("x") or "")
    y = str(chart.get("Y") or chart.get("y") or "")
    series = str(chart.get("SERIES") or chart.get("series") or "")
    chart_type = str(chart.get("TYPE") or chart.get("type") or "BAR").lower()
    title = str(chart.get("TITLE") or chart.get("title") or chart.get("name") or "Chart")
    max_points = int(chart.get("MAX_POINTS") or chart.get("max_points") or 80)
    for column, value in chart.get('FILTER', {}).items():
        if column not in frame:
            frame = frame.iloc[:0]; break
        matches = frame[column].map(lambda v: str(v).strip().lower() in {'true','1'}) if value is True else frame[column].map(lambda v: str(v).strip().lower() in {'false','0'}) if value is False else frame[column].eq(value)
        frame = frame.loc[matches]
    frame = frame.head(max_points).copy()

    fig, ax = plt.subplots(figsize=(8, 4.4))
    if frame.empty or x not in frame.columns or y not in frame.columns:
        ax.text(0.5, 0.5, "No chart data", ha="center", va="center")
        ax.set_axis_off()
    elif chart_type == "scatter":
        def scatter(part, label):
            numeric_x = pd.to_numeric(part[x], errors="coerce")
            axis_x = numeric_x if numeric_x.notna().all() else part[x].astype(str)
            ax.scatter(axis_x, pd.to_numeric(part[y], errors="coerce"), label=label, s=12, alpha=.45)
        _plot_by_series(ax, frame, x, y, series, scatter)
    elif chart_type == "interval":
        lower, upper = str(chart.get("Y_LOW", "ci95_low")), str(chart.get("Y_HIGH", "ci95_high"))
        center = pd.to_numeric(frame[y], errors="raise").to_numpy()
        low = pd.to_numeric(frame[lower], errors="raise").to_numpy()
        high = pd.to_numeric(frame[upper], errors="raise").to_numpy()
        ax.errorbar(frame[x].astype(str), center, yerr=[center-low, high-center], fmt="o", capsize=4)
    elif chart_type == "line":
        _plot_by_series(ax, frame, x, y, series, lambda part, label: ax.plot(part[x].astype(str), pd.to_numeric(part[y], errors="coerce"), marker="o", label=label))
    else:
        if series and series in frame.columns:
            pivot = frame.assign(**{y: pd.to_numeric(frame[y], errors="coerce")}).pivot_table(
                index=x,
                columns=series,
                values=y,
                aggfunc="first",
            )
            pivot.plot(kind="bar", ax=ax)
        else:
            values = pd.to_numeric(frame[y], errors="coerce")
            labels = frame[x].astype(str)
            ax.bar(labels, values)
    from textwrap import fill
    ax.set_title(fill(title, width=88))
    ax.set_xlabel(str(chart.get("X_LABEL", x)))
    ax.set_ylabel(str(chart.get("Y_LABEL", y)))
    ax.tick_params(axis="x", rotation=45)
    if series and series in frame.columns:
        handles, labels = ax.get_legend_handles_labels()
        if handles:
            ax.legend(handles, labels, loc="best")
    fig.tight_layout()
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=150)
    plt.close(fig)
    return output_path


def _configure_matplotlib_font(plt, font_manager) -> None:
    global _FONT_WARNING_EMITTED

    available = {font.name for font in font_manager.fontManager.ttflist}
    for candidate in (
        "NanumGothic",
        "NanumBarunGothic",
        "Noto Sans CJK KR",
        "Noto Sans CJK JP",
        "UnDotum",
        "Malgun Gothic",
        "AppleGothic",
    ):
        if candidate in available:
            plt.rcParams["font.family"] = candidate
            plt.rcParams["axes.unicode_minus"] = False
            return
    plt.rcParams["axes.unicode_minus"] = False
    if not _FONT_WARNING_EMITTED:
        warnings.warn(
            "No Korean-capable Matplotlib font was found. Korean labels in chart PNGs may render as square boxes. "
            "Install Malgun Gothic, NanumGothic, or Noto Sans CJK KR on the report machine.",
            RuntimeWarning,
            stacklevel=2,
        )
        _FONT_WARNING_EMITTED = True


def _chart_frame(chart: dict[str, Any], tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows = chart.get("ROWS") or chart.get("rows")
    if isinstance(rows, list) and rows:
        return pd.DataFrame(rows)
    table_name = str(chart.get("TABLE") or chart.get("table") or "").strip()
    if table_name and table_name in tables:
        return tables[table_name]
    if "main" in tables:
        return tables["main"]
    if "dataset" in tables:
        return tables["dataset"]
    return next(iter(tables.values()), pd.DataFrame())


def _plot_by_series(ax, frame: pd.DataFrame, x: str, y: str, series: str, plotter) -> None:
    if series and series in frame.columns:
        for label, part in frame.groupby(series, dropna=False, observed=False):
            plotter(part, str(label))
        return
    plotter(frame, None)


def _add_table(document, name: str, frame: pd.DataFrame, *, max_rows: int) -> None:
    document.add_heading(name, level=2)
    if frame is None or frame.empty:
        document.add_paragraph("(empty)")
        return
    preview = frame.head(max_rows)
    table = document.add_table(rows=1, cols=len(preview.columns))
    table.style = "Table Grid"
    for index, column in enumerate(preview.columns):
        table.rows[0].cells[index].text = str(column)
    for _, row in preview.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(preview.columns):
            cells[index].text = "" if pd.isna(row[column]) else str(row[column])
    if len(frame) > max_rows:
        document.add_paragraph(f"... {len(frame) - max_rows} more row(s)")
