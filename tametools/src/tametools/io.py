from __future__ import annotations

from copy import deepcopy
import csv
from dataclasses import dataclass
from io import StringIO
from pathlib import Path
import re
from typing import Any

import openpyxl
import pandas as pd

from .cellstate import CellTokenConfig, parse_serialized_cell, serialize_cell, token_config
from .config import ci_get, ci_get_nested
from .exceptions import TameFormatError, TameImportError
from .metadata import extract_column_tags, format_settings
from .merge import drop_absent_only_columns, merge_datasets
from .models import TameDataset, merged_column_specs
from .sex import standardize_sex_dataset
from .tag_placement import dataset_with_tags_in_headers, dataset_with_tags_in_meta
from .tags import merge_tags
from .toml_compat import dumps as dumps_toml
from .toml_compat import loads as loads_toml


SECTION_NAMES = ("SCHEMA", "JOB", "META", "DATA")
CONTROL_SECTION_NAMES = ("SCHEMA", "JOB", "META")
SECTION_RE_TEMPLATE = r"<\s*{name}\s*>(.*?)<\s*/\s*{name}\s*>"
DATA_SHEET_RE = re.compile(r"^DATA(?:_(\d+))?$", re.IGNORECASE)
CONTROL_SHEETS = {"META", "SCHEMA", "JOB"}


@dataclass(frozen=True)
class ControlBundle:
    meta: dict[str, Any]
    schema: dict[str, Any]
    job: dict[str, Any]
    raw_sections: dict[str, str]


def read_tame(path: str | Path, *, meta_path: str | Path | None = None) -> TameDataset:
    kind = _path_kind(path)
    if kind == "meta_tame":
        raise TameFormatError("Meta-only .meta.tame files cannot be loaded as datasets without a data file.")

    raw_sections = _read_tame_sections(path, allow_implicit_data=(kind == "data_tame"))
    if "DATA" not in raw_sections:
        raise TameFormatError("TAME dataset files require a DATA section.")

    sidecar = _load_sidecar_bundle(path, meta_path)
    current = _bundle_from_raw_sections({name: raw_sections[name] for name in CONTROL_SECTION_NAMES if name in raw_sections})
    merged = _merge_control_bundles(sidecar, current)
    return _dataset_from_parts(
        data_text=raw_sections["DATA"],
        controls=merged,
        source_path=str(path),
    )


def write_tame(
    path: str | Path,
    dataset: TameDataset,
    *,
    include_schema: bool = True,
    include_job: bool = True,
    tag_storage: str = "preserve",
) -> None:
    path = Path(path)
    dataset = standardize_sex_dataset(dataset)
    dataset = _dataset_for_tag_storage(dataset, tag_storage)
    include_header_tags = _include_header_tags(tag_storage)
    sections = _control_sections(dataset, include_schema=include_schema, include_job=include_job)
    sections.append(f"<DATA>\n{_write_tsv(dataset, include_tags=include_header_tags)}\n</DATA>")

    path.write_text("\n".join(sections) + "\n", encoding="utf-8")


def write_meta_tame(path: str | Path, dataset: TameDataset, *, include_schema: bool = True, include_job: bool = True) -> None:
    path = Path(path)
    meta = _materialized_meta_for_sidecar(dataset)
    sections = _control_sections(
        dataset,
        include_schema=include_schema,
        include_job=include_job,
        meta_text=dumps_toml(meta),
    )
    path.write_text("\n".join(sections) + "\n", encoding="utf-8")


def write_data_tame(path: str | Path, dataset: TameDataset) -> None:
    path = Path(path)
    dataset = standardize_sex_dataset(dataset)
    path.write_text(f"<DATA>\n{_write_tsv(dataset)}\n</DATA>\n", encoding="utf-8")


def import_xlsx_into_tame(template: TameDataset, xlsx_path: str | Path) -> TameDataset:
    controls_applied = _read_xlsx_with_controls(xlsx_path, _control_bundle_from_dataset(template))
    _validate_import_headers(template, controls_applied)

    return TameDataset(
        df=controls_applied.df.copy(),
        columns=deepcopy(template.columns),
        meta=deepcopy(controls_applied.meta),
        schema=deepcopy(controls_applied.schema),
        job=deepcopy(controls_applied.job),
        raw_sections=_control_raw_sections(template, controls_applied),
        source_path=str(xlsx_path),
    )


def read_xlsx(
    path: str | Path,
    *,
    meta_path: str | Path | None = None,
    source_column_name: str | None = None,
) -> TameDataset:
    workbook = openpyxl.load_workbook(path, data_only=True)
    meta_text = _sheet_text(workbook, "META")
    schema_text = _sheet_text(workbook, "SCHEMA")
    job_text = _sheet_text(workbook, "JOB")
    sidecar = _load_sidecar_bundle(path, meta_path)
    workbook_bundle = _bundle_from_raw_sections(
        {
            key: value
            for key, value in {"META": meta_text, "SCHEMA": schema_text, "JOB": job_text}.items()
            if value.strip()
        }
    )
    merged_controls = _merge_control_bundles(sidecar, workbook_bundle)

    data_sheet_names = _data_sheet_names(workbook)
    if not data_sheet_names:
        data_sheet_names = [workbook.active.title]

    datasets = [
        _sheet_dataset(
            workbook,
            sheet_name,
            meta=merged_controls.meta,
            schema=merged_controls.schema,
            job=merged_controls.job,
            source_path=str(path),
        )
        for sheet_name in data_sheet_names
    ]
    meta = merged_controls.meta
    if len(datasets) == 1:
        frame = datasets[0].df
        columns = datasets[0].columns
    else:
        sheet_source_column = _source_column_name(merged_controls.meta, source_column_name)
        merged = merge_datasets(
            datasets,
            source_labels=data_sheet_names,
            add_source_column=True,
            source_column_name=sheet_source_column,
            source_column_tags=("SHEET", "STR"),
            numeric_conflict="promote",
            prefer_tag_names=True,
        ).dataset
        frame = merged.df
        columns = merged.columns
        meta = _augment_multisheet_meta(merged_controls.meta, data_sheet_names, source_column_name=sheet_source_column)
        merged_controls = _merge_control_bundles(
            None,
            ControlBundle(
                meta=meta,
                schema=merged_controls.schema,
                job=merged_controls.job,
                raw_sections=_serialize_control_sections(meta, merged_controls.schema, merged_controls.job),
            ),
        )

    return TameDataset(
        df=frame,
        columns=columns,
        meta=meta,
        schema=merged_controls.schema,
        job=merged_controls.job,
        raw_sections=merged_controls.raw_sections,
        source_path=str(path),
    )


def write_xlsx(
    path: str | Path,
    dataset: TameDataset,
    *,
    include_schema: bool = True,
    include_job: bool = True,
    tag_storage: str = "preserve",
) -> None:
    dataset = standardize_sex_dataset(dataset)
    dataset = _dataset_for_tag_storage(dataset, tag_storage)
    include_header_tags = _include_header_tags(tag_storage)
    workbook = openpyxl.Workbook()
    sheet_column = dataset.first_column_with_tag("SHEET")
    if sheet_column is None:
        data_sheet = workbook.active
        data_sheet.title = "DATA"
        _write_data_sheet(data_sheet, dataset, include_tags=include_header_tags)
    else:
        groups = _sheet_groups(dataset, sheet_column.name)
        first_sheet = workbook.active
        first_sheet.title = _excel_sheet_name(groups[0][0], used_names=set())
        group_name, group_dataset = groups[0]
        _write_data_sheet(first_sheet, group_dataset, include_tags=include_header_tags)

        used_names = {first_sheet.title}
        for raw_name, group_dataset in groups[1:]:
            sheet_name = _excel_sheet_name(raw_name, used_names=used_names)
            used_names.add(sheet_name)
            sheet = workbook.create_sheet(sheet_name)
            _write_data_sheet(sheet, group_dataset, include_tags=include_header_tags)

    _write_text_sheet(workbook, "META", dumps_toml(dataset.meta))
    if include_schema and (dataset.raw_sections.get("SCHEMA") or dataset.schema):
        _write_text_sheet(workbook, "SCHEMA", dataset.raw_sections.get("SCHEMA") or dumps_toml(dataset.schema))
    if include_job and (dataset.raw_sections.get("JOB") or dataset.job):
        _write_text_sheet(workbook, "JOB", dataset.raw_sections.get("JOB") or dumps_toml(dataset.job))

    workbook.save(path)


def _extract_section(text: str, name: str) -> str | None:
    match = re.search(SECTION_RE_TEMPLATE.format(name=name), text, flags=re.DOTALL | re.IGNORECASE)
    if not match:
        return None
    return match.group(1).strip()


def _path_kind(path: str | Path) -> str:
    lowered = str(path).lower()
    if lowered.endswith(".meta.tame"):
        return "meta_tame"
    if lowered.endswith(".data.tame"):
        return "data_tame"
    suffix = Path(path).suffix.lower()
    if suffix == ".tame":
        return "tame"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    raise TameFormatError(f"Unsupported input format: {path}")


def _read_tame_sections(
    path: str | Path,
    *,
    allow_implicit_data: bool = False,
    allow_implicit_meta: bool = False,
) -> dict[str, str]:
    text = Path(path).read_text(encoding="utf-8")
    raw_sections = {name: value for name in SECTION_NAMES if (value := _extract_section(text, name)) is not None}
    if raw_sections:
        return raw_sections
    stripped = text.strip()
    if allow_implicit_data and stripped:
        return {"DATA": stripped}
    if allow_implicit_meta and stripped:
        return {"META": stripped}
    return {}


def _bundle_from_raw_sections(raw_sections: dict[str, str]) -> ControlBundle:
    meta_text = raw_sections.get("META", "")
    schema_text = raw_sections.get("SCHEMA", "")
    job_text = raw_sections.get("JOB", "")
    return ControlBundle(
        meta=loads_toml(meta_text),
        schema=loads_toml(schema_text),
        job=loads_toml(job_text),
        raw_sections={key: value for key, value in raw_sections.items() if key in CONTROL_SECTION_NAMES and value.strip()},
    )


def _load_control_bundle(path: str | Path) -> ControlBundle:
    kind = _path_kind(path)
    if kind == "xlsx":
        workbook = openpyxl.load_workbook(path, data_only=True)
        return _bundle_from_raw_sections(
            {
                key: value
                for key, value in {
                    "META": _sheet_text(workbook, "META"),
                    "SCHEMA": _sheet_text(workbook, "SCHEMA"),
                    "JOB": _sheet_text(workbook, "JOB"),
                }.items()
                if value.strip()
            }
        )
    raw_sections = _read_tame_sections(path, allow_implicit_meta=(kind == "meta_tame"))
    return _bundle_from_raw_sections(raw_sections)


def _control_bundle_from_dataset(dataset: TameDataset) -> ControlBundle:
    raw_sections = _serialize_control_sections(dataset.meta, dataset.schema, dataset.job)
    for name in CONTROL_SECTION_NAMES:
        text = dataset.raw_sections.get(name, "")
        if text.strip():
            raw_sections[name] = text
    return ControlBundle(
        meta=deepcopy(dataset.meta),
        schema=deepcopy(dataset.schema),
        job=deepcopy(dataset.job),
        raw_sections=raw_sections,
    )


def _load_sidecar_bundle(path: str | Path, meta_path: str | Path | None) -> ControlBundle | None:
    resolved = _resolve_meta_path(path, meta_path)
    if resolved is None:
        return None
    return _load_control_bundle(resolved)


def _resolve_meta_path(path: str | Path, meta_path: str | Path | None) -> str | None:
    if meta_path is not None:
        return str(meta_path)

    kind = _path_kind(path)
    candidate: Path | None = None
    source = Path(path)
    if kind == "xlsx":
        candidate = source.with_suffix("")
        candidate = candidate.parent / f"{candidate.name}.meta.tame"
    elif kind == "data_tame":
        name = source.name[: -len(".data.tame")]
        candidate = source.parent / f"{name}.meta.tame"
    if candidate is not None and candidate.exists():
        return str(candidate)
    return None


def _merge_control_bundles(base: ControlBundle | None, override: ControlBundle | None) -> ControlBundle:
    base = base or ControlBundle(meta={}, schema={}, job={}, raw_sections={})
    override = override or ControlBundle(meta={}, schema={}, job={}, raw_sections={})
    meta = _merge_nested_dict(base.meta, override.meta)
    schema = _merge_nested_dict(base.schema, override.schema)
    job = _merge_nested_dict(base.job, override.job)
    return ControlBundle(meta=meta, schema=schema, job=job, raw_sections=_serialize_control_sections(meta, schema, job))


def _control_raw_sections(template: TameDataset, loaded: TameDataset) -> dict[str, str]:
    raw_sections = _control_bundle_from_dataset(template).raw_sections
    if loaded.meta != template.meta:
        if loaded.meta:
            raw_sections["META"] = dumps_toml(loaded.meta)
        else:
            raw_sections.pop("META", None)
    if loaded.schema != template.schema:
        if loaded.schema:
            raw_sections["SCHEMA"] = dumps_toml(loaded.schema)
        else:
            raw_sections.pop("SCHEMA", None)
    if loaded.job != template.job:
        if loaded.job:
            raw_sections["JOB"] = dumps_toml(loaded.job)
        else:
            raw_sections.pop("JOB", None)
    return raw_sections


def _merge_nested_dict(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(base)
    for key, value in override.items():
        if key in result and isinstance(result[key], dict) and isinstance(value, dict):
            result[key] = _merge_nested_dict(result[key], value)
        else:
            result[key] = deepcopy(value)
    return result


def _serialize_control_sections(meta: dict[str, Any], schema: dict[str, Any], job: dict[str, Any]) -> dict[str, str]:
    sections: dict[str, str] = {}
    if schema:
        sections["SCHEMA"] = dumps_toml(schema)
    if job:
        sections["JOB"] = dumps_toml(job)
    if meta:
        sections["META"] = dumps_toml(meta)
    return sections


def _control_sections(
    dataset: TameDataset,
    *,
    include_schema: bool = True,
    include_job: bool = True,
    meta_text: str | None = None,
) -> list[str]:
    sections: list[str] = []
    if include_schema and (dataset.raw_sections.get("SCHEMA") or dataset.schema):
        schema_text = dataset.raw_sections.get("SCHEMA") or dumps_toml(dataset.schema)
        sections.append(_wrapped_section("SCHEMA", schema_text))
    if include_job and (dataset.raw_sections.get("JOB") or dataset.job):
        job_text = dataset.raw_sections.get("JOB") or dumps_toml(dataset.job)
        sections.append(_wrapped_section("JOB", job_text))
    resolved_meta_text = meta_text if meta_text is not None else dumps_toml(dataset.meta)
    sections.append(_wrapped_section("META", resolved_meta_text))
    return sections


def _wrapped_section(name: str, text: str) -> str:
    return f"<{name}>\n{text.rstrip()}\n</{name}>"


def _dataset_for_tag_storage(dataset: TameDataset, tag_storage: str) -> TameDataset:
    normalized = _normalize_tag_storage(tag_storage)
    if normalized == "preserve":
        return dataset
    if normalized == "meta":
        return dataset_with_tags_in_meta(dataset)
    if normalized == "header":
        return dataset_with_tags_in_headers(dataset)[0]
    raise ValueError(f"Unsupported tag_storage: {tag_storage}")


def _include_header_tags(tag_storage: str) -> bool:
    return _normalize_tag_storage(tag_storage) != "meta"


def _normalize_tag_storage(tag_storage: str) -> str:
    normalized = str(tag_storage or "preserve").strip().lower().replace("-", "_")
    aliases = {
        "preserve": "preserve",
        "both": "preserve",
        "meta": "meta",
        "metadata": "meta",
        "header": "header",
        "headers": "header",
        "data": "header",
    }
    if normalized not in aliases:
        raise ValueError("tag_storage must be one of: preserve, meta, header")
    return aliases[normalized]


def _dataset_from_parts(*, data_text: str, controls: ControlBundle, source_path: str) -> TameDataset:
    df = _read_tsv(data_text, _header_rows(controls.meta), _cell_settings(controls.meta))
    columns = merged_column_specs(
        headers=list(df.columns),
        meta_tags=_extract_meta_tags(controls.meta),
        schema_tags=_extract_schema_tags(controls.schema),
    )
    df.columns = [column.name for column in columns]
    raw_sections = dict(controls.raw_sections)
    raw_sections["DATA"] = data_text.strip()
    return TameDataset(
        df=df,
        columns=columns,
        meta=controls.meta,
        schema=controls.schema,
        job=controls.job,
        raw_sections=raw_sections,
        source_path=source_path,
    )


def _materialized_meta_for_sidecar(dataset: TameDataset) -> dict[str, Any]:
    meta = deepcopy(dataset.meta)
    existing = _extract_meta_tags(meta)
    for key in list(meta):
        if str(key).upper() == "TAGS":
            meta.pop(key, None)
    column_section = ci_get(meta, "COLUMN", {})
    column_section = dict(column_section) if isinstance(column_section, dict) else {}
    for column in dataset.columns:
        tags = merge_tags(existing.get(column.name, ()), column.tags)
        if tags:
            entry = column_section.get(column.name, {})
            entry = dict(entry) if isinstance(entry, dict) else {}
            entry["TAGS"] = list(tags)
            column_section[column.name] = entry
    for name, tags in existing.items():
        if name not in column_section and tags:
            column_section[name] = {"TAGS": list(tags)}
    if column_section:
        meta["COLUMN"] = column_section
    return meta


def _validate_import_headers(template: TameDataset, incoming: TameDataset) -> None:
    expected = template.columns
    actual = incoming.columns
    errors: list[str] = []

    if len(expected) != len(actual):
        errors.append(f"Column count mismatch: expected {len(expected)} columns but got {len(actual)}.")

    width = min(len(expected), len(actual))
    for index in range(width):
        expected_column = expected[index]
        actual_column = actual[index]
        if expected_column.name != actual_column.name:
            errors.append(
                f"Column {index + 1} header mismatch: expected '{expected_column.name}' but got '{actual_column.name}'."
            )
            continue
        if actual_column.tags:
            actual_tags = set(actual_column.tags)
            expected_tags = set(expected_column.tags)
            if not actual_tags.issubset(expected_tags):
                errors.append(
                    f"Column {index + 1} tag mismatch for '{actual_column.name}': incoming tags {sorted(actual_tags)} are not compatible with template tags {sorted(expected_tags)}."
                )

    if errors:
        raise TameImportError("XLSX import header validation failed.\n" + "\n".join(errors))


def _read_xlsx_with_controls(path: str | Path, controls: ControlBundle) -> TameDataset:
    workbook = openpyxl.load_workbook(path, data_only=True)
    meta = deepcopy(controls.meta)
    schema = deepcopy(controls.schema)
    job = deepcopy(controls.job)

    data_sheet_names = _data_sheet_names(workbook)
    if not data_sheet_names:
        data_sheet_names = [workbook.active.title]

    datasets = [
        _sheet_dataset(
            workbook,
            sheet_name,
            meta=meta,
            schema=schema,
            job=job,
            source_path=str(path),
        )
        for sheet_name in data_sheet_names
    ]
    if len(datasets) == 1:
        frame = datasets[0].df
        columns = datasets[0].columns
    else:
        merged = merge_datasets(
            datasets,
            source_labels=data_sheet_names,
            add_source_column=True,
            source_column_name="시트명",
            source_column_tags=("SHEET", "STR"),
            numeric_conflict="promote",
            prefer_tag_names=True,
        ).dataset
        frame = merged.df
        columns = merged.columns
        meta = _augment_multisheet_meta(meta, data_sheet_names)

    return TameDataset(
        df=frame,
        columns=columns,
        meta=meta,
        schema=schema,
        job=job,
        raw_sections=_serialize_control_sections(meta, schema, job),
        source_path=str(path),
    )


def _header_rows(meta: dict[str, Any]) -> int:
    rows = ci_get_nested(meta, "DATA", "headers", default=None)
    if rows is None:
        rows = ci_get_nested(meta, "data", "headers", default=1)
    try:
        return int(rows)
    except (TypeError, ValueError) as exc:
        raise TameFormatError(f"META[DATA].headers must be an integer, got {rows!r}.") from exc


def _cell_settings(meta: dict[str, Any]) -> dict[str, Any]:
    return format_settings(meta)


def _source_column_name(meta: dict[str, Any], override: str | None) -> str:
    if override and str(override).strip():
        return str(override).strip()
    multisheet = ci_get(meta, "MULTISHEET", {})
    if isinstance(multisheet, dict):
        configured = ci_get(multisheet, "SOURCE_COLUMN_NAME", ci_get(multisheet, "COLUMN", ""))
        if str(configured).strip():
            return str(configured).strip()
    return "시트명"


def _read_tsv(text: str, header_rows: int, cell_settings: dict[str, Any]) -> pd.DataFrame:
    reader = csv.reader(StringIO(text), delimiter="\t", quotechar='"')
    rows = [row for row in reader]
    headers, data_rows = _split_header_rows(rows, header_rows)
    cfg = token_config(cell_settings)
    parsed_rows = [[parse_serialized_cell(value, cfg) for value in row] for row in data_rows]
    return _frame_from_rows(headers, parsed_rows)


def _write_tsv(dataset: TameDataset, *, include_tags: bool = True) -> str:
    buffer = StringIO()
    writer = csv.writer(buffer, delimiter="\t", quotechar='"', lineterminator="\n")
    writer.writerow(dataset.tagged_headers() if include_tags else [column.name for column in dataset.columns])
    cfg = token_config(_cell_settings(dataset.meta))
    for row in dataset.df.itertuples(index=False):
        writer.writerow([_tsv_cell_value(value, cfg) for value in row])
    return buffer.getvalue().rstrip("\n")


def _split_header_rows(rows: list[list[Any]], header_rows: int) -> tuple[list[str], list[list[Any]]]:
    if not rows:
        return [], []

    max_len = max(len(row) for row in rows)
    padded_rows = [list(row) + [None] * (max_len - len(row)) for row in rows]
    if header_rows <= 0:
        headers = [f"column_{idx + 1}" for idx in range(max_len)]
        return headers, padded_rows

    header_parts = padded_rows[:header_rows]
    headers: list[str] = []
    for idx in range(max_len):
        pieces = [str(row[idx]).strip() for row in header_parts if row[idx] not in (None, "")]
        headers.append("__".join(pieces) if pieces else f"column_{idx + 1}")
    return headers, padded_rows[header_rows:]


def _frame_from_rows(headers: list[str], data_rows: list[list[Any]]) -> pd.DataFrame:
    width = len(headers)
    normalized_rows = [list(row)[:width] + [None] * max(0, width - len(row)) for row in data_rows]
    return pd.DataFrame(normalized_rows, columns=headers, dtype=object)


def _extract_meta_tags(meta: dict[str, Any]) -> dict[str, tuple[str, ...]]:
    return extract_column_tags(meta)


def _extract_schema_tags(schema: dict[str, Any]) -> dict[str, tuple[str, ...]]:
    section = ci_get(schema, "column", {})
    result: dict[str, tuple[str, ...]] = {}
    if not isinstance(section, dict):
        return result
    for name, entry in section.items():
        if isinstance(entry, dict):
            tags = entry.get("tags", [])
            if isinstance(tags, list):
                result[str(name)] = merge_tags(tags)
    return result


def _sheet_text(workbook: openpyxl.Workbook, sheet_name: str) -> str:
    if sheet_name not in workbook.sheetnames:
        return ""
    values: list[str] = []
    for row in workbook[sheet_name].iter_rows(values_only=True):
        value = row[0] if row else None
        values.append("" if value is None else str(value))
    return "\n".join(values).strip()


def _write_text_sheet(workbook: openpyxl.Workbook, sheet_name: str, text: str) -> None:
    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row in range(1, sheet.max_row + 1):
            sheet.cell(row=row, column=1, value=None)
    else:
        sheet = workbook.create_sheet(sheet_name)
    for row_idx, line in enumerate(text.splitlines(), start=1):
        sheet.cell(row=row_idx, column=1, value=line)


def _sheet_rows(sheet: openpyxl.worksheet.worksheet.Worksheet, cell_settings: dict[str, Any]) -> list[list[Any]]:
    cfg = token_config(cell_settings)
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if row is None:
            continue
        rows.append([parse_serialized_cell(value, cfg) for value in row])
    return rows


def _data_sheet_names(workbook: openpyxl.Workbook) -> list[str]:
    matches: list[tuple[int, str]] = []
    for name in workbook.sheetnames:
        match = DATA_SHEET_RE.match(name)
        if match:
            order = int(match.group(1) or 0)
            matches.append((order, name))
    if matches:
        return [name for _, name in sorted(matches, key=lambda item: item[0])]
    return [name for name in workbook.sheetnames if name.upper() not in CONTROL_SHEETS]


def _tsv_cell_value(value: Any, cfg: CellTokenConfig) -> Any:
    return serialize_cell(value, cfg, for_excel=False)


def _xlsx_cell_value(value: Any, cfg: CellTokenConfig) -> Any:
    return serialize_cell(value, cfg, for_excel=True)


def _sheet_dataset(
    workbook: openpyxl.Workbook,
    sheet_name: str,
    *,
    meta: dict[str, Any],
    schema: dict[str, Any],
    job: dict[str, Any],
    source_path: str,
) -> TameDataset:
    rows = _sheet_rows(workbook[sheet_name], _cell_settings(meta))
    headers, data_rows = _split_header_rows(rows, _header_rows(meta))
    frame = _frame_from_rows(headers, data_rows)
    columns = merged_column_specs(
        headers=headers,
        meta_tags=_extract_meta_tags(meta),
        schema_tags=_extract_schema_tags(schema),
    )
    frame.columns = [column.name for column in columns]
    return TameDataset(
        df=frame,
        columns=columns,
        meta=meta,
        schema=schema,
        job=job,
        raw_sections={},
        source_path=source_path,
    )


def _augment_multisheet_meta(meta: dict[str, Any], sheet_names: list[str], *, source_column_name: str) -> dict[str, Any]:
    updated = dict(meta)
    updated["MULTISHEET"] = {
        "COLUMN": source_column_name,
        "SOURCE_COLUMN_NAME": source_column_name,
        "SHEETS": list(sheet_names),
    }
    return updated


def _write_data_sheet(sheet: openpyxl.worksheet.worksheet.Worksheet, dataset: TameDataset, *, include_tags: bool = True) -> None:
    headers = dataset.tagged_headers() if include_tags else [column.name for column in dataset.columns]
    for col_idx, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=col_idx, value=header)

    cfg = token_config(_cell_settings(dataset.meta))
    for row_idx, row in enumerate(dataset.df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=_xlsx_cell_value(value, cfg))


def _sheet_groups(dataset: TameDataset, sheet_column_name: str) -> list[tuple[str, TameDataset]]:
    groups: list[tuple[str, TameDataset]] = []
    seen: list[str] = []
    series = dataset.df[sheet_column_name]
    for value in series.tolist():
        name = "DATA" if value in (None, "") else str(value)
        if name not in seen:
            seen.append(name)

    for name in seen:
        if name == "DATA":
            mask = series.isna() | (series == "")
        else:
            mask = series == name
        frame = dataset.df.loc[mask].reset_index(drop=True).copy()
        columns = [column for column in dataset.columns if column.name != sheet_column_name]
        frame = frame[[column.name for column in columns]]
        frame, columns = drop_absent_only_columns(frame, columns)
        groups.append((name, dataset.replace(df=frame, columns=columns)))
    return groups


def _excel_sheet_name(raw_name: str, *, used_names: set[str]) -> str:
    text = re.sub(r"[:\\\\/?*\\[\\]]", "_", str(raw_name)).strip() or "DATA"
    text = text[:31]
    candidate = text
    suffix = 1
    while candidate in used_names:
        suffix_text = f"_{suffix}"
        candidate = (text[: 31 - len(suffix_text)] + suffix_text).strip() or f"DATA{suffix_text}"
        suffix += 1
    return candidate
