"""Korean, self-contained study reports and complete machine-readable audit tables."""
from __future__ import annotations

import base64
from datetime import datetime, timezone
from hashlib import sha256
from html import escape
from importlib import metadata
import json
from pathlib import Path

import numpy as np
import pandas as pd


CAUTIONS = [
    "참고구간은 정의된 참고집단의 분포를 기술합니다. 질병 진단·치료 결정을 위한 임상 결정한계와 구분해야 합니다.",
    "참고 대상자의 건강 상태, 선정·제외 기준, 검체 종류, 채혈 시간·공복·자세·운동·약물, 보관·운송 조건과 분석법·보정·QC를 확인해야 합니다. 통계 계산만으로 이 조건을 입증할 수 없습니다.",
    "비모수 중앙 95% 구간은 분할 집단마다 적합한 독립 참고 대상자 120명 이상을 권장합니다. 제외한 관측값이 있으면 필요한 대상자 수를 보충합니다. n은 결측·검열·이상치 처리 후 분모와 함께 확인합니다.",
    "EP28 9.1은 robust 방법의 고정 최소 표본 수를 정하지 않습니다. 작은 표본에서 계산된다는 사실은 신뢰성을 보장하지 않습니다. 부록 B의 20명 예제는 계산 설명용이며 해당 크기의 연구를 권장하는 예제가 아닙니다.",
    "90% 참고한계 CI 폭이 RI 폭의 0.2보다 작도록 권고한 기준을 기본 검토값으로 사용합니다(EP28 9.1의 Harris–Boyd 권고). 사용자가 변경한 임계값은 설정표에 남습니다. 참고구간과 참고한계의 신뢰구간은 서로 다릅니다.",
    "이상치는 오류나 다른 모집단의 관측인지 확인합니다. 원인을 모르는 극단값의 기계적 제거는 생리적인 꼬리를 잘라낼 수 있습니다. 기본 FLAG는 값을 유지하며 REMOVE는 제외 기록과 민감도 분석을 함께 검토해야 합니다.",
    "REED는 D/R ≥ 1/3 규칙이며, R outliers 패키지의 Dixon Q 검정(3–30개, n별 통계량·임계값)과 다릅니다. 반복/블록 검출은 마스킹에 대응하지만 과도한 꼬리 제거도 점검해야 합니다.",
    "정규·로그정규·Box-Cox 모수법은 변환 후의 분포 가정을 점검합니다. Shapiro–Wilk 비유의는 정규성의 증명이 아닙니다. 기본 robust는 부록 B의 반복 추정법이며 비대칭 분포에서는 변환·대안 추정과 꼬리 민감도를 검토합니다.",
    "Harris–Boyd는 정규성 또는 적절한 변환과 집단 구성에 의존합니다. p값의 유의성만으로 분할하지 않습니다. 성별·연령 등 분할은 생리학적 근거와 임상적 유용성을 사전에 정하며, 여러 집단 비교에는 다중 비교와 작은 분모 문제가 있습니다.",
    "공통 구간 밖 각 꼬리의 비율과 이항 신뢰구간을 제공합니다. Lahti 4.1%/3.2% 기준은 중앙 95% 구간에서의 보조 검토 신호이며, 경험적 소표본 비율만으로 통합 또는 분할을 확정하지 않습니다. 집단의 상대 빈도도 영향을 줍니다.",
    "검출한계 미만·초과 값을 경계값이나 공개 대체값으로 바꾸거나 제외하면 참고한계가 편향될 수 있습니다. 그러한 결과는 검열 민감도 분석으로 표시하며 실제 농도가 복원된 것으로 해석하지 않습니다.",
    "20명 검증은 사전에 정해진 외부 참고구간·비교 가능한 방법과 모집단에 적용합니다. 확인된 이상치를 제외했다면 새 대상자로 보충하여 적합한 20개 결과를 확보합니다. 첫 20명 중 구간 밖 0–2명이면 기준 충족, 3–4명이면 새 20명, 5명 이상이면 재검토합니다. 두 번째 20명은 0–2명만 기준을 충족합니다.",
    "검증용 20명을 보고 유리한 대상자나 구간을 사후 선택하면 안 됩니다. 두 번째 20명은 첫 집단과 독립적이어야 합니다. 반복 측정은 사람 수를 늘리지 않습니다. 조사표본 가중치를 사용하는 모집단 추정은 이 직접 참고구간 모듈의 기능이 아닙니다.",
    "LAVE는 다른 검사값으로 대상자를 선별합니다. 대상 검사 자체는 선별에 쓰지 않으며, 참조 검사·허용 이상 수·구간 확장·결측 처리·성별 층화·반복 수를 명시합니다. 개별 검사 이상치 제거와 별개이고, 제외가 임상적으로 정당한지 검토해야 합니다.",
    "이동 Box–Cox의 원점·power 추정은 선언한 탐색 범위에 의존합니다. 경계 도달과 변환 후 정규성을 확인하십시오.",
    "BOOTSTRAP_MEAN은 재표집 참고한계의 평균을 점추정으로 보고합니다. 재표집 횟수가 적으면 참고한계와 CI가 불안정할 수 있습니다. R과 Python은 같은 seed만으로 동일 재표집이 되지 않습니다. R_BOOT와 TYPE7의 CI 끝점 보간도 구분합니다.",
    "Wilcoxon 검정은 분포 차이의 보조 지표입니다. SDR 0.4, BR 0.375, 보고 단위 3배는 케냐 연구의 검토 기준이며 CLSI의 보편적 자동 분할 규칙이 아닙니다. nested SDR의 요인 순서·불균형 처리도 기록해야 합니다.",
    "부트스트랩은 현재 선택된 집단과 이상치 처리 후 표본에 조건부입니다. 인구 선정·이상치 선택·partition 선택의 전체 불확실성을 포함하지 않습니다. Box-Cox 변환은 각 재표집에서 다시 적합하며, 실패한 재표집 건수를 기록합니다.",
]

SOURCES = [
    ("Omuse et al. 2020: Kenya RI methods","https://doi.org/10.1371/journal.pone.0235234"),
    ("Ichihara et al. 2017: modified Box-Cox and LAVE","https://doi.org/10.1016/j.cca.2016.09.016"),
    ("Abebe et al. 2018: Ethiopia RI methods","https://doi.org/10.1371/journal.pone.0201782"),
    ("CLSI EP28-A3c: 9.1–9.5, 11.2, Appendix B","https://clsi.org/shop/standards/ep28/"),
    ("R referenceIntervals 1.3.1 공식 명세·소스","https://cran.r-project.org/package=referenceIntervals"),
    ("Horn, Pesce & Copeland, 1998: robust reference interval","https://pubmed.ncbi.nlm.nih.gov/9510871/"),
    ("Harris & Boyd, 1990: partitioning","https://pubmed.ncbi.nlm.nih.gov/2302771/"),
    ("Lahti et al., 2002: subgroup tail criteria","https://pubmed.ncbi.nlm.nih.gov/11805016/"),
    ("Lahti et al., 2004: non-Gaussian partitioning","https://pubmed.ncbi.nlm.nih.gov/15010425/"),
]


def _safe(value):
    if isinstance(value,dict):return {str(k):_safe(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)):return [_safe(v) for v in value]
    if isinstance(value,np.ndarray):return _safe(value.tolist())
    if isinstance(value,np.generic):return _safe(value.item())
    if isinstance(value,float) and not np.isfinite(value):return None
    if isinstance(value,(str,int,float,bool)) or value is None:return value
    return str(value)


def _text(value):
    if value is None or (isinstance(value,float) and not np.isfinite(value)):return "—"
    if isinstance(value,float):return f"{value:.6g}"
    return str(value)


def _group_label(value):
    if value == 'ALL': return '전체'
    try:
        return ' · '.join(f"{ {'SEX':'성별','AGE':'연령'}.get(k,k)}: { {'male':'남성','female':'여성'}.get(v,v)}" for k,v in json.loads(value).items())
    except (ValueError, TypeError, AttributeError): return str(value)


def _display(frame,columns,maximum,labels=None):
    frame = frame.copy()
    if 'test_name' in frame and 'result_id' in frame and labels:
        frame['test_name'] = [labels.get(i,n) if n == i else n for i,n in zip(frame.result_id,frame.test_name)]
    if 'unit' in frame: frame['unit'] = frame['unit'].map(lambda x:str(x).replace('umol/L','µmol/L'))
    if 'group' in frame: frame['group'] = frame['group'].map(_group_label)
    columns=[c for c in columns if c in frame]
    return frame[columns].head(maximum).apply(lambda column:column.map(_text)) if columns else pd.DataFrame()


def _interpretation(tables,cfg):
    main=tables['reference_intervals']
    if main.empty:return ['선정 조건을 충족하는 정량 결과가 없어 참고구간을 산출하지 못했습니다. 입력·제외 분모를 확인하십시오.']
    messages=[]
    if not cfg['POPULATION'].get('REFERENCE_INDIVIDUALS_CONFIRMED',False):
        messages.append('이 자료는 건강 참고집단으로 확인되지 않았습니다. 아래 수치는 탐색적 후보 구간이며 검사실 운영 참고구간으로 바로 채택할 수 없습니다.')
    small=int(main.n.lt(cfg['MIN_N']).sum())
    messages.append(f"방법·집단 조합 {len(main)}개 중 {small}개가 처리 후 n<{cfg['MIN_N']}입니다. 같은 대상자가 전체·분할·방법 비교에 반복 표시되므로 표의 n을 합산하지 마십시오.")
    for status,label in [('not_evaluable','수치 산출 또는 신뢰구간 계산 불가'),('outlier_method_not_evaluable','선택한 이상치 방법 적용 불가'),
                         ('insufficient_order_resolution','요청 백분위의 순위 해상도 부족'),('censoring_sensitivity_only','검열 처리로 꼬리 편향 가능'),
                         ('degenerate_distribution','퇴화한 분포')]:
        count=int(main.status.eq(status).sum())
        if count:messages.append(f'{label}: {count}개 결과입니다. 해당 행의 사유를 확인하고 자료·방법을 재검토하십시오.')
    if 'max_ci_width_ratio' in main:
        count=int(pd.to_numeric(main.max_ci_width_ratio,errors='coerce').gt(cfg['CI_WIDTH_WARN_RATIO']).sum())
        if count:messages.append(f"참고한계 CI 폭이 구간 폭의 {cfg['CI_WIDTH_WARN_RATIO']:.0%}를 넘는 결과가 {count}개입니다. 분할 집단별 추가 대상자 확보와 정밀도를 검토하십시오.")
    if 'shapiro_p' in main:
        mask=main.method.isin(['PARAMETRIC','LOG_PARAMETRIC','BOXCOX_PARAMETRIC','BOXCOX_SHIFTED_PARAMETRIC']) & pd.to_numeric(main.shapiro_p,errors='coerce').lt(.05)
        if mask.any():messages.append(f'모수법 {int(mask.sum())}개 결과에서 적합 척도의 정규성 검토 신호가 있습니다. 그래프와 꼬리를 함께 확인하고 사전에 정한 비모수·변환·robust 대안과 비교하십시오.')
    negative=int(main.get('notes',pd.Series(dtype=str)).fillna('').str.contains('Negative lower limit',regex=False).sum())
    if negative:messages.append(f'음수가 없는 입력에서 음수 하한을 산출한 결과가 {negative}개입니다. 해당 검사의 물리적 허용 범위와 분포 가정을 확인하십시오. 하한을 임의로 0으로 바꾸지는 않았습니다.')
    partition=tables['partition_tests']
    if len(partition) and 'statistical_signal' in partition:
        count=int(partition.statistical_signal.fillna(False).astype(bool).sum())
        messages.append(f'Harris–Boyd 비교 {len(partition)}개 중 {count}개에서 분할 검토 신호가 있습니다. 생리학적 근거, 분포 가정, 각 집단의 n과 공통 구간 밖 꼬리 비율을 검토한 뒤 분할 여부를 결정하십시오.')
    verification=tables['verification']
    if len(verification):
        decisions=verification.get('decision',pd.Series(dtype=str))
        passed=int(decisions.eq('verification_criterion_met').sum())
        messages.append(f'외부 참고구간 검증 {len(verification)}개 중 통계적 기준을 충족한 결과는 {passed}개입니다. 나머지는 두 번째 독립 집단 확보, 이상치 검토·보충 또는 구간·검사법 재평가가 필요할 수 있습니다. 아래 개별 판정을 확인하십시오.')
    return messages


def _figures(groups,directory):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from tametools.reporting import _configure_matplotlib_font
    _configure_matplotlib_font(plt, font_manager)
    from scipy import stats
    paths=[]
    for i,group in enumerate(groups,1):
        raw=np.asarray(group["raw_values"]); x=np.asarray(group["values"])
        fig,axes=plt.subplots(1,2,figsize=(10,3.5),layout="constrained")
        axes[0].hist(raw,bins="auto",color="#dbe3eb",label="Before outlier action")
        axes[0].hist(x,bins="auto",histtype="step",color="#176579",linewidth=1.5,label="Used in RI")
        axes[0].set(xlabel=f"Result ({group['unit']})",ylabel="Count")
        axes[0].legend(fontsize=7)
        stats.probplot(x,dist="norm",plot=axes[1])
        axes[1].set_title("Raw-scale normal Q-Q (diagnostic)")
        fig.suptitle(f"Group {i}: {group.get('display_name', group['test_name'])}  n={len(x)}",fontsize=10)
        path=directory/f"distribution_{i:02d}.png"
        fig.savefig(path,dpi=140)
        plt.close(fig)
        paths.append((group,path))
    return paths


def write_reports(output,cfg,snapshot,groups):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches,Pt
    from tametools.io import write_tame
    report_path=Path(cfg["REPORT_PATH"]).expanduser().resolve() if cfg["REPORT_PATH"] else None
    directory=Path(cfg["OUTPUT_DIR"]).expanduser().resolve() if cfg["OUTPUT_DIR"] else report_path.with_name(report_path.stem+"_files")
    if directory.exists():raise ValueError("Choose a new OUTPUT_DIR; existing reports are preserved")
    if report_path is not None and report_path.exists():raise ValueError("REPORT_PATH already exists")
    if report_path is not None and report_path.suffix.lower()!=".docx":raise ValueError("REPORT_PATH must end in .docx")
    directory.mkdir(parents=True)
    report_path=report_path or directory/"reference_interval_report_ko.docx"
    report_path.parent.mkdir(parents=True,exist_ok=True)
    csv_dir=directory/"tables"; csv_dir.mkdir()
    for name,frame in output.tables.items():frame.to_csv(csv_dir/(name+".csv"),index=False,encoding="utf-8-sig",float_format="%.17g")
    write_tame(directory/"reference_intervals.tame",output.dataset)
    (directory/"effective_input.json").write_text(json.dumps(_safe(snapshot),ensure_ascii=False,sort_keys=True,indent=2)+"\n",encoding="utf-8")
    figure_dir=directory/"figures"; figure_dir.mkdir()
    labels=output.dataset.meta['RI_EP28_RESULT'].get('TEST_LABELS',{})
    display_groups=[dict(g, display_name=labels.get(g['result_id'],g['test_name']) if g['test_name']==g['result_id'] else g['test_name']) for g in groups]
    figures=_figures(display_groups,figure_dir)
    main=output.tables["reference_intervals"]
    columns=["test_name","unit","group","method","n","ref_low","ref_high","status"]
    ci_columns=["test_name","group","method","ci_method","low_ci_low","low_ci_high","high_ci_low","high_ci_high"]
    sections=[
        ("참고구간 산출과 방법 비교",main,columns,"reference_intervals"),
        ("참고한계의 신뢰구간",main,ci_columns,"reference_intervals"),
        ("입력·제외 분모",output.tables["input_summary"],["result_id","reason","n"],"input_summary"),
        ("이상치 감사",output.tables["outliers"],["test_name","group","source_row","value","method","action"],"outliers"),
        ("이상치 방법별 민감도",output.tables["outlier_sensitivity"],["test_name","group","outlier_method","n","removed_n","ref_low","ref_high","status"],"outlier_sensitivity"),
        ("Partition 평가: Harris–Boyd",output.tables["partition_tests"],["test_name","factor","stratum","group_a","group_b","n_a","n_b","z","z_critical","sd_ratio","decision"],"partition_tests"),
        ("Partition 평가: 공통 구간 밖 꼬리",output.tables["partition_tails"],["test_name","factor","stratum","group","tail","outside_n","n","proportion","ci_low","ci_high","signal"],"partition_tails"),
        ("연구별 평가: Wilcoxon·SDR·BR",output.tables["partition_tests"],["test_name","factor","group_a","group_b","wilcoxon_p","sdr","br_low","br_high","reporting_unit"],"partition_tests"),
        ("연구별 평가: nested SDR",output.tables.get("nested_sdr",pd.DataFrame()),["test_name","outer","inner","n","sdr","nested_sdr","negative_component","status"],"nested_sdr"),
        ("LAVE 대상자 선별",output.tables.get("lave_membership",pd.DataFrame()),["result_id","stratum","subject_id","abnormal_other_tests","included","reason","converged"],"lave_membership"),
        ("LAVE 반복 과정",output.tables.get("lave_iterations",pd.DataFrame()),["stratum","iteration","reference_id","n","ref_low","ref_high","screen_low","screen_high"],"lave_iterations"),
        ("변환 후 일회 제외",output.tables.get("parametric_trim",pd.DataFrame()),["test_name","group","subject_id","value","removed"],"parametric_trim"),
        ("정규성: 적합 평균·분산에 대한 KS 보정",main,["test_name","group","method","ks_d","ks_lilliefors_p","ks_status"],"reference_intervals"),
        ("외부 참고구간의 20명 검증",output.tables["verification"],["test_name","group","n","batch","outside_n","decision","status","notes"],"verification"),
        ("구간별 해석 주의사항",main,["test_name","group","method","status","notes"],"reference_intervals"),
    ]
    document=Document()
    section=document.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Inches(11.69);section.page_height=Inches(8.27)
    section.left_margin=section.right_margin=Inches(.55)
    section.top_margin=section.bottom_margin=Inches(.55)
    style=document.styles["Normal"];style.font.name="Malgun Gothic";style.font.size=Pt(9)
    style._element.rPr.rFonts.set(qn("w:eastAsia"),"Malgun Gothic")
    document.add_heading("참고구간 설정·평가 보고서",0)
    intro=(f"중앙 {cfg['COVERAGE']*100:g}% 참고구간 · 참고한계 {cfg['CI_LEVEL']*100:g}% 신뢰구간. "
           f"주 방법: {cfg['METHOD']}; 이상치: {cfg['OUTLIER_METHOD']} / {cfg['OUTLIER_ACTION']}. "
           "CLSI EP28-A3c 9장·11.2절·부록 B를 근거로 계산하며, 검사실의 임상 채택 판단을 지원합니다.")
    document.add_paragraph(intro)
    population=cfg["POPULATION"]
    document.add_paragraph("참고집단 확인: "+("문서화됨" if population.get("REFERENCE_INDIVIDUALS_CONFIRMED",False) else "미확인 — 표본 기반 탐색적 후보 구간"))
    for key,value in population.items():document.add_paragraph(f"{key}: {value}")
    html=["<!doctype html><html lang='ko'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>",
          "<title>참고구간 설정·평가 보고서</title><style>body{font:15px/1.7 system-ui,sans-serif;max-width:1400px;margin:40px auto;padding:0 24px;color:#22313f}h1,h2{color:#155d6b}table{border-collapse:collapse;font-size:12px;width:100%}td,th{border:1px solid #cdd8dc;padding:7px;text-align:left;vertical-align:top}th{background:#e8f1f3}section{overflow:auto;margin:30px 0}img{max-width:100%}.note{background:#f1f5f6;padding:16px}pre{white-space:pre-wrap}a{color:#155d6b}@media print{body{margin:0}section{break-inside:auto}thead{display:table-header-group}}</style>",
          "<h1>참고구간 설정·평가 보고서</h1><p>"+escape(intro)+"</p>",
          "<div class='note'><b>참고집단 문서화</b><pre>"+escape(json.dumps(population,ensure_ascii=False,indent=2))+"</pre></div>"]
    document.add_heading('이 결과의 평가와 해석',1)
    html.append('<h2>이 결과의 평가와 해석</h2><ul>')
    for message in _interpretation(output.tables,cfg):
        document.add_paragraph(message)
        html.append('<li>'+escape(message)+'</li>')
    html.append('</ul>')
    def add_table(frame):
        if frame.empty:
            document.add_paragraph("해당 결과 없음 / 설정하지 않음.");return
        table=document.add_table(rows=1,cols=len(frame.columns));table.style="Table Grid"
        for cell,name in zip(table.rows[0].cells,frame.columns):cell.text=name
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
        repeat=OxmlElement("w:tblHeader");table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in frame.itertuples(index=False,name=None):
            for cell,value in zip(table.add_row().cells,row):cell.text=str(value)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:run.font.size=Pt(8)
    for title,frame,selected,name in sections:
        document.add_heading(title,1)
        display=_display(frame,selected,cfg["REPORT_MAX_ROWS"], output.dataset.meta["RI_EP28_RESULT"].get("TEST_LABELS",{}))
        add_table(display)
        total=f"전체 {len(frame)}행. 모든 열·행은 tables/{name}.csv에 보존합니다."
        if len(frame)>len(display):total+=" 문서에는 처음 "+str(len(display))+"행만 표시합니다."
        document.add_paragraph(total)
        html.append("<section><h2>"+escape(title)+"</h2>"+(display.to_html(index=False,escape=True) if len(display) else "<p>해당 결과 없음 / 설정하지 않음.</p>")+"<p>"+escape(total)+"</p></section>")
    document.add_heading("분포와 Q-Q 진단",1)
    html.append("<h2>분포와 Q-Q 진단</h2>")
    for group,path in figures:
        caption=f"{group['display_name']} ({str(group['unit']).replace('umol/L','µmol/L')}) / {_group_label(group['group'])}"
        document.add_paragraph(caption);document.add_picture(str(path),width=Inches(9.8))
        html.append("<p>"+escape(caption)+"</p><img alt='분포 및 정규 Q-Q 진단' src='data:image/png;base64,"+base64.b64encode(path.read_bytes()).decode()+"'>")
    document.add_heading("해석상 주의사항",1);html.append("<h2>해석상 주의사항</h2><ol>")
    for text in CAUTIONS:
        document.add_paragraph(text,style="List Number");html.append("<li>"+escape(text)+"</li>")
    html.append("</ol>")
    document.add_heading("실행 경고",1)
    for warning in output.warnings:document.add_paragraph(warning)
    html.append("<h2>실행 경고</h2><ul>"+"".join("<li>"+escape(w)+"</li>" for w in output.warnings)+"</ul>")
    document.add_heading("설정 및 재현 정보",1)
    settings=pd.DataFrame([dict(option=k,value=json.dumps(v,ensure_ascii=False)) for k,v in cfg.items()])
    add_table(settings)
    html.append("<details><summary>적용된 전체 설정</summary><pre>"+escape(json.dumps(_safe(cfg),ensure_ascii=False,indent=2))+"</pre></details>")
    document.add_heading("근거",1);html.append("<h2>근거</h2><ul>")
    for title,url in SOURCES:
        document.add_paragraph(title+" — "+url);html.append("<li><a href='"+escape(url,quote=True)+"'>"+escape(title)+"</a></li>")
    html.append("</ul></html>")
    document.save(report_path)
    (directory/"reference_interval_report_ko.html").write_text("\n".join(html),encoding="utf-8")
    versions={}
    for name in ["tametools","numpy","pandas","scipy","python-docx","matplotlib"]:
        try:versions[name]=metadata.version(name)
        except metadata.PackageNotFoundError:versions[name]="unknown"
    import tametools.ri_ep28 as engine
    code_files=[Path(__file__),Path(engine.__file__),Path(__file__).parent/"plugins/reference_interval_ep28.py",Path(__file__).parent/"analysis.py",Path(__file__).parent/"sex_normalization.py",Path(__file__).parent/"reporting.py"]
    from tametools import __version__
    versions['tametools_distribution'] = versions.get('tametools', 'unknown')
    versions['tametools'] = __version__
    manifest=dict(created_utc=datetime.now(timezone.utc).isoformat(),versions=versions,
                  implementation_sha256={p.name:sha256(p.read_bytes()).hexdigest() for p in code_files},
                  source_data_sha256=snapshot["source_data_sha256"],
                  files=[dict(path=p.relative_to(directory).as_posix(),sha256=sha256(p.read_bytes()).hexdigest())
                         for p in sorted(directory.rglob("*")) if p.is_file()])
    (directory/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return [str(p) for p in sorted(directory.rglob("*")) if p.is_file()]+([] if report_path.is_relative_to(directory) else [str(report_path)])
