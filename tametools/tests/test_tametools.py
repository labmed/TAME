from __future__ import annotations

import contextlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import textwrap
import unicodedata
import unittest
from unittest.mock import patch

import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from tametools.age import age_band_label, age_value_profile, parse_age, parse_age_to_years, standardize_age_dataset
from tametools.analysis import describe_dataset, exploratory_data_analysis, reference_interval_plan, validate_dataset
from tametools.actions import ActionError, available_action_pipelines, available_actions, execute_action, execute_action_pipeline
from tametools.cellstate import NULL, cell_state
from tametools.cli import main as cli_main
from tametools.command_pipelines import CommandPipelineError, available_command_pipelines, execute_command_pipeline
from tametools.evaluation import evaluate_dataset, write_evaluation_report
from tametools.exporters import export_dataset
from tametools.images import embed_image_columns, extract_image_columns
from tametools.io import import_xlsx_into_tame, read_tame, read_xlsx, write_data_tame, write_meta_tame, write_tame, write_xlsx
from tametools.merge import merge_datasets
from tametools.models import make_unique_column_names
from tametools.pipeline import execute_work
from tametools.plugin_base import list_plugins, run_plugin
from tametools.presets import apply_clinical_chemistry_preset, apply_preset
from tametools.provenance import append_log_entry
from tametools.review_profiles import category_value_profile, datetime_value_profile
from tametools.sex import normalize_sex, sex_value_profile, standardize_sex_dataset
from tametools.tag_placement import tags_to_header, tags_to_meta
from tametools.tag_catalog import flat_tag_names
from tametools.tags import build_header, parse_header
from tametools.transforms import anonymize_dataset, fix_num_comparator_values, harmonize_comparator_thresholds, sample_dataset, split_comparator_columns, write_mapping_tables


SAMPLE_TAME = textwrap.dedent(
    """
    <META>
    [INFO]
    DESCRIPTION = "sample"

    [TAGS]
    "진료과" = ["BY"]

    [SETTINGS]
    VALIDATE_ERROR = "DELETE"
    CRR = "VALUE"

    [WORKS]
    DEFAULT = ["VALIDATE", "DESCRIBE"]
    RI = ["VALIDATE", "REFERENCE_INTERVAL"]
    </META>
    <DATA>
    [[ID::STR]]등록번호\t[[SEX]]성별\t[[AGE]]나이\t[[ITEM]]검사항목명\t[[RESULT::<NUM>]]보고값\t진료과
    A0001\tF\t32\tAST\t25\tIM
    A0002\t남\t2m\tAST\t<3\tGS
    A0003\tbadsex\tbad\tALT\tpositive\tIM
    </DATA>
    """
).strip()

SAMPLE_ANON_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    </META>
    <DATA>
    [[HOSPITAL_ID]]병원ID\t[[NAME]]이름\t[[ID::STR]]등록번호\t[[RESULT::<NUM>]]보고값
    H01\t홍길동\tP001\t<30
    H01\t김영희\tP002\t25
    </DATA>
    """
).strip()

SAMPLE_NUM_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    </META>
    <DATA>
    [[HOSPITAL_ID]]병원ID\t[[ITEM]]검사항목명\t[[RESULT::NUM]]보고값
    A\tAST\t30
    A\tALT\t40
    </DATA>
    """
).strip()

SAMPLE_CNUM_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    CRR = "VALUE"
    </META>
    <DATA>
    [[HOSPITAL_ID]]병원ID\t[[ITEM]]검사항목명\t[[RESULT::<NUM>]]보고값
    B\tAST\t<30
    B\tAST\t<25
    B\tALT\t35
    </DATA>
    """
).strip()

SAMPLE_HARMONIZE_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    CRR = "HARMONIZE"
    </META>
    <DATA>
    [[ITEM]]검사항목명\t[[RESULT::<NUM>]]보고값
    AST\t<30
    AST\t<20
    AST\t22
    AST\t18
    ALT\t35
    </DATA>
    """
).strip()

SAMPLE_RI_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    CRR = "VALUE"

    [WORKS]
    RI = ["REFERENCE_INTERVAL"]
    </META>
    <DATA>
    [[TESTNAME]]검사항목명\t[[SEX]]성별\t[[AGE]]나이\t[[RESULT::<NUM>]]보고값
    AST\tM\t5\t10
    AST\tF\t15\t12
    AST\tM\t25\t<30
    AST\tF\t72\t20
    ALT\tM\t35\t30
    ALT\tF\t82\t40
    </DATA>
    """
).strip()

SAMPLE_RI_MISSING_TAGS_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    CRR = "VALUE"
    </META>
    <DATA>
    [[TESTNAME]]검사항목명\t[[RESULT::NUM]]보고값
    AST\t10
    AST\t12
    ALT\t20
    </DATA>
    """
).strip()

SAMPLE_VALUE_STATES_TAME = textwrap.dedent(
    """
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    </META>
    <DATA>
    [[COMMENT::STR::NULLABLE::EMPTY_OK::WS_OK]]비고\t[[CODE::STR::REQUIRED]]코드\t[[NOTE::STR]]메모\t[[RESULT::NUM::NULLABLE]]수치
    <<ABSENT>>\tA01\t정상\t1
    <<NULL>>\tA02\t<<NULL>>\t<<NULL>>
    <<EMPTY>>\tA03\t<<EMPTY>>\t2
    <<WS:3>>\tA04\t<<WS:2>>\t3
    \\<<NULL>>\tA05\t\\<<EMPTY>>\t4
    </DATA>
    """
).strip()

SAMPLE_IMAGE_DATA_URL = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO7Z6gAAAABJRU5ErkJggg=="

SAMPLE_IMAGE_TAME = textwrap.dedent(
    f"""
    <META>
    [SETTINGS]
    VALIDATE_ERROR = "REPORT"
    </META>
    <DATA>
    [[ID::STR]]등록번호\t[[IMAGE::B64]]썸네일\t[[LABEL::STR]]라벨
    P001\t{SAMPLE_IMAGE_DATA_URL}\t정상
    </DATA>
    """
).strip()

SAMPLE_FUNCTION_TAME = textwrap.dedent(
    """
    <META>
    [WORKS]
    DEFAULT = ["TRIM_RESULT"]

    [TRIM_RESULT]
    FUNCTION = "trim_result"

    [FUNCTIONS.trim_result]
    LANG = "python"
    ENTRY = "run"
    SOURCE = '''
def run(dataset, options, api):
    result_column = dataset.first_column_with_tag("RESULT")
    frame = dataset.df.copy()
    frame[result_column.name] = frame[result_column.name].map(
        lambda value: value.strip() if isinstance(value, str) else value
    )
    return dataset.replace(df=frame)
'''
    </META>
    <DATA>
    [[ID::STR]]등록번호\t[[RESULT::STR]]보고값
    P001\t  alpha  
    P002\tbeta
    </DATA>
    """
).strip()

SAMPLE_ACTION_TAME = textwrap.dedent(
    """
    <META>
    [ACTIONS.ADD_SEX_AGE_TAGS]
    LABEL = "Add SEX/AGE tags"
    TYPE = "TAG_COLUMNS"
    DESCRIPTION = "Add SEX tag to sex column and AGE tag to age column."

    [[ACTIONS.ADD_SEX_AGE_TAGS.COLUMNS]]
    COLUMN = "sex"
    TAGS = ["SEX"]

    [[ACTIONS.ADD_SEX_AGE_TAGS.COLUMNS]]
    COLUMN = "age"
    TAGS = ["AGE"]
    </META>
    <DATA>
    sex\tage\tvalue
    M\t30\t1
    F\t2mo\t2
    </DATA>
    """
).strip()

SAMPLE_R_FUNCTION_TAME = textwrap.dedent(
    """
    <META>
    [WORKS]
    DEFAULT = ["TRIM_RESULT_R"]

    [TRIM_RESULT_R]
    FUNCTION = "trim_result_r"

    [FUNCTIONS.trim_result_r]
    LANG = "r"
    ENTRY = "run"
    SOURCE = '''
run <- function(input_path, output_path) {
  df <- read.delim(input_path, check.names = FALSE, quote = "", comment.char = "", stringsAsFactors = FALSE)
  result_column <- grep("보고값$", names(df), value = TRUE)[1]
  df[[result_column]] <- trimws(df[[result_column]])
  write.table(df, output_path, sep = "\\t", row.names = FALSE, quote = FALSE, fileEncoding = "UTF-8")
  "trimmed by R"
}
'''
    </META>
    <DATA>
    [[ID::STR]]등록번호\t[[RESULT::STR]]보고값
    P001\t  gamma  
    P002\tdelta
    </DATA>
    """
).strip()


def _display_width_until(line: str, needle: str) -> int:
    prefix = line[: line.index(needle)]
    width = 0
    for char in prefix:
        if unicodedata.combining(char):
            continue
        width += 2 if unicodedata.east_asian_width(char) in {"F", "W"} else 1
    return width


class TagParsingTest(unittest.TestCase):
    def test_parse_and_build_header(self) -> None:
        name, tags = parse_header("[[RESULT::<NUM>]]보고값")
        self.assertEqual(name, "보고값")
        self.assertEqual(tags, ("RESULT", "<NUM>"))
        self.assertEqual(build_header(name, tags), "[[RESULT::<NUM>]]보고값")

    def test_parse_header_with_date_format_tag(self) -> None:
        # Regression (BUG-1): parameterized DATE/DATETIME format tags must parse in headers.
        name, tags = parse_header("[[RECEIVED_AT::DATETIME(%Y-%m-%d %H:%M:%S)]]접수일시")
        self.assertEqual(name, "접수일시")
        self.assertEqual(tags, ("RECEIVED_AT", "DATETIME(%Y-%m-%d %H:%M:%S)"))
        name2, tags2 = parse_header("[[DATE(%Y-%m-%d)]]생년월일")
        self.assertEqual(name2, "생년월일")
        self.assertEqual(tags2, ("DATE(%Y-%m-%d)",))

    def test_parse_header_with_qualified_tag(self) -> None:
        name, tags = parse_header("[[ID(PATIENT)::STR]]등록번호")
        self.assertEqual(name, "등록번호")
        self.assertEqual(tags, ("ID(patient)", "STR"))
        self.assertEqual(build_header(name, tags), "[[ID(patient)::STR]]등록번호")

    def test_parse_legacy_header_and_build_canonical_header(self) -> None:
        header = "_AGE_age\\_years"
        name, tags = parse_header(header)
        self.assertEqual(name, "age_years")
        self.assertEqual(tags, ("AGE",))
        self.assertEqual(build_header(name, tags), "[[AGE]]age_years")

    def test_parse_canonical_header_with_underscore_tag(self) -> None:
        header = "[[HOSPITAL_ID]]병원ID"
        name, tags = parse_header(header)
        self.assertEqual(name, "병원ID")
        self.assertEqual(tags, ("HOSPITAL_ID",))

    def test_parse_tag_only_header_uses_tag_block_as_column_name(self) -> None:
        name, tags = parse_header("[[SEX]]")
        self.assertEqual(name, "[[SEX]]")
        self.assertEqual(tags, ("SEX",))
        self.assertEqual(build_header(name, tags), "[[SEX]]")

    def test_make_unique_column_names_uses_pandas_style_suffixes(self) -> None:
        self.assertEqual(
            make_unique_column_names(["SEX", "SEX", "SEX.1", "SEX"]),
            ["SEX", "SEX.2", "SEX.1", "SEX.3"],
        )


class TameReadTest(unittest.TestCase):
    def test_read_tame_merges_tags(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        self.assertEqual(dataset.columns[0].name, "등록번호")
        self.assertIn("ID", dataset.columns[0].tags)
        self.assertIn("BY", dataset.first_column_with_tag("BY").tags)
        self.assertEqual(dataset.first_column_with_tag("RESULT").name, "보고값")

    def test_qualified_id_tags_inherit_generic_and_legacy_aliases(self) -> None:
        text = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"
            </META>
            <DATA>
            [[ID(hospital)::STR]]기관ID\t[[ID(patient)::STR]]등록번호\t[[ID(sample)::STR]]검체번호
            H01\tP001\tS001
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "qualified_ids.tame"
            path.write_text(text, encoding="utf-8")
            dataset = read_tame(path)

        self.assertEqual([column.name for column in dataset.columns_with_tag("ID")], ["기관ID", "등록번호", "검체번호"])
        self.assertEqual(dataset.first_column_with_tag("ID(patient)").name, "등록번호")
        self.assertEqual(dataset.first_column_with_tag("PATIENT_ID").name, "등록번호")
        self.assertEqual(dataset.first_column_with_tag("ID(hospital)").name, "기관ID")
        self.assertEqual(dataset.first_column_with_tag("HOSPITAL_ID").name, "기관ID")
        self.assertEqual(dataset.first_column_with_tag("ID(sample)").name, "검체번호")
        self.assertEqual(dataset.first_column_with_tag("SAMPLE_ID").name, "검체번호")

    def test_read_tame_supports_tag_only_and_duplicate_headers(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            SEX\tSEX\t[[SEX]]\t[[SEX]]
            M\tF\tM\tF
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "duplicate_headers.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        self.assertEqual([column.name for column in dataset.columns], ["SEX", "SEX.1", "[[SEX]]", "[[SEX]].1"])
        self.assertEqual(dataset.columns[2].tags, ("SEX",))
        self.assertEqual(dataset.columns[3].tags, ("SEX",))

    def test_write_tame_canonicalizes_legacy_headers(self) -> None:
        legacy_tame = textwrap.dedent(
            """
            <META>
            </META>
            <DATA>
            _HOSPITAL_ID_병원ID\t_RESULT::<NUM>_보고값
            H01\t<30
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "legacy.tame"
            output_path = Path(tmpdir) / "canonical.tame"
            input_path.write_text(legacy_tame, encoding="utf-8")
            dataset = read_tame(input_path)
            write_tame(output_path, dataset)
            written = output_path.read_text(encoding="utf-8")

        self.assertIn("[[HOSPITAL_ID]]병원ID", written)
        self.assertIn("[[RESULT::<NUM>]]보고값", written)
        self.assertNotIn("_HOSPITAL_ID_병원ID", written)

    def test_read_tame_distinguishes_absent_null_empty_and_ws(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "states.tame"
            path.write_text(SAMPLE_VALUE_STATES_TAME, encoding="utf-8")
            dataset = read_tame(path)

        self.assertIsNone(dataset.df.loc[0, "비고"])
        self.assertIs(dataset.df.loc[1, "비고"], NULL)
        self.assertEqual(dataset.df.loc[2, "비고"], "")
        self.assertEqual(dataset.df.loc[3, "비고"], "   ")
        self.assertEqual(dataset.df.loc[4, "비고"], "<<NULL>>")
        self.assertEqual(dataset.df.loc[4, "메모"], "<<EMPTY>>")
        self.assertEqual(cell_state(dataset.df.loc[0, "비고"]), "ABSENT")
        self.assertEqual(cell_state(dataset.df.loc[1, "비고"]), "NULL")
        self.assertEqual(cell_state(dataset.df.loc[2, "비고"]), "EMPTY")
        self.assertEqual(cell_state(dataset.df.loc[3, "비고"]), "WS")

    def test_append_log_entry_updates_meta_and_raw_meta_section(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")
            dataset = read_tame(path)

        logged = append_log_entry(
            dataset,
            action="fix",
            message="Applied standardize-sex.",
            parent="sample.tame",
            output="sample.fixed.tame",
            parameters={"action": "standardize-sex"},
        )

        self.assertEqual(len(logged.meta["LOG"]), 1)
        self.assertEqual(logged.meta["LOG"][0]["OPERATION"], "FIX")
        self.assertEqual(logged.meta["LOG"][0]["PARAMS"]["action"], "standardize-sex")
        self.assertIn("[[LOG]]", logged.raw_sections["META"])
        self.assertNotIn("[[LOG.ENTRIES]]", logged.raw_sections["META"])

    def test_meta_tags_are_base_and_header_tags_are_additive(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [TAGS]
            value = ["RESULT"]
            </META>
            <DATA>
            [[NUM]]value
            1
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "tag_merge.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        self.assertEqual(dataset.columns[0].tags, ("RESULT", "NUM"))

    def test_column_metadata_tags_are_supported(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [COLUMN.value]
            TAGS = ["RESULT"]
            UNIT = "U/L"
            LOINC = "1742-6"
            </META>
            <DATA>
            [[NUM]]value
            1
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "column_meta.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        self.assertEqual(dataset.columns[0].tags, ("RESULT", "NUM"))
        self.assertEqual(dataset.column_metadata("value")["UNIT"], "U/L")

    def test_format_section_controls_cell_tokens(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [FORMAT]
            NULL_TOKEN = "<NULL>"

            [SETTINGS]
            VALIDATE_ERROR = "REPORT"
            </META>
            <DATA>
            [[NULLABLE]]note
            <NULL>
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "format_tokens.tame"
            output_path = Path(tmpdir) / "format_tokens_out.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)
            write_tame(output_path, dataset)
            written = output_path.read_text(encoding="utf-8")

        self.assertIs(dataset.df.loc[0, "note"], NULL)
        self.assertIn("[FORMAT]", written)
        self.assertIn("<NULL>", written)

    def test_tags_can_be_materialized_to_meta_or_headers(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [TAGS]
            sex = ["CATEGORY"]
            age = ["NUM"]
            </META>
            <DATA>
            [[SEX]]sex\t[[AGE]]age
            M\t30
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "mixed_tags.tame"
            meta_path = Path(tmpdir) / "tags_meta.tame"
            header_path = Path(tmpdir) / "tags_header.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

            meta_output = tags_to_meta(dataset)
            header_output = tags_to_header(dataset)
            write_tame(meta_path, meta_output.dataset, tag_storage="meta")
            write_tame(header_path, header_output.dataset, tag_storage="header")
            meta_text = meta_path.read_text(encoding="utf-8")
            header_text = header_path.read_text(encoding="utf-8")
            meta_reloaded = read_tame(meta_path)
            header_reloaded = read_tame(header_path)

        self.assertIn("[COLUMN.sex]", meta_text)
        self.assertIn('TAGS = ["CATEGORY", "SEX"]', meta_text)
        self.assertIn("sex\tage", meta_text)
        self.assertNotIn("[[SEX]]sex", meta_text)
        self.assertNotIn("[TAGS]", header_text)
        self.assertIn("[[CATEGORY::SEX]]sex\t[[NUM::AGE]]age", header_text)
        self.assertEqual(meta_reloaded.columns[0].tags, ("CATEGORY", "SEX"))
        self.assertEqual(header_reloaded.columns[1].tags, ("NUM", "AGE"))


class MetaActionTest(unittest.TestCase):
    def test_available_actions_and_execute_tag_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "action.tame"
            input_path.write_text(SAMPLE_ACTION_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

        actions = available_actions(dataset.meta)
        output = execute_action(dataset, "ADD_SEX_AGE_TAGS")
        tagged = output.dataset

        self.assertEqual([action.name for action in actions], ["ADD_SEX_AGE_TAGS"])
        self.assertIsNotNone(tagged)
        self.assertEqual(tagged.columns[0].tags, ("SEX",))
        self.assertEqual(tagged.columns[1].tags, ("AGE",))
        self.assertEqual(output.message, "tagged_columns=2")
        self.assertEqual(output.table["column"].tolist(), ["sex", "age"])

    def test_row_filter_numeric_and_between_comparison(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.HIGH]
            TYPE = "ROW.INCLUDE"
            CONDITIONS = [ { COLUMN = "result", OP = "gt", VALUE = 100 } ]

            [ACTIONS.MID]
            TYPE = "ROW.INCLUDE"
            CONDITIONS = [ { COLUMN = "age", OP = "between", VALUES = [40, 60] } ]
            </META>
            <DATA>
            [[RESULT::NUM]]result\t[[AGE]]age
            50\t30
            150\t45
            200\t70
            10\t55
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "numfilter.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        high = execute_action(dataset, "HIGH").dataset
        self.assertEqual(high.df["result"].tolist(), ["150", "200"])
        mid = execute_action(dataset, "MID").dataset
        self.assertEqual(mid.df["age"].tolist(), ["45", "55"])

    def test_derive_action_computes_arithmetic_and_blocks_unsafe_expr(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.RATIO]
            TYPE = "DERIVE"
            NAME = "ratio"
            EXPR = "{ast} / {alt}"
            ROUND = 2
            TAGS = ["NUM"]

            [ACTIONS.BAD]
            TYPE = "DERIVE"
            NAME = "x"
            EXPR = "__import__('os')"
            </META>
            <DATA>
            [[RESULT::NUM]]ast\t[[RESULT::NUM]]alt
            40\t20
            30\t60
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "derive.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        derived = execute_action(dataset, "RATIO").dataset
        self.assertEqual(derived.columns[-1].name, "ratio")
        self.assertEqual(derived.columns[-1].tags, ("NUM",))
        self.assertEqual([round(float(v), 2) for v in derived.df["ratio"]], [2.0, 0.5])
        with self.assertRaises(ActionError):
            execute_action(dataset, "BAD")

    def test_dedup_action_removes_duplicate_rows(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.DD_ALL]
            TYPE = "DEDUP"
            KEEP = "first"

            [ACTIONS.DD_KEY]
            TYPE = "DEDUP"
            COLUMNS = ["pid", "test"]
            </META>
            <DATA>
            [[ID::STR]]pid\t[[ITEM]]test\t[[RESULT::NUM]]value
            P1\tAST\t10
            P1\tAST\t10
            P1\tAST\t99
            P2\tALT\t20
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "dedup.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        all_dd = execute_action(dataset, "DD_ALL")
        self.assertEqual(len(all_dd.dataset.df), 3)  # one exact dup removed
        self.assertEqual(int(all_dd.table["removed"].iloc[0]), 1)
        key_dd = execute_action(dataset, "DD_KEY")
        self.assertEqual(len(key_dd.dataset.df), 2)  # P1/AST collapses to one
        self.assertEqual(key_dd.dataset.df["value"].tolist(), ["10", "20"])

    def test_tier2_date_derive_impute_outlier_recode(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.AGE]
            TYPE = "DATE_DERIVE"
            KIND = "age"
            NAME = "나이"
            FROM = "birth"
            AS_OF = "visit"

            [ACTIONS.TAT]
            TYPE = "DATE_DERIVE"
            KIND = "diff_minutes"
            NAME = "tat"
            START = "recv"
            END = "done"

            [ACTIONS.FILL]
            TYPE = "IMPUTE"
            COLUMNS = ["result"]
            METHOD = "median"

            [ACTIONS.MAP_DEPT]
            TYPE = "RECODE"
            COLUMN = "dept"
            MAP = { IM = "내과", GS = "외과" }
            </META>
            <DATA>
            [[DATE]]birth\t[[DATE]]visit\t[[DATETIME]]recv\t[[DATETIME]]done\t[[RESULT::NUM]]result\t[[CATEGORY]]dept
            2000-01-01\t2020-01-01\t2024-01-01 08:00\t2024-01-01 08:30\t10\tIM
            1990-06-01\t2020-06-01\t2024-01-01 09:00\t2024-01-01 10:00\t12\tGS
            1980-01-01\t2020-01-01\t2024-01-01 09:00\t2024-01-01 09:15\t\tIM
            1970-01-01\t2020-01-01\t2024-01-01 09:00\t2024-01-01 09:15\t999\tGS
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "tier2.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        aged = execute_action(dataset, "AGE")
        self.assertEqual(aged.dataset.columns[-1].tags, ("AGE",))
        self.assertEqual([int(v) for v in aged.dataset.df["나이"]], [20, 30, 40, 50])

        tat = execute_action(dataset, "TAT")
        self.assertEqual([float(v) for v in tat.dataset.df["tat"]], [30.0, 60.0, 15.0, 15.0])

        filled = execute_action(dataset, "FILL")
        self.assertEqual(int(filled.table["filled"].iloc[0]), 1)
        self.assertFalse(any(str(v).strip() == "" for v in filled.dataset.df["result"]))

        recoded = execute_action(dataset, "MAP_DEPT")
        self.assertEqual(recoded.dataset.df["dept"].tolist(), ["내과", "외과", "내과", "외과"])

    def test_unit_convert_action_normalizes_result_units(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.CONVERT_UNITS]
            TYPE = "UNIT.CONVERT"
            SOURCE = "raw_result"
            UNIT_COLUMN = "unit"
            TARGET_UNIT_BY_TEST = { GLU = "mg/dL", CREA = "mg/dL" }
            OUTPUT_VALUE = "result"
            OUTPUT_UNIT = "unit_standard"
            OUTPUT_STATUS = "parse_status"
            DECIMALS = 2
            </META>
            <DATA>
            [[TESTNAME]]test\t[[RAW_RESULT::STR]]raw_result\t[[UNIT]]unit
            GLU\t5.6 mmol/L\tmmol/L
            GLU\t101 mg/dL\tmg/dL
            CREA\t88 umol/L\tumol/L
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "units.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        output = execute_action(dataset, "CONVERT_UNITS")
        converted = output.dataset

        self.assertEqual(converted.df["unit_standard"].tolist(), ["mg/dL", "mg/dL", "mg/dL"])
        self.assertEqual(converted.df["parse_status"].tolist(), ["CONVERTED", "SAME_UNIT", "CONVERTED"])
        self.assertAlmostEqual(float(converted.df.loc[0, "result"]), 100.90, places=2)
        self.assertAlmostEqual(float(converted.df.loc[2, "result"]), 1.00, places=2)
        self.assertEqual(validate_dataset(converted).issues, [])

    def test_outlier_filter_removes_iqr_outliers(self) -> None:
        values = "\n".join(str(v) for v in [10, 11, 12, 10, 11, 12, 13, 9, 10, 11, 500])
        sample = (
            "<META>\n[ACTIONS.CUT]\nTYPE = \"OUTLIER_FILTER\"\nCOLUMN = \"result\"\nMETHOD = \"iqr\"\n</META>\n"
            "<DATA>\n[[RESULT::NUM]]result\n" + values + "\n</DATA>\n"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "outlier.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)
        cut = execute_action(dataset, "CUT")
        self.assertEqual(int(cut.table["removed"].iloc[0]), 1)
        self.assertNotIn("500", cut.dataset.df["result"].tolist())

    def test_tier4_sort_round_bin_actions(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.SORTV]
            TYPE = "SORT"
            BY = ["v"]
            NUMERIC = true
            ASCENDING = false

            [ACTIONS.RND]
            TYPE = "ROUND"
            COLUMNS = ["v"]
            DECIMALS = 1

            [ACTIONS.AGEBIN]
            TYPE = "BIN"
            COLUMN = "age"
            NAME = "ageband"
            BINS = [0, 40, 60, 200]
            LABELS = ["<40", "40-59", "60+"]
            </META>
            <DATA>
            [[RESULT::NUM]]v\t[[AGE]]age
            3.456\t25
            10.1\t55
            7.0\t65
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "tier4.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        sorted_ds = execute_action(dataset, "SORTV").dataset
        self.assertEqual([float(v) for v in sorted_ds.df["v"]], [10.1, 7.0, 3.456])
        rounded = execute_action(dataset, "RND").dataset
        self.assertEqual(rounded.df["v"].tolist(), ["3.5", "10.1", "7.0"])
        binned = execute_action(dataset, "AGEBIN").dataset
        self.assertEqual(binned.df["ageband"].tolist(), ["<40", "40-59", "60+"])
        self.assertEqual(binned.columns[-1].tags, ("CATEGORY",))

    def test_tier2_join_action_merges_external_table(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.JOINDEMO]
            TYPE = "JOIN"
            RIGHT = "{right}"
            ON = ["pid"]
            HOW = "left"
            </META>
            <DATA>
            [[ID::STR]]pid\t[[RESULT::NUM]]result
            P1\t10
            P2\t20
            P3\t30
            </DATA>
            """
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            right_path = Path(tmpdir) / "demo.csv"
            pd.DataFrame({"pid": ["P1", "P2"], "plan": ["A", "B"]}).to_csv(right_path, index=False)
            input_path = Path(tmpdir) / "join.tame"
            input_path.write_text(textwrap.dedent(sample).strip().replace("{right}", right_path.as_posix()), encoding="utf-8")
            dataset = read_tame(input_path)
            joined = execute_action(dataset, "JOINDEMO")

        self.assertIn("plan", [c.name for c in joined.dataset.columns])
        self.assertEqual(joined.dataset.df["plan"].tolist()[:2], ["A", "B"])
        self.assertEqual(len(joined.dataset.df), 3)

    def test_join_action_select_accepts_right_tag_selector(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.JOINDEMO]
            TYPE = "JOIN"
            RIGHT = "{right}"
            ON = ["tag:ID(patient)"]
            SELECT = ["tag:INSURANCE"]
            HOW = "left"
            </META>
            <DATA>
            [[ID(patient)::STR]]pid\t[[RESULT::NUM]]result
            P1\t10
            P2\t20
            </DATA>
            """
        )
        right = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[ID(patient)::STR]]right_pid\t[[INSURANCE::CATEGORY]]plan\tignored
            P1\tA\tx
            P2\tB\ty
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            right_path = Path(tmpdir) / "demo.tame"
            right_path.write_text(right, encoding="utf-8")
            input_path = Path(tmpdir) / "join.tame"
            input_path.write_text(textwrap.dedent(sample).strip().replace("{right}", right_path.as_posix()), encoding="utf-8")
            dataset = read_tame(input_path)
            joined = execute_action(dataset, "JOINDEMO")

        self.assertIn("plan", [c.name for c in joined.dataset.columns])
        self.assertNotIn("ignored", [c.name for c in joined.dataset.columns])
        self.assertEqual(joined.dataset.df["plan"].tolist(), ["A", "B"])

    def test_parameterized_age_tags_behave_like_css_class_and_qualifier(self) -> None:
        sample = textwrap.dedent(
            """
            <DATA>
            [[AGE(baseline)::NUM]]age_a\t[[AGE(visit)::NUM]]age_b
            10\t11
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "age.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        self.assertEqual([c.name for c in dataset.columns_with_tag("AGE")], ["age_a", "age_b"])
        self.assertEqual([c.name for c in dataset.columns_with_tag("AGE(baseline)")], ["age_a"])
        self.assertEqual([c.name for c in dataset.columns_with_tag("AGE(visit)")], ["age_b"])

    def test_execute_preprocessing_action_pipeline(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.KEEP_GROUP_A]
            TYPE = "ROW.INCLUDE"
            COLUMN = "group"
            VALUES = ["A"]

            [ACTIONS.SPLIT_SEX_AGE]
            TYPE = "COLUMN.SPLIT"
            COLUMN = "sex_age"
            SEP = "/"
            INTO = ["sex", "age"]
            DROP_SOURCE = true

            [ACTIONS.SPLIT_SEX_AGE.COLUMN_TAGS]
            sex = ["SEX"]
            age = ["AGE"]

            [ACTIONS.CONCAT_SEX_AGE]
            TYPE = "COLUMN.CONCAT"
            COLUMNS = ["sex", "age"]
            OUTPUT = "sex_age_text"
            SEP = "/"
            TAGS = ["STR"]

            [ACTIONS.DROP_UNUSED]
            TYPE = "COLUMN.EXCLUDE"
            COLUMNS = ["dropme"]

            [ACTIONS.ADD_SOURCE]
            TYPE = "ADD_COLUMN"
            NAME = "source"
            VALUE = "web"
            TAGS = ["CATEGORY"]

            [ACTION_PIPELINES]
            DEFAULT = ["KEEP_GROUP_A", "SPLIT_SEX_AGE", "CONCAT_SEX_AGE", "DROP_UNUSED", "ADD_SOURCE"]
            </META>
            <DATA>
            group\tsex_age\tdropme\tvalue
            A\tM/10\tx\t1
            B\tF/20\ty\t2
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "preprocess.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        result = execute_action_pipeline(dataset)
        transformed = result.final_dataset

        self.assertEqual(available_action_pipelines(dataset.meta), ["DEFAULT"])
        self.assertEqual([output.name for output in result.outputs], ["KEEP_GROUP_A", "SPLIT_SEX_AGE", "CONCAT_SEX_AGE", "DROP_UNUSED", "ADD_SOURCE"])
        self.assertEqual(list(transformed.df.columns), ["group", "sex", "age", "value", "sex_age_text", "source"])
        self.assertEqual(len(transformed.df), 1)
        self.assertEqual(transformed.df.loc[0, "sex"], "M")
        self.assertEqual(transformed.df.loc[0, "age"], "10")
        self.assertEqual(transformed.df.loc[0, "sex_age_text"], "M/10")
        self.assertEqual(transformed.df.loc[0, "source"], "web")
        self.assertEqual(transformed.first_column_with_tag("SEX").name, "sex")
        self.assertEqual(transformed.first_column_with_tag("AGE").name, "age")
        self.assertEqual(transformed.first_column_with_tag("STR").name, "sex_age_text")
        self.assertEqual(transformed.first_column_with_tag("CATEGORY").name, "source")
        self.assertEqual(
            [entry["OPERATION"] for entry in transformed.meta["LOG"]],
            [
                "ACTION:KEEP_GROUP_A",
                "ACTION:SPLIT_SEX_AGE",
                "ACTION:CONCAT_SEX_AGE",
                "ACTION:DROP_UNUSED",
                "ACTION:ADD_SOURCE",
            ],
        )
        self.assertEqual(transformed.meta["LOG"][0]["PARAMS"]["type"], "ROW.INCLUDE")
        self.assertIn("input_hash", transformed.meta["LOG"][0]["PARAMS"])
        self.assertIn("output_hash", transformed.meta["LOG"][-1]["PARAMS"])

    def test_add_column_template_resolves_tag_references(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.ADD_LABEL]
            TYPE = "ADD_COLUMN"
            NAME = "label"
            TEMPLATE = "{tag:ITEM}:{tag:RESULT}"
            TAGS = ["CATEGORY"]
            </META>
            <DATA>
            [[ITEM]]test\t[[RESULT::NUM]]value
            AST\t10
            ALT\t20
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "template.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        result = execute_action(dataset, "ADD_LABEL")
        transformed = result.dataset

        self.assertEqual(transformed.df["label"].tolist(), ["AST:10", "ALT:20"])
        self.assertEqual(transformed.columns[-1].tags, ("CATEGORY",))

    def test_single_tag_selector_rejects_ambiguous_matches(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.SPLIT_ID]
            TYPE = "COLUMN.SPLIT"
            COLUMN = "tag:ID"
            SEP = "-"
            INTO = ["left", "right"]
            </META>
            <DATA>
            [[ID(patient)::STR]]patient_id\t[[ID(sample)::STR]]sample_id
            P-001\tS-001
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "ambiguous.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        with self.assertRaisesRegex(ActionError, "Multiple columns found with tag ID"):
            execute_action(dataset, "SPLIT_ID")

        try:
            execute_action(dataset, "SPLIT_ID")
        except ActionError as exc:
            self.assertIn("This action requires a single column", str(exc))
            self.assertIn("exact column name", str(exc))
            self.assertNotIn("TAGS selector", str(exc))
        else:
            self.fail("expected ambiguous tag selector to raise ActionError")

    def test_pivot_longer_and_pivot_wider_actions_roundtrip_shape(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.PIVOT_RESULTS_LONGER]
            TYPE = "PIVOT_LONGER"
            ID_COLS = ["id"]
            COLS = ["AST", "ALT"]
            NAMES_TO = "test"
            VALUES_TO = "result"
            VALUES_TAGS = ["RESULT", "NUM"]

            [ACTIONS.PIVOT_RESULTS_WIDER]
            TYPE = "PIVOT_WIDER"
            ID_COLS = ["id"]
            NAMES_FROM = "test"
            VALUES_FROM = "result"

            [ACTION_PIPELINES]
            DEFAULT = ["PIVOT_RESULTS_LONGER", "PIVOT_RESULTS_WIDER"]
            </META>
            <DATA>
            id\tAST\tALT
            P1\t10\t20
            P2\t30\t40
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "pivot.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        result = execute_action_pipeline(dataset)
        wide = result.final_dataset.df.sort_values("id").reset_index(drop=True)

        self.assertEqual(len(result.outputs), 2)
        self.assertEqual(set(wide.columns), {"id", "AST", "ALT"})
        self.assertEqual(wide.loc[0, "AST"], "10")
        self.assertEqual(wide.loc[0, "ALT"], "20")
        self.assertEqual(wide.loc[1, "AST"], "30")
        self.assertEqual(wide.loc[1, "ALT"], "40")

    def test_pivot_actions_preserve_source_metadata(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.WIDE]
            TYPE = "PIVOT_WIDER"
            ID_COLS = ["tag:ID(sample)"]
            NAMES_FROM = "tag:ITEM"
            VALUES_FROM = "tag:RESULT"

            [ACTIONS.LONG_REF]
            TYPE = "PIVOT_LONGER"
            ID_COLS = ["tag:ID(sample)", "tag:ITEM"]
            COLS = ["tag:REF_LOW", "tag:REF_HIGH"]
            NAMES_TO = "bound"
            VALUES_TO = "limit"
            VALUES_TAGS = ["NUM"]
            </META>
            <DATA>
            [[ID(sample)::STR]]sample\t[[ITEM]]test\t[[RESULT::<NUM>]]value\t[[UNIT]]unit\t[[REF_LOW::NUM]]low\t[[REF_HIGH::NUM]]high
            S1\tAST\t10\tU/L\t0\t40
            S2\tAST\t20\tU/L\t0\t40
            S3\tCr\t1.0\tmg/dL\t0.7\t1.3
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "pivot_meta.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        wide = execute_action(dataset, "WIDE").dataset
        ast_meta = wide.column_metadata("AST")
        cr_meta = wide.column_metadata("Cr")
        self.assertEqual(ast_meta["PIVOT"]["NAME_VALUE"], "AST")
        self.assertEqual(ast_meta["PIVOT_CONTEXT"]["UNIT"], "U/L")
        self.assertEqual(ast_meta["PIVOT_CONTEXT"]["REF_LOW"], "0")
        self.assertEqual(ast_meta["PIVOT_CONTEXT"]["REF_HIGH"], "40")
        self.assertEqual(cr_meta["PIVOT_CONTEXT"]["UNIT"], "mg/dL")

        longer = execute_action(dataset, "LONG_REF").dataset
        bound_meta = longer.column_metadata("bound")
        limit_meta = longer.column_metadata("limit")
        self.assertEqual(bound_meta["PIVOT"]["SOURCE_TAGS"]["low"], ["REF_LOW", "NUM"])
        self.assertEqual(bound_meta["PIVOT"]["SOURCE_TAGS"]["high"], ["REF_HIGH", "NUM"])
        self.assertEqual(limit_meta["PIVOT"]["TYPE"], "LONGER")


class AnalysisTest(unittest.TestCase):
    def test_parse_age_uses_hl7_ucum_units_and_default_years(self) -> None:
        self.assertEqual(parse_age("10").canonical_value, "10")
        self.assertEqual(parse_age("10a").unit_code, "a")
        self.assertEqual(parse_age("10a").canonical_value, "10")
        self.assertAlmostEqual(parse_age_to_years("2mo"), 2 / 12)
        self.assertAlmostEqual(parse_age_to_years("2m"), 2 / 12)
        self.assertAlmostEqual(parse_age_to_years("1d"), 1 / 365.25)
        self.assertEqual(parse_age("0").canonical_value, "0")
        self.assertEqual(age_band_label(parse_age_to_years("0"), width=10), "<1")
        self.assertEqual(age_band_label(parse_age_to_years("6mo"), width=5), "<1")
        self.assertEqual(age_band_label(parse_age_to_years("1"), width=5), "1-4")
        self.assertEqual(age_band_label(parse_age_to_years("9"), width=10), "1-9")
        self.assertEqual(age_band_label(parse_age_to_years("10"), width=10), "10-19")
        self.assertIsNone(parse_age("-1a"))
        self.assertIsNone(parse_age("2wk"))
        self.assertIsNone(parse_age("bad"))

    def test_age_value_profile_groups_convertible_and_unrecognized_values(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[AGE]]나이
            0
            10a
            2m
            2mo
            1d
            bad
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "age_profile.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        profile = age_value_profile(dataset)
        rows = profile.set_index("raw_value")

        self.assertEqual(rows.loc["0", "canonical_value"], "0")
        self.assertEqual(rows.loc["0", "band_10"], "<1")
        self.assertFalse(bool(rows.loc["0", "will_change"]))
        self.assertEqual(rows.loc["10a", "canonical_value"], "10")
        self.assertTrue(bool(rows.loc["10a", "will_change"]))
        self.assertEqual(rows.loc["2m", "canonical_value"], "2mo")
        self.assertEqual(rows.loc["2mo", "unit_code"], "mo")
        self.assertEqual(rows.loc["1d", "unit_code"], "d")
        self.assertFalse(bool(rows.loc["bad", "recognized"]))

    def test_standardize_age_dataset_writes_storage_age_format(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[AGE]]나이
            10a
            2m
            2개월
            1day
            0
            bad
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "age_standardize.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        standardized = standardize_age_dataset(dataset)

        self.assertEqual(standardized.df["나이"].tolist(), ["10", "2mo", "2mo", "1d", "0", "bad"])

    def test_category_and_datetime_profiles_summarize_tagged_columns(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[CATEGORY]]진단군\t[[DATETIME]]채혈시각
            A\t2024-01-01
            B\t2024-01-02 13:30
            A\tbad-date
            B\t20240103
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "review_profiles.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        category = category_value_profile(dataset).set_index("raw_value")
        datetimes = datetime_value_profile(dataset)
        validation = validate_dataset(dataset)

        self.assertEqual(int(category.loc["A", "raw_count"]), 2)
        self.assertEqual(int(category.loc["B", "raw_count"]), 2)
        self.assertEqual(int(datetimes.loc[datetimes["recognized"], "raw_count"].sum()), 3)
        self.assertEqual(int(datetimes.loc[~datetimes["recognized"], "raw_count"].sum()), 1)
        self.assertEqual(datetimes.loc[~datetimes["recognized"], "raw_value"].tolist(), ["bad-date"])
        self.assertEqual({issue.tag for issue in validation.issues}, {"DATETIME"})

    def test_date_format_tag_is_validated_and_profiled(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[DOB::DATE(%y%m%d)]]birthdate
            230101
            2023-01-01
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "date_format.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        profile = datetime_value_profile(dataset).set_index("raw_value")
        validation = validate_dataset(dataset)

        self.assertTrue(bool(profile.loc["230101", "recognized"]))
        self.assertFalse(bool(profile.loc["2023-01-01", "recognized"]))
        self.assertEqual({issue.tag for issue in validation.issues}, {"DATE"})

    def test_temporal_limits_can_be_configured_in_column_meta(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"

            [COLUMN."채혈시각"]
            TAGS = ["COLLECTION_AT", "DATETIME"]
            MIN = "2026-01-01 00:00:00"
            MAX = "2026-01-31 23:59:59"

            [COLUMN."보고일"]
            TAGS = ["DATE"]
            LOWER_LIMIT = "2026-01-01"
            UPPER_LIMIT = "2026-01-31"

            [COLUMN."보고시간"]
            TAGS = ["TIME"]
            MIN = "08:00"
            MAX = "18:00"
            </META>
            <DATA>
            채혈시각\t보고일\t보고시간
            2026-01-15 08:30:00\t2026-01-15\t12:00
            2025-12-31 23:59:59\t2025-12-31\t07:59
            2026-02-01 00:00:00\t2026-02-01\t18:01
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "temporal_limits.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        issues = validate_dataset(dataset).issues

        self.assertEqual(
            [(issue.row_number, issue.column, issue.tag) for issue in issues],
            [
                (3, "채혈시각", "MIN"),
                (4, "채혈시각", "MAX"),
                (3, "보고일", "MIN"),
                (4, "보고일", "MAX"),
                (3, "보고시간", "MIN"),
                (4, "보고시간", "MAX"),
            ],
        )

    def test_invalid_column_meta_temporal_limit_is_reported(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [COLUMN.when]
            TAGS = ["DATETIME"]
            MIN = "not-a-date"
            </META>
            <DATA>
            when
            2026-01-15 08:30:00
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "bad_temporal_limit.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        issues = validate_dataset(dataset).issues

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].row_number, 1)
        self.assertEqual(issues[0].column, "when")
        self.assertEqual(issues[0].tag, "LIMIT")
        self.assertIn("MIN", issues[0].message)

    def test_normalize_sex_recognizes_canonical_text_and_binary_map(self) -> None:
        self.assertEqual(normalize_sex("M"), "male")
        self.assertEqual(normalize_sex("female"), "female")
        self.assertEqual(normalize_sex("X"), "other")
        self.assertEqual(normalize_sex("미상"), "unknown")
        self.assertIsNone(normalize_sex("1"))
        self.assertEqual(normalize_sex("1", binary_map={"1": "male", "0": "female"}), "male")
        self.assertEqual(normalize_sex(0, binary_map={"1": "male", "0": "female"}), "female")

    def test_write_tame_standardizes_sex_values(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별
            M
            F
            남
            여
            X
            unknown
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sex.tame"
            output_path = Path(tmpdir) / "sex_out.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)
            write_tame(output_path, dataset)
            reloaded = read_tame(output_path)

        self.assertEqual(
            reloaded.df["성별"].tolist(),
            ["male", "female", "male", "female", "other", "unknown"],
        )

    def test_sex_value_profile_keeps_raw_values_grouped_by_canonical_value(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별
            M
            M
            male
            F
            badsex
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sex_profile.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        profile = sex_value_profile(dataset)
        male_rows = profile.loc[profile["canonical_value"] == "male"].set_index("raw_value")

        self.assertEqual(int(male_rows.loc["M", "raw_count"]), 2)
        self.assertEqual(int(male_rows.loc["male", "raw_count"]), 1)
        self.assertEqual(int(male_rows.loc["M", "canonical_count"]), 3)
        self.assertTrue(bool(male_rows.loc["M", "will_change"]))
        self.assertFalse(bool(male_rows.loc["male", "will_change"]))
        self.assertEqual(profile.loc[~profile["recognized"], "raw_value"].tolist(), ["badsex"])

    def test_standardize_sex_dataset_uses_confirmed_binary_map(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별
            1
            0
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "binary_sex.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        self.assertEqual({issue.tag for issue in validate_dataset(dataset).issues}, {"SEX"})
        standardized = standardize_sex_dataset(dataset, binary_map={"1": "female", "0": "male"})
        self.assertEqual(standardized.df["성별"].tolist(), ["female", "male"])
        self.assertEqual(validate_dataset(standardized).issues, [])

    def test_validate_and_describe(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        validation = validate_dataset(dataset)
        self.assertEqual(len(validation.issues), 3)
        self.assertEqual(len(validation.cleaned_dataset.df), 2)

        described = describe_dataset(validation.cleaned_dataset)
        result_row = described.loc[described["column"] == "보고값"].iloc[0]
        age_row = described.loc[described["column"] == "나이"].iloc[0]
        sex_row = described.loc[described["column"] == "성별"].iloc[0]
        self.assertEqual(result_row["kind"], "numeric")
        self.assertEqual(int(result_row["count"]), 2)
        self.assertEqual(age_row["kind"], "numeric")
        self.assertEqual(sex_row["kind"], "categorical")
        for field in ("top", "mean", "median", "q2_5", "q97_5"):
            self.assertTrue(pd.isna(sex_row[field]), f"{field} should be NaN for categorical rows")

    def test_reference_interval_plan_finds_tagged_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        plan = reference_interval_plan(dataset)
        self.assertEqual(plan.item_column.name, "검사항목명")
        self.assertEqual(plan.sex_column.name, "성별")
        self.assertEqual(plan.age_column.name, "나이")
        self.assertEqual(plan.by_columns[0].name, "진료과")

    def test_execute_work_uses_meta_pipeline(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        result = execute_work(dataset, "DEFAULT")
        self.assertEqual(result.work_name, "DEFAULT")
        self.assertEqual(len(result.outputs), 2)
        self.assertEqual(result.outputs[0].name, "VALIDATE")
        self.assertEqual(result.outputs[1].name, "DESCRIBE")
        self.assertEqual(len(result.final_dataset.df), 2)
        self.assertEqual([entry["OPERATION"] for entry in result.final_dataset.meta["LOG"]], ["WORK:VALIDATE", "WORK:DESCRIBE"])
        self.assertEqual(result.final_dataset.meta["LOG"][0]["PARAMS"]["step"], "VALIDATE")

    def test_eda_reports_comparator_profile_and_warning(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "cnum.tame"
            path.write_text(SAMPLE_CNUM_TAME, encoding="utf-8")
            dataset = read_tame(path)

        report = exploratory_data_analysis(dataset, comparator_policy="VALUE")
        self.assertFalse(report.summary.empty)
        self.assertFalse(report.comparator_profile.empty)
        self.assertTrue(any(raw in set(report.comparator_profile["raw_value"]) for raw in ["<30", "<25"]))
        self.assertTrue(any("Multiple comparator thresholds detected" in warning for warning in report.warnings))
        impact = report.comparator_policy_impact.iloc[0]
        self.assertEqual(impact["policy"], "VALUE")
        self.assertGreaterEqual(int(impact["numeric_used_under_policy"]), 3)
        self.assertFalse(report.harmonization_preview.empty)

    def test_eda_reports_category_percentiles_and_result_by_groups(self) -> None:
        sample = textwrap.dedent(
            """
            <DATA>
            [[TESTNAME::CATEGORY]]검사\t[[SEX::CATEGORY::BY]]성별\t[[INSTRUMENT::CATEGORY::BY]]장비\t[[RESULT::NUM]]결과
            AST\tmale\tA\t10
            AST\tfemale\tA\t20
            ALT\tmale\tB\t30
            ALT\tfemale\tB\t40
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "eda_groups.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        report = exploratory_data_analysis(dataset)
        ast = report.category_distribution.loc[
            (report.category_distribution["column"] == "검사")
            & (report.category_distribution["category"] == "AST")
        ].iloc[0]
        self.assertEqual(int(ast["count"]), 2)
        self.assertEqual(float(ast["percent_of_column"]), 50.0)
        self.assertIn("P97_5", set(report.numeric_percentiles["stat"]))
        self.assertFalse(report.numeric_percentile_bands.empty)
        by_columns = set(report.result_by_summary["by_column"])
        self.assertEqual(by_columns, {"성별", "장비"})
        equipment_a = report.result_by_summary.loc[
            (report.result_by_summary["by_column"] == "장비")
            & (report.result_by_summary["by_value"] == "A")
        ].iloc[0]
        self.assertEqual(int(equipment_a["count"]), 2)
        self.assertEqual(float(equipment_a["median"]), 15.0)
        equipment_all_ast = report.result_by_summary.loc[
            (report.result_by_summary["by_column"] == "장비")
            & (report.result_by_summary["by_value"] == "all")
            & (report.result_by_summary["item"] == "AST")
        ].iloc[0]
        self.assertEqual(int(equipment_all_ast["count"]), 2)
        self.assertEqual(float(equipment_all_ast["median"]), 15.0)

    def test_custom_age5_tag_inherits_age_and_controls_grouping(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [TAG_DEFINITIONS.AGE5]
            LABEL = "Age with 5-year groups"
            DESCRIPTION = "AGE-compatible custom tag using 5-year grouping."
            INHERITS = ["AGE"]
            AGE_BIN_WIDTH = 5
            </META>
            <DATA>
            [[TESTNAME::CATEGORY]]검사항목명\t[[SEX::CATEGORY::BY]]성별\t[[AGE5]]나이\t[[RESULT::NUM]]보고값
            AST\tmale\t1\t10
            AST\tfemale\t4\t20
            AST\tmale\t6\t30
            AST\tfemale\t9\t40
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "age5.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        self.assertIn("AGE5", flat_tag_names(dataset.meta))
        self.assertEqual(dataset.first_column_with_tag("AGE").name, "나이")
        self.assertEqual(validate_dataset(dataset).issues, [])

        age_profile = age_value_profile(dataset)
        self.assertEqual(set(age_profile["preferred_band_width"]), {5})
        self.assertEqual(set(age_profile["preferred_band"]), {"1-4", "5-9"})

        report = exploratory_data_analysis(dataset)
        age_summary = report.result_by_summary.loc[report.result_by_summary["by_column"] == "나이"]
        self.assertEqual(set(age_summary["by_grouping"]), {"age_band_5"})
        self.assertEqual(set(age_summary["by_value"]), {"all", "1-4", "5-9"})
        age_all = age_summary.loc[age_summary["by_value"] == "all"].iloc[0]
        self.assertEqual(int(age_all["count"]), 4)
        self.assertEqual(float(age_all["median"]), 25.0)

        chemistry = run_plugin(dataset, "CHEMISTRY_ANALYSIS", dataset.meta, "AGE_SEX_RESULT", {"MODE": "AGE_SEX_RESULT"})
        self.assertEqual(set(chemistry.dataset.df["연령그룹"]), {"1-4", "5-9"})

        reference = run_plugin(dataset, "REFERENCE_INTERVAL", dataset.meta, "REFERENCE_INTERVAL", {})
        self.assertEqual(reference.dataset.meta["PLUGIN"]["AGE_BIN_WIDTH"], 5)
        self.assertIn("1-4", set(reference.dataset.df["age_group"]))

    def test_eda_harmonize_policy_updates_impact(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "harmonize.tame"
            path.write_text(SAMPLE_HARMONIZE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        report = exploratory_data_analysis(dataset, comparator_policy="HARMONIZE")
        impact = report.comparator_policy_impact.loc[report.comparator_policy_impact["result_column"] == "보고값"].iloc[0]
        self.assertEqual(impact["policy"], "HARMONIZE")
        self.assertGreaterEqual(int(impact["harmonized_rows"]), 3)

    def test_validate_recognizes_cell_state_policy_tags(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "states.tame"
            path.write_text(SAMPLE_VALUE_STATES_TAME, encoding="utf-8")
            dataset = read_tame(path)

        validation = validate_dataset(dataset)
        issues = {(issue.row_number, issue.column, issue.tag) for issue in validation.issues}
        self.assertIn((3, "메모", "NULL"), issues)
        self.assertIn((4, "메모", "EMPTY"), issues)
        self.assertIn((5, "메모", "WS"), issues)
        self.assertNotIn((3, "비고", "NULL"), issues)
        self.assertNotIn((4, "비고", "EMPTY"), issues)
        self.assertNotIn((5, "비고", "WS"), issues)

        described = describe_dataset(dataset)
        row = described.loc[described["column"] == "비고"].iloc[0]
        self.assertEqual(int(row["absent_cells"]), 1)
        self.assertEqual(int(row["null_cells"]), 1)
        self.assertEqual(int(row["empty_cells"]), 1)
        self.assertEqual(int(row["ws_cells"]), 1)


class CliTest(unittest.TestCase):
    def _run_cli(self, *args: str) -> tuple[int, str]:
        output = io.StringIO()
        error = io.StringIO()
        with (
            patch.object(sys, "argv", ["tametools", *args]),
            contextlib.redirect_stdout(output),
            contextlib.redirect_stderr(error),
        ):
            code = cli_main()
        return code, output.getvalue() + error.getvalue()

    def _run_cli_with_stdin(self, stdin_text: str, *args: str) -> tuple[int, str]:
        output = io.StringIO()
        error = io.StringIO()
        with (
            patch.object(sys, "argv", ["tametools", *args]),
            patch.object(sys, "stdin", io.StringIO(stdin_text)),
            contextlib.redirect_stdout(output),
            contextlib.redirect_stderr(error),
        ):
            code = cli_main()
        return code, output.getvalue() + error.getvalue()

    def test_cli_help_is_beginner_oriented(self) -> None:
        code, output = self._run_cli("--help")

        self.assertEqual(code, 0)
        self.assertIn("Recommended beginner workflow:", output)
        self.assertIn("tametools check FILE", output)
        self.assertIn("tametools analyze fixed.tame PLUGIN", output)
        self.assertIn("Command groups:", output)
        self.assertIn("Compatibility notes:", output)
        self.assertIn("Run a declared analysis plan or a named plugin.", output)

    def test_cli_help_command_shows_command_specific_examples(self) -> None:
        code, output = self._run_cli("help", "analyze")

        self.assertEqual(code, 0)
        self.assertIn("usage: tametools analyze", output)
        self.assertIn("Use 'tametools plugins FILE' to discover plugin names.", output)
        self.assertIn("tametools analyze data.tame CLINICAL_STATS", output)
        self.assertIn("Plugin name from 'tametools plugins FILE'.", output)

    def test_cli_quickstart_doctor_and_version_are_available(self) -> None:
        quickstart_code, quickstart_output = self._run_cli("quickstart")
        doctor_code, doctor_output = self._run_cli("doctor")
        version_code, version_output = self._run_cli("--version")

        self.assertEqual(quickstart_code, 0)
        self.assertIn("tametools quickstart", quickstart_output)
        self.assertIn("tametools doctor", quickstart_output)
        self.assertEqual(doctor_code, 0)
        self.assertIn("[doctor]", doctor_output)
        self.assertIn("status: ok", doctor_output)
        self.assertEqual(version_code, 0)
        self.assertRegex(version_output, r"tametools \d+\.\d+\.\d+")

    def test_cli_unknown_command_prints_short_beginner_error(self) -> None:
        code, output = self._run_cli("nonesuch")

        self.assertEqual(code, 2)
        self.assertIn("Error: unknown command: nonesuch", output)
        self.assertIn("tametools --help", output)
        self.assertIn("tametools quickstart", output)
        self.assertNotIn("choose from", output)

    def test_cli_inspection_commands_support_json_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")

            info_code, info_output = self._run_cli("info", str(input_path), "--json")
            columns_code, columns_output = self._run_cli("columns", str(input_path), "--json")
            validate_code, validate_output = self._run_cli("validate", str(input_path), "--json")
            review_code, review_output = self._run_cli("review", str(input_path), "--json")

        self.assertEqual(info_code, 0)
        self.assertEqual(json.loads(info_output)["rows"], 2)
        self.assertEqual(columns_code, 0)
        self.assertIn("columns", json.loads(columns_output))
        self.assertEqual(validate_code, 0)
        self.assertIn("issues", json.loads(validate_output))
        self.assertEqual(review_code, 0)
        review_payload = json.loads(review_output)
        self.assertIn("dataset", review_payload)
        self.assertIn("validation", review_payload)

    def test_cli_columns_aligns_wide_korean_column_names(self) -> None:
        sample = textwrap.dedent(
            """\
            <DATA>
            [[ID::STR]]등록번호\t[[SEX::BY::CATEGORY]]성별\t[[AGE]]나이\t[[ITEM::CATEGORY]]검사항목명\t[[RESULT::<NUM>]]보고값
            P001\t남\t45\tAST\t30
            </DATA>
            """
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "korean.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("columns", str(input_path))

        self.assertEqual(code, 0, output)
        lines = [line for line in output.splitlines() if line and not line.startswith("[")]
        header = lines[0]
        tags_column = _display_width_until(header, "tags")
        expected_tags = ["ID::STR", "SEX::BY::CATEGORY", "AGE", "ITEM::CATEGORY", "RESULT::<NUM>"]
        for line, tag in zip(lines[1:], expected_tags):
            self.assertEqual(_display_width_until(line, tag), tags_column, line)

    def test_cli_review_categorical_summary_hides_top_and_numeric_stats(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")

            code, output = self._run_cli("review", str(input_path))

        self.assertEqual(code, 0, output)
        summary_lines = []
        in_summary = False
        for line in output.splitlines():
            if line == "[review:summary]":
                in_summary = True
                continue
            if in_summary and line.startswith("["):
                break
            if in_summary and line.strip():
                summary_lines.append(line)
        hospital_row = next(line for line in summary_lines if line.startswith("병원ID"))
        self.assertTrue(hospital_row.rstrip().endswith("1.0"), hospital_row)

    def test_cli_review_sections_are_spaced_and_age_standard_line_is_hidden(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")

            code, output = self._run_cli("review", str(input_path))

        self.assertEqual(code, 0, output)
        for section in ("review:columns", "review:summary", "review:SEX", "review:AGE", "review:validation-summary"):
            self.assertIn(f"\n\n[{section}]", output)
        self.assertNotIn("standard: HL7 FHIR Age Quantity", output)
        self.assertNotIn("[review:DATETIME]", output)
        self.assertNotIn("(no DATE/DATETIME columns)", output)

    def test_cli_reads_tame_from_stdin_dash(self) -> None:
        code, output = self._run_cli_with_stdin(SAMPLE_NUM_TAME, "info", "-", "--json")

        payload = json.loads(output)
        self.assertEqual(code, 0)
        self.assertEqual(payload["source"], "-")
        self.assertEqual(payload["rows"], 2)

    def test_cli_convert_writes_tame_to_stdout_dash_without_status_text(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")

            code, output = self._run_cli("convert", str(input_path), "-")

        self.assertEqual(code, 0)
        self.assertTrue(output.startswith("<META>"))
        self.assertIn("<DATA>", output)
        self.assertNotIn("saved:", output)

    def test_cli_completion_command_outputs_shell_script(self) -> None:
        code, output = self._run_cli("completion", "bash")

        self.assertEqual(code, 0)
        self.assertIn("complete -F _tametools_completion tametools", output)
        self.assertIn("analyze", output)
        self.assertIn("--json", output)

    def test_cli_beginner_aliases_preserve_existing_behavior(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            output_path = Path(tmpdir) / "converted.tame"
            input_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")

            check_code, check_output = self._run_cli("check", str(input_path))
            inspect_code, inspect_output = self._run_cli("inspect", str(input_path))
            convert_code, convert_output = self._run_cli("convert", str(input_path), str(output_path))
            analyze_code, analyze_output = self._run_cli("analyze", str(input_path), "NO_SUCH_PLUGIN")
            converted_exists = output_path.exists()

        self.assertEqual(check_code, 0)
        self.assertIn("[check:columns]", check_output)
        self.assertEqual(inspect_code, 0)
        self.assertIn("[review:dataset]", inspect_output)
        self.assertEqual(convert_code, 0)
        self.assertIn("saved:", convert_output)
        self.assertTrue(converted_exists)
        self.assertEqual(analyze_code, 1)
        self.assertIn("error: unknown plugin: NO_SUCH_PLUGIN", analyze_output)

    def test_cli_help_path_does_not_import_pandas_runtime(self) -> None:
        repo_root = Path(__file__).resolve().parents[2]
        env = os.environ.copy()
        env["PYTHONPATH"] = f"{repo_root / 'tametools' / 'src'}:{repo_root}"
        script = textwrap.dedent(
            """
            import contextlib
            import io
            import sys
            from tametools.cli import main

            sys.argv = ["tametools", "--help"]
            with contextlib.redirect_stdout(io.StringIO()):
                code = main()
            print(code)
            print("pandas" in sys.modules)
            """
        )
        result = subprocess.run(
            [sys.executable, "-c", script],
            check=True,
            text=True,
            capture_output=True,
            env=env,
        )

        self.assertEqual(result.stdout.strip().splitlines(), ["0", "False"])

    def test_review_prints_sex_profile_without_changing_input(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별
            M
            M
            male
            badsex
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sex_review.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("review", str(input_path))
            reloaded = read_tame(input_path)

        self.assertEqual(code, 0)
        self.assertIn("male: 3 (M: 2, male: 1)", output)
        self.assertIn("unrecognized: badsex: 1", output)
        self.assertIn("detail_command: tametools review detail", output)
        self.assertNotIn("detail_command: tametools validate", output)
        self.assertEqual(reloaded.df["성별"].tolist(), ["M", "M", "male", "badsex"])

    def test_run_plugin_unknown_plugin_prints_clean_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")

            code, output = self._run_cli("run-plugin", str(input_path), "NO_SUCH_PLUGIN")

        self.assertEqual(code, 1)
        self.assertIn("error: unknown plugin: NO_SUCH_PLUGIN", output)
        self.assertIn("available_plugins:", output)
        self.assertNotIn("Traceback", output)

    def test_cli_global_error_handler_hides_tracebacks_for_common_user_errors(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            missing_path = Path(tmpdir) / "missing.tame"
            broken_path = Path(tmpdir) / "broken.tame"
            sample_path = Path(tmpdir) / "sample.tame"
            broken_path.write_text("<META>\n[bad\n</META>\n<DATA>\na\n1\n</DATA>", encoding="utf-8")
            sample_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")

            cases = [
                self._run_cli("validate", str(missing_path)),
                self._run_cli("run-action", str(missing_path), "FOO"),
                self._run_cli("review", str(broken_path)),
                self._run_cli("run-plugin", str(sample_path), "REFERENCE_INTERVAL", "--option", "BAD_OPTION"),
            ]

        for code, output in cases:
            self.assertNotEqual(code, 0)
            self.assertIn("Error:", output)
            self.assertIn("--debug", output)
            self.assertNotIn("Traceback", output)
            self.assertNotIn("FileNotFoundError", output)
            self.assertNotIn("TOMLDecodeError", output)

    def test_review_detail_prints_column_values_and_issues(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")

            code, output = self._run_cli("review", "detail", str(input_path), "나이")

        self.assertEqual(code, 0)
        self.assertIn("[review:detail]", output)
        self.assertIn("column: 나이", output)
        self.assertIn("[review:detail:values]", output)
        self.assertIn("bad: 1", output)
        self.assertIn("[review:detail:issues]", output)
        self.assertIn("row=4 tag=AGE value='bad'", output)

    def test_review_prints_age_profile_and_tag_selector(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[AGE]]나이\t[[AGE]]추가나이
            0\t2m
            10a\t2mo
            bad\t1d
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "age_review.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("review", str(input_path), "tag:AGE")

        self.assertEqual(code, 0)
        self.assertIn("[review:tag]", output)
        self.assertIn("tag: AGE", output)
        self.assertIn("[review:AGE]", output)
        self.assertIn("default_unit: a (year), stored without suffix", output)
        self.assertIn("distribution_10y: <1: 1, 10-19: 1", output)
        self.assertIn("year (a): 2 (0: 1, 10a -> 10: 1)", output)
        self.assertIn("month (mo): 2 (2m -> 2mo: 1, 2mo: 1)", output)
        self.assertIn("day (d): 1 (1d: 1)", output)
        self.assertIn("unrecognized: bad: 1", output)

    def test_fix_standardize_age_writes_new_dataset(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[AGE]]나이
            10a
            2m
            1day
            0
            bad
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "age_fix.tame"
            output_path = Path(tmpdir) / "age_fixed.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("fix", str(input_path), "--standardize-age", "--output", str(output_path))
            fixed = read_tame(output_path)

        self.assertEqual(code, 0)
        self.assertIn("fix: standardize-age changed_values=3", output)
        self.assertEqual(fixed.df["나이"].tolist(), ["10", "2mo", "1d", "0", "bad"])

    def test_review_prints_category_and_datetime_profiles(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[CATEGORY]]진단군\t[[DATETIME]]채혈시각
            A\t2024-01-01
            B\t2024-01-02 13:30
            A\tbad-date
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "category_datetime_review.tame"
            input_path.write_text(sample, encoding="utf-8")

            category_code, category_output = self._run_cli("review", str(input_path), "tag:CATEGORY")
            datetime_code, datetime_output = self._run_cli("review", str(input_path), "tag:DATETIME")

        self.assertEqual(category_code, 0)
        self.assertIn("[review:CATEGORY]", category_output)
        self.assertIn("categories: 2", category_output)
        self.assertIn("counts: A: 2, B: 1", category_output)
        self.assertEqual(datetime_code, 0)
        self.assertIn("[review:DATETIME]", datetime_output)
        self.assertIn("parseable: 2", datetime_output)
        self.assertIn("unparseable: 1", datetime_output)
        self.assertIn("unrecognized: bad-date: 1", datetime_output)

    def test_review_tag_selector_prints_all_matching_columns(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별\t[[SEX]]보호자성별
            M\tF
            male\tfemale
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "multi_sex.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("review", str(input_path), "tag:SEX")

        self.assertEqual(code, 0)
        self.assertIn("[review:tag]", output)
        self.assertIn("tag: SEX", output)
        self.assertIn("columns: 2", output)
        self.assertIn("column: 성별", output)
        self.assertIn("column: 보호자성별", output)

    def test_review_reports_duplicate_column_resolution(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            SEX\tSEX
            M\tF
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "duplicate_review.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("review", str(input_path))

        self.assertEqual(code, 0)
        self.assertIn("[review:duplicate-columns]", output)
        self.assertIn("original_name: SEX -> SEX, SEX.1", output)
        self.assertIn("SEX.1", output)

    def test_fix_standardize_sex_writes_new_dataset(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별
            M
            F
            male
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sex_fix.tame"
            output_path = Path(tmpdir) / "sex_fixed.tame"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("fix", str(input_path), "--standardize-sex", "--output", str(output_path))
            fixed = read_tame(output_path)

        self.assertEqual(code, 0)
        self.assertIn("fix: standardize-sex changed_values=2", output)
        self.assertEqual(fixed.df["성별"].tolist(), ["male", "female", "male"])

    def test_validate_can_fail_on_issues_without_saving(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")

            code, output = self._run_cli("validate", str(input_path), "--fail-on-issues")

        self.assertEqual(code, 1)
        self.assertIn("issues: 3", output)
        self.assertNotIn("saved:", output)

    def test_save_converts_dataset_with_writer_standardization(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[SEX]]성별
            M
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sex_save.tame"
            output_path = Path(tmpdir) / "sex_save.xlsx"
            input_path.write_text(sample, encoding="utf-8")

            code, output = self._run_cli("save", str(input_path), str(output_path))
            reloaded = read_xlsx(output_path)

        self.assertEqual(code, 0)
        self.assertIn("saved:", output)
        self.assertEqual(reloaded.df["성별"].tolist(), ["male"])

    def test_actions_and_run_action_commands(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "action.tame"
            output_path = Path(tmpdir) / "tagged.tame"
            detailed_path = Path(tmpdir) / "tagged_detailed.tame"
            input_path.write_text(SAMPLE_ACTION_TAME, encoding="utf-8")

            list_code, list_output = self._run_cli("actions", str(input_path))
            run_code, run_output = self._run_cli(
                "run-action",
                str(input_path),
                "ADD_SEX_AGE_TAGS",
                "--output",
                str(output_path),
            )
            detailed_code, detailed_output = self._run_cli(
                "run-action",
                str(input_path),
                "ADD_SEX_AGE_TAGS",
                "--output",
                str(detailed_path),
                "--log-level",
                "detailed",
            )
            tagged = read_tame(output_path)
            detailed = read_tame(detailed_path)

        self.assertEqual(list_code, 0)
        self.assertIn("[ADD_SEX_AGE_TAGS]", list_output)
        self.assertIn("type: TAG_COLUMNS", list_output)
        self.assertEqual(run_code, 0)
        self.assertEqual(detailed_code, 0)
        self.assertIn("tagged_columns=2", run_output)
        self.assertIn("tagged_columns=2", detailed_output)
        self.assertIn("saved:", run_output)
        self.assertEqual(tagged.columns[0].tags, ("SEX",))
        self.assertEqual(tagged.columns[1].tags, ("AGE",))
        self.assertEqual([entry["OPERATION"] for entry in tagged.meta["LOG"]], ["CLI:RUN-ACTION", "EXPORT_PREPARATION"])
        self.assertEqual(
            [entry["OPERATION"] for entry in detailed.meta["LOG"]],
            ["ACTION:ADD_SEX_AGE_TAGS", "CLI:RUN-ACTION", "EXPORT_PREPARATION"],
        )
        self.assertIn("input_hash", detailed.meta["LOG"][0]["PARAMS"])

    def test_run_plugin_output_inherits_existing_log_chain(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "logged_input.tame"
            output_path = Path(tmpdir) / "plugin_output.tame"
            input_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")
            logged = append_log_entry(read_tame(input_path), action="previous-step", message="existing chain")
            write_tame(input_path, logged)

            code, output = self._run_cli(
                "run-plugin",
                str(input_path),
                "CHEMISTRY_ANALYSIS",
                "--option",
                "MODE=ITEM_COUNTS",
                "--output",
                str(output_path),
            )
            result = read_tame(output_path)

        self.assertEqual(code, 0)
        self.assertIn("saved:", output)
        self.assertEqual([entry["OPERATION"] for entry in result.meta["LOG"]], ["PREVIOUS-STEP", "CLI:RUN-PLUGIN"])
        self.assertEqual(result.meta["LOG"][0]["NOTES"], "existing chain")
        self.assertNotIn("PARENT", result.meta["LOG"][-1])
        self.assertNotIn("OUTPUT", result.meta["LOG"][-1])
        self.assertEqual(
            result.meta["LOG"][-1]["PARAMS"],
            {
                "command": (
                    f"tametools run-plugin {input_path} CHEMISTRY_ANALYSIS --option MODE=ITEM_COUNTS "
                    f"--output {output_path}"
                )
            },
        )

    def test_action_pipeline_commands(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.KEEP_A]
            TYPE = "ROW.INCLUDE"
            COLUMN = "group"
            VALUE = "A"

            [ACTIONS.ADD_BATCH]
            TYPE = "ADD_COLUMN"
            NAME = "batch"
            VALUE = "one"

            [ACTION_PIPELINES]
            DEFAULT = ["KEEP_A", "ADD_BATCH"]
            </META>
            <DATA>
            group\tvalue
            A\t1
            B\t2
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "action_pipeline.tame"
            output_path = Path(tmpdir) / "action_pipeline_out.tame"
            detailed_path = Path(tmpdir) / "action_pipeline_detailed.tame"
            input_path.write_text(sample, encoding="utf-8")
            original_text = input_path.read_text(encoding="utf-8")

            list_code, list_output = self._run_cli("action-pipelines", str(input_path))
            run_code, run_output = self._run_cli(
                "run-action-pipeline",
                str(input_path),
                "DEFAULT",
                "--output",
                str(output_path), "--log-level", "simple",
            )
            detailed_code, detailed_output = self._run_cli(
                "run-action-pipeline",
                str(input_path),
                "DEFAULT",
                "--output",
                str(detailed_path),
                "--log-level",
                "detailed",
            )
            logs_code, logs_output = self._run_cli("logs", str(output_path), "--full")
            logs_json_code, logs_json_output = self._run_cli("logs", str(output_path), "--json", "--limit", "1")
            clear_path = Path(tmpdir) / "action_pipeline_public.tame"
            clear_code, clear_output = self._run_cli("clear-log", str(output_path), "--output", str(clear_path))
            transformed = read_tame(output_path)
            detailed = read_tame(detailed_path)
            cleared = read_tame(clear_path)
            original = read_tame(input_path)
            overwrite_code, overwrite_output = self._run_cli(
                "run-action-pipeline",
                str(input_path),
                "DEFAULT",
                "--output",
                str(input_path),
            )
            final_input_text = input_path.read_text(encoding="utf-8")

        self.assertEqual(list_code, 0)
        self.assertIn("[DEFAULT]", list_output)
        self.assertIn("KEEP_A", list_output)
        self.assertEqual(run_code, 0)
        self.assertEqual(detailed_code, 0)
        self.assertIn("action_pipeline: DEFAULT", run_output)
        self.assertIn("action_pipeline: DEFAULT", detailed_output)
        self.assertIn("saved:", run_output)
        self.assertEqual(final_input_text, original_text)
        self.assertNotIn("LOG", original.meta)
        self.assertEqual(original.df["group"].tolist(), ["A", "B"])
        self.assertNotIn("batch", original.df.columns)
        self.assertEqual(transformed.df["group"].tolist(), ["A"])
        self.assertEqual(transformed.df["batch"].tolist(), ["one"])
        self.assertEqual(
            [entry["OPERATION"] for entry in transformed.meta["LOG"]],
            ["CLI:RUN-ACTION-PIPELINE"],
        )
        self.assertEqual(
            [entry["OPERATION"] for entry in detailed.meta["LOG"]],
            ["ACTION:KEEP_A", "ACTION:ADD_BATCH", "CLI:RUN-ACTION-PIPELINE"],
        )
        save_log = transformed.meta["LOG"][0]
        self.assertNotIn("PARENT", save_log)
        self.assertNotIn("OUTPUT", save_log)
        self.assertEqual(save_log["PARAMS"]["command"], f"tametools run-action-pipeline {input_path} DEFAULT --output {output_path} --log-level simple")
        self.assertEqual(list(save_log["PARAMS"].keys()), ["command"])
        detailed_save_log = detailed.meta["LOG"][-1]
        self.assertEqual(detailed_save_log["PARENT"], str(input_path))
        self.assertEqual(detailed_save_log["OUTPUT"], str(detailed_path))
        self.assertEqual(detailed_save_log["PARAMS"]["source_paths"], [str(input_path)])
        self.assertEqual(detailed_save_log["PARAMS"]["output_path"], str(detailed_path))
        self.assertTrue(detailed_save_log["PARAMS"]["meta_preserved"])
        self.assertIn("ACTIONS", detailed_save_log["PARAMS"]["meta_keys"])
        self.assertIn("ACTION_PIPELINES", transformed.meta)
        self.assertEqual(logs_code, 0)
        self.assertIn("[logs]", logs_output)
        self.assertNotIn("ACTION:KEEP_A", logs_output)
        self.assertIn("CLI:RUN-ACTION-PIPELINE", logs_output)
        self.assertIn(str(input_path), logs_output)
        self.assertIn(str(output_path), logs_output)
        self.assertIn("entries: 1", logs_output)
        self.assertEqual(logs_json_code, 0)
        self.assertEqual(json.loads(logs_json_output)[0]["OPERATION"], "CLI:RUN-ACTION-PIPELINE")
        self.assertEqual(json.loads(logs_json_output)[0]["PARAMS"]["command"], f"tametools run-action-pipeline {input_path} DEFAULT --output {output_path} --log-level simple")
        self.assertEqual(clear_code, 0)
        self.assertIn("removed_log_entries: 1", clear_output)
        self.assertNotIn("LOG", cleared.meta)
        self.assertNotIn("[[LOG]]", cleared.raw_sections["META"])
        self.assertEqual(cleared.df["batch"].tolist(), ["one"])
        self.assertEqual(overwrite_code, 1)
        self.assertIn("Output path must be different from input path", overwrite_output)

    def test_cli_tutorial14_clean_pipeline_keeps_age_selector_unambiguous(self) -> None:
        repo_root = Path(__file__).resolve().parents[2]
        base = repo_root / "tutorial" / "14_capability_probe" / "clinical_chem.tame"
        meta = repo_root / "tutorial" / "14_capability_probe" / "examples" / "preprocessing.meta.tame"
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "cleaned.tame"
            code, output = self._run_cli(
                "run-action-pipeline",
                str(base),
                "CLEAN",
                "--meta",
                str(meta),
                "--output",
                str(output_path),
            )
            cleaned = read_tame(output_path)

        self.assertEqual(code, 0, output)
        self.assertIn("action_pipeline: CLEAN", output)
        self.assertEqual([column.name for column in cleaned.columns_with_tag("AGE(baseline)")], ["나이"])
        visit_age = [column for column in cleaned.columns if column.name == "방문시연령"][0]
        self.assertEqual(visit_age.tags, ("AGE(visit)", "NUM"))
        self.assertEqual([column.name for column in cleaned.columns_with_tag("AGE")], ["나이", "방문시연령"])
        self.assertEqual([column.name for column in cleaned.columns_with_tag("AGE_GROUP")], ["연령군"])

    def test_apply_preset_command_writes_reusable_meta(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            검사항목명\t보고값\t장비명\t접수일\t검사시간\t등록번호\tsex\tage
            AST\t10\tA1\t2024-01-01 08:00\t2024-01-01 08:30\tP1\tM\t30
            ALT\t20\tA2\t2024-01-01 09:00\t2024-01-01 09:30\tP2\tF\t40
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "raw_chem.tame"
            output_path = Path(tmpdir) / "preset_chem.tame"
            input_path.write_text(sample, encoding="utf-8")

            list_code, list_output = self._run_cli("presets")
            run_code, run_output = self._run_cli("apply-preset", str(input_path), "--output", str(output_path))
            transformed = read_tame(output_path)

        self.assertEqual(list_code, 0)
        self.assertIn("[CLINICAL_CHEMISTRY]", list_output)
        self.assertEqual(run_code, 0)
        self.assertIn("analyses=13", run_output)
        self.assertIn("saved:", run_output)
        self.assertIsNotNone(transformed.first_column_with_tag("TESTNAME"))
        self.assertIn("TAT_BY_TEST", transformed.meta["ANALYSES"])

    def test_fix_num_comparator_command_supports_delete_and_value_handling(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[RESULT::NUM]]보고값\t[[ID]]등록번호
            <3\tA
            4\tB
            >5\tC
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "num_fix.tame"
            delete_path = Path(tmpdir) / "num_delete.tame"
            value_path = Path(tmpdir) / "num_value.tame"
            input_path.write_text(sample, encoding="utf-8")

            delete_code, delete_output = self._run_cli("fix", str(input_path), "--fix-num-comparator", "--output", str(delete_path))
            value_code, value_output = self._run_cli("fix", str(input_path), "--fix-num-comparator", "value", "--output", str(value_path))
            deleted = read_tame(delete_path)
            valued = read_tame(value_path)

        self.assertEqual(delete_code, 0)
        self.assertIn("handling=delete", delete_output)
        self.assertEqual(deleted.df["등록번호"].tolist(), ["B"])
        self.assertEqual(value_code, 0)
        self.assertIn("handling=value", value_output)
        self.assertEqual(valued.df["보고값"].tolist(), ["3", "4", "5"])

    def test_tags_to_meta_and_tags_to_header_commands(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [TAGS]
            sex = ["CATEGORY"]
            </META>
            <DATA>
            [[SEX]]sex
            M
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "mixed_tags.tame"
            meta_path = Path(tmpdir) / "tags_meta.tame"
            header_path = Path(tmpdir) / "tags_header.tame"
            input_path.write_text(sample, encoding="utf-8")

            meta_code, meta_output = self._run_cli("tags-to-meta", str(input_path), "--output", str(meta_path))
            header_code, header_output = self._run_cli("tags-to-header", str(input_path), "--output", str(header_path))
            meta_text = meta_path.read_text(encoding="utf-8")
            header_text = header_path.read_text(encoding="utf-8")

        self.assertEqual(meta_code, 0)
        self.assertIn("tag_storage=meta", meta_output)
        self.assertIn("[COLUMN.sex]", meta_text)
        self.assertIn('TAGS = ["CATEGORY", "SEX"]', meta_text)
        self.assertNotIn("[[SEX]]sex", meta_text)
        self.assertEqual(header_code, 0)
        self.assertIn("tag_storage=header", header_output)
        self.assertNotIn("[TAGS]", header_text)
        self.assertIn("[[CATEGORY::SEX]]sex", header_text)


class EvaluationTest(unittest.TestCase):
    def test_evaluate_dataset_builds_quality_workflow_and_output_tables(self) -> None:
        baseline_tame = textwrap.dedent(
            """
            <META>
            </META>
            <DATA>
            등록번호\t성별\t나이\t검사항목명\t보고값\t진료과
            A0001\tF\t32\tAST\t26\tIM
            A0002\t남\t2m\tAST\t<3\tGS
            A0003\tbadsex\tbad\tALT\tpositive\tIM
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            baseline_path = Path(tmpdir) / "baseline.tame"
            tame_path = Path(tmpdir) / "sample.tame"
            report_dir = Path(tmpdir) / "report"
            baseline_path.write_text(baseline_tame, encoding="utf-8")
            tame_path.write_text(SAMPLE_TAME, encoding="utf-8")
            baseline = read_tame(baseline_path)
            dataset = read_tame(tame_path)

            report = evaluate_dataset(
                dataset,
                baseline_dataset=baseline,
                baseline_steps=8,
                tame_steps=2,
                baseline_minutes=30.0,
                tame_minutes=5.0,
            )
            written = write_evaluation_report(report, report_dir)
            dataset_quality_written = (report_dir / "dataset_quality.csv").exists()
            workflow_path = str(report_dir / "workflow_comparison.csv")

        quality = report.dataset_quality.set_index("dataset")
        self.assertEqual(int(quality.loc["existing", "tagged_columns"]), 0)
        self.assertGreater(int(quality.loc["TAME", "tagged_columns"]), 0)
        self.assertEqual(int(report.workflow_comparison.loc[0, "preprocessing_steps"]), 8)
        self.assertEqual(int(report.workflow_comparison.loc[1, "preprocessing_steps"]), 2)
        self.assertEqual(int(report.output_comparison.loc[0, "cell_mismatches"]), 1)
        self.assertTrue(dataset_quality_written)
        self.assertIn(workflow_path, written)

    def test_execute_work_runs_evaluate_step(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        meta = dict(dataset.meta)
        meta["WORKS"] = {"DEFAULT": ["EVALUATE"]}
        meta["EVALUATE"] = {"TAME_STEPS": 2, "TAME_MINUTES": 5.0}
        result = execute_work(dataset.replace(meta=meta), "DEFAULT")

        self.assertEqual(result.outputs[0].name, "EVALUATE")
        self.assertIn("dataset_quality", result.outputs[0].tables)
        self.assertIn("workflow_comparison", result.outputs[0].tables)
        self.assertEqual(len(result.final_dataset.df), len(dataset.df))


class PluginTest(unittest.TestCase):
    def test_list_plugins_includes_reference_interval(self) -> None:
        plugins = {plugin.name for plugin in list_plugins()}
        self.assertIn("REFERENCE_INTERVAL", plugins)
        self.assertIn("RI", plugins)
        self.assertIn("CHEMISTRY_ANALYSIS", plugins)

    def test_abnormal_flag_plugin_flags_and_rates(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"
            </META>
            <DATA>
            [[ITEM]]검사항목명\t[[RESULT::NUM]]보고값\t[[REF_LOW::NUM]]하한\t[[REF_HIGH::NUM]]상한
            AST\t10\t0\t40
            AST\t60\t0\t40
            ALT\t5\t10\t40
            ALT\t20\t10\t40
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "flag.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        dataset.meta.setdefault("COLUMN", {}).update({name: {"UNIT": "U/L"} for name in ("보고값", "하한", "상한")})
        flagged = run_plugin(dataset, "ABNORMAL_FLAG", dataset.meta, "ABNORMAL_FLAG", {"MODE": "FLAG"})
        self.assertIsNotNone(flagged)
        self.assertEqual(flagged.dataset.df["이상플래그"].tolist(), ["N", "H", "L", "N"])
        self.assertEqual(flagged.dataset.columns[-2].tags, ("FLAG", "INTERPRETATION", "CATEGORY"))
        self.assertIn("이상플래그_reason", flagged.dataset.df)

        rate = run_plugin(dataset, "ABNORMAL_FLAG", dataset.meta, "ABNORMAL_FLAG", {"MODE": "RATE"})
        rates = {row["검사항목명"]: row for _, row in rate.table.iterrows()}
        self.assertEqual(int(rates["AST"]["abnormal_n"]), 1)
        self.assertEqual(float(rates["AST"]["abnormal_rate"]), 50.0)
        self.assertEqual(int(rates["ALT"]["low_n"]), 1)

    def test_abnormal_flag_plugin_uses_pivot_context_reference_limits(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"

            [COLUMN.AST.PIVOT_CONTEXT]
            UNIT = "U/L"
            REF_LOW = "0"
            REF_HIGH = "40"

            [COLUMN.ALT.PIVOT_CONTEXT]
            UNIT = "U/L"
            REF_LOW = "10"
            REF_HIGH = "40"
            </META>
            <DATA>
            [[ID(sample)::STR]]sample\t[[RESULT::NUM]]AST\t[[RESULT::NUM]]ALT
            S1\t10\t5
            S2\t60\t20
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "wide_flag.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        flagged = run_plugin(dataset, "ABNORMAL_FLAG", dataset.meta, "ABNORMAL_FLAG", {"MODE": "FLAG"})
        self.assertEqual(flagged.dataset.df["AST_이상플래그"].tolist(), ["N", "H"])
        self.assertEqual(flagged.dataset.df["ALT_이상플래그"].tolist(), ["L", "N"])
        self.assertEqual(flagged.dataset.columns[-2].tags, ("FLAG", "INTERPRETATION", "CATEGORY"))
        self.assertEqual(flagged.dataset.column_metadata("AST_이상플래그")["SOURCE_RESULT"], "AST")
        self.assertEqual(flagged.dataset.column_metadata("AST_이상플래그")["REFERENCE_SOURCE"], "PIVOT_CONTEXT")

        rate = run_plugin(dataset, "ABNORMAL_FLAG", dataset.meta, "ABNORMAL_FLAG", {"MODE": "RATE"})
        rates = {row["test"]: row for _, row in rate.table.iterrows()}
        self.assertEqual(int(rates["AST"]["high_n"]), 1)
        self.assertEqual(int(rates["ALT"]["low_n"]), 1)
        self.assertEqual(float(rates["AST"]["abnormal_rate"]), 50.0)

    def test_autoverification_plugin_outputs_decisions_and_validates(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"
            </META>
            <DATA>
            [[PATIENT_ID::ID]]pid\t[[SAMPLE_ID]]sample\t[[TESTNAME]]test\t[[RESULT_TIME::DATETIME]]time\t[[INSTRUMENT]]inst\t[[RESULT::NUM]]result\t[[REF_LOW::NUM]]low\t[[REF_HIGH::NUM]]high
            P1\tS1\tGLU\t2024-01-01 08:00\tA\t90\t70\t110
            P1\tS2\tGLU\t2024-01-02 08:00\tA\t180\t70\t110
            P2\tS3\tGLU\t2024-01-01 08:00\tB\t450\t70\t110
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "autoverify.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)
        dataset.meta["COLUMN"] = {name: {"UNIT": "mg/dL"} for name in ("result", "low", "high")}

        output = run_plugin(
            dataset,
            "AUTOVERIFICATION",
            dataset.meta,
            "AUTOVERIFICATION",
            {"CRITICAL_HIGH_BY_TEST": {"GLU": 400}, "HOLD_INSTRUMENTS": "B"},
        )
        decisions = output.table["decision"].tolist()

        self.assertEqual(decisions, ["PASS", "REVIEW", "HOLD"])
        self.assertIn("DELTA_PERCENT", output.table.loc[1, "rule_id"])
        self.assertIn("CRITICAL_HIGH", output.table.loc[2, "rule_id"])
        self.assertEqual(validate_dataset(output.dataset).issues, [])

    def test_method_comparison_recovers_known_slope(self) -> None:
        lines = ["[[ID::STR]]pid\t[[ITEM]]test\t[[INSTRUMENT]]inst\t[[RESULT::NUM]]val"]
        for i in range(40):
            x = 10 + i * 2
            y = 1.1 * x + 1
            lines.append(f"P{i}\tAST\tA\t{x}")
            lines.append(f"P{i}\tAST\tB\t{y:.2f}")
        sample = "<META>\n[INFO]\n</META>\n<DATA>\n" + "\n".join(lines) + "\n</DATA>\n"
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "mc.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)
        dataset.meta["COLUMN"] = {"val": {"UNIT": "U/L"}}
        out = run_plugin(dataset, "METHOD_COMPARISON", dataset.meta, "METHOD_COMPARISON", {"METHOD_A": "A", "METHOD_B": "B"})
        row = out.table.iloc[0]
        self.assertEqual(int(row["n"]), 40)
        self.assertAlmostEqual(float(row["pb_slope"]), 1.1, places=2)
        self.assertAlmostEqual(float(row["deming_slope"]), 1.1, places=2)

    def test_group_test_correlation_qc_trend_plugins(self) -> None:
        rows = ["[[ID::STR]]pid\t[[ITEM]]test\t[[BY]]hosp\t[[RESULT::NUM]]val\t[[DATETIME]]when"]
        import datetime as _dt
        base = _dt.date(2024, 1, 1)
        for i in range(40):
            hosp = "H1" if i % 2 == 0 else "H2"
            val = 20 + (i % 10) + (0 if hosp == "H1" else 15)
            day = base + _dt.timedelta(days=i * 3)
            rows.append(f"P{i}\tAST\t{hosp}\t{val}\t{day.isoformat()} 09:00:00")
        sample = "<META>\n[SETTINGS]\nVALIDATE_ERROR = \"REPORT\"\n</META>\n<DATA>\n" + "\n".join(rows) + "\n</DATA>\n"
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "stats.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        gt = run_plugin(dataset, "GROUP_TEST", dataset.meta, "GROUP_TEST", {"GROUP": "hosp"})
        self.assertEqual(gt.table.iloc[0]["method"], "mann_whitney")
        self.assertLess(float(gt.table.iloc[0]["p_value"]), 0.05)

        qc = run_plugin(dataset, "QC_ANALYSIS", dataset.meta, "QC_ANALYSIS", {"MODE": "PRECISION"})
        self.assertIn("cv_percent", list(qc.table.columns))
        self.assertGreater(float(qc.table.iloc[0]["cv_percent"]), 0)

        with self.assertRaisesRegex(ValueError, "RUN_ID"):
            run_plugin(dataset, "QC_ANALYSIS", dataset.meta, "QC_ANALYSIS", {"MODE": "WESTGARD"})
        westgard = run_plugin(dataset, "QC_ANALYSIS", dataset.meta, "QC_ANALYSIS", {"MODE": "LEVEY_JENNINGS"})
        self.assertIn("NONE", westgard.table["violation"].tolist())
        self.assertEqual(validate_dataset(westgard.dataset).issues, [])

        trend = run_plugin(dataset, "RESULT_TREND", dataset.meta, "RESULT_TREND", {"PERIOD": "M"})
        self.assertIn("moving_avg", list(trend.table.columns))
        self.assertGreater(len(trend.table), 1)

        corr = run_plugin(dataset, "CORRELATION", dataset.meta, "CORRELATION", {})
        self.assertIsNotNone(corr)

    def test_roc_analysis_plugin_summary(self) -> None:
        rows = ["[[RESULT::NUM]]score\t[[LABEL]]label"]
        for i in range(60):
            disease = i % 2
            score = (70 if disease else 40) + (i % 5)
            rows.append(f"{score}\t{'1' if disease else '0'}")
        sample = "<META>\n[INFO]\n</META>\n<DATA>\n" + "\n".join(rows) + "\n</DATA>\n"
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "roc.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)
        out = run_plugin(dataset, "ROC_ANALYSIS", dataset.meta, "ROC_ANALYSIS", {"LABEL": "label", "POSITIVE": "1"})
        row = out.table.iloc[0]
        self.assertEqual(int(row["n_pos"]), 30)
        self.assertGreater(float(row["auc"]), 0.95)
        self.assertGreaterEqual(float(row["best_cutoff"]), 44)
        self.assertEqual(out.dataset.first_column_with_tag("AUC").name, "auc")
        self.assertEqual(out.dataset.first_column_with_tag("THRESHOLD").name, "best_cutoff")
        self.assertEqual(out.dataset.first_column_with_tag("SENSITIVITY").name, "sensitivity")

    def test_clinical_chemistry_plugin_outputs_tat_with_chart_meta(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"
            </META>
            <DATA>
            [[TESTNAME]]검사항목명\t[[RESULT::NUM]]보고값\t[[INSTRUMENT]]장비명\t[[RECEIVED_AT::DATETIME]]접수일\t[[RESULT_TIME::DATETIME]]검사시간\t[[ID]]등록번호\t[[SEX]]sex\t[[AGE]]age
            AST\t10\tA1\t2024-01-01 08:00\t2024-01-01 08:30\tP1\tM\t30
            AST\t20\tA1\t2024-01-01 09:00\t2024-01-01 09:45\tP2\tF\t40
            ALT\t30\tA2\t2024-01-01 10:00\t2024-01-01 10:20\tP3\tM\t50
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "chem_tat.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        output = run_plugin(dataset, "CHEMISTRY_ANALYSIS", dataset.meta, "CHEMISTRY_ANALYSIS", {"MODE": "TAT_BY_TEST"})

        self.assertIsNotNone(output)
        self.assertEqual(set(output.dataset.df["검사항목명"]), {"AST", "ALT"})
        ast = output.dataset.df.loc[output.dataset.df["검사항목명"] == "AST"].iloc[0]
        self.assertEqual(float(ast["중앙값TAT분"]), 37.5)
        self.assertEqual(output.dataset.meta["VISUALIZATIONS"]["MAIN"]["Y"], "중앙값TAT분")
        self.assertEqual(output.dataset.first_column_with_tag("MEDIAN").name, "중앙값TAT분")
        self.assertEqual(output.dataset.first_column_with_tag("DURATION").name, "평균TAT분")

    def test_clinical_chemistry_preset_adds_tags_actions_and_analyses(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            검사항목명\t보고값\t장비명\t접수일\t검사시간\t등록번호\tsex\tage
            AST\t10\tA1\t2024-01-01 08:00\t2024-01-01 08:30\tP1\tM\t30
            AST\t20\tA1\t2024-01-01 09:00\t2024-01-01 09:45\tP2\tF\t40
            ALT\t30\tA2\t2024-01-01 10:00\t2024-01-01 10:20\tP3\tM\t50
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "raw_chem.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        output = apply_clinical_chemistry_preset(dataset)
        preset_dataset = output.dataset
        analysis = preset_dataset.meta["ANALYSES"]["TAT_BY_TEST"]
        analysis_output = run_plugin(preset_dataset, analysis["PLUGIN"], preset_dataset.meta, "TAT_BY_TEST", analysis["OPTIONS"])

        self.assertIsNotNone(preset_dataset.first_column_with_tag("TESTNAME"))
        self.assertIsNotNone(preset_dataset.first_column_with_tag("RESULT"))
        self.assertIsNotNone(preset_dataset.first_column_with_tag("RECEIVED_AT"))
        self.assertIn("CORE_ANALYSIS_VIEW", preset_dataset.meta["ACTION_PIPELINES"])
        self.assertEqual(len(preset_dataset.meta["ANALYSES"]), 13)
        self.assertEqual(analysis_output.dataset.meta["VISUALIZATIONS"]["MAIN"]["Y"], "중앙값TAT분")

    def test_reference_interval_plugin_outputs_chainable_tame_dataset(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "ri_input.tame"
            output_path = Path(tmpdir) / "ri_output.tame"
            input_path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

            output = run_plugin(dataset, "REFERENCE_INTERVAL", dataset.meta, "REFERENCE_INTERVAL", {})
            self.assertIsNotNone(output)
            self.assertIsNotNone(output.dataset)
            result_dataset = output.dataset
            write_tame(output_path, result_dataset)
            reloaded = read_tame(output_path)

        self.assertTrue(any(column.has_tag("TESTNAME") for column in reloaded.columns))
        self.assertTrue(any(column.has_tag("SEX_GROUP") for column in reloaded.columns))
        self.assertTrue(any(column.has_tag("AGE_GROUP") for column in reloaded.columns))
        self.assertEqual(set(reloaded.df["group_level"]), {"TESTNAME", "TESTNAME+SEX", "TESTNAME+AGE", "TESTNAME+SEX+AGE"})
        self.assertIn("male", set(reloaded.df["sex_group"]))
        self.assertIn("female", set(reloaded.df["sex_group"]))
        self.assertIn("70+", set(reloaded.df["age_group"]))
        self.assertEqual(reloaded.column_metadata("ref_width")["LABEL"], "참고치폭")
        validation = validate_dataset(reloaded)
        self.assertEqual(validation.issues, [])

    def test_reference_interval_plugin_groups_zero_month_day_ages_with_five_year_bins(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            VALIDATE_ERROR = "REPORT"
            </META>
            <DATA>
            [[TESTNAME]]검사항목명\t[[AGE]]나이\t[[RESULT::NUM]]보고값
            AST\t0\t10
            AST\t6mo\t11
            AST\t10d\t12
            AST\t1\t13
            AST\t5\t14
            AST\t10\t15
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "ri_age_bins.tame"
            input_path.write_text(sample, encoding="utf-8")
            dataset = read_tame(input_path)

        output = run_plugin(dataset, "REFERENCE_INTERVAL", dataset.meta, "REFERENCE_INTERVAL", {"AGE_BIN_WIDTH": 5})
        age_rows = output.dataset.df.loc[output.dataset.df["group_level"] == "TESTNAME+AGE"]

        self.assertEqual(set(age_rows["age_group"]), {"<1", "1-4", "5-9", "10-14"})
        under_one = age_rows.loc[age_rows["age_group"] == "<1"].iloc[0]
        self.assertEqual(int(under_one["n"]), 3)
        self.assertEqual(float(under_one["age_low"]), 0.0)
        self.assertEqual(float(under_one["age_high"]), 1.0)

    def test_reference_interval_plugin_warns_when_sex_age_tags_missing(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "ri_missing.tame"
            input_path.write_text(SAMPLE_RI_MISSING_TAGS_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

        output = run_plugin(dataset, "REFERENCE_INTERVAL", dataset.meta, "REFERENCE_INTERVAL", {})
        self.assertIsNotNone(output)
        self.assertIsNotNone(output.dataset)
        self.assertTrue(any("Missing SEX tag" in warning for warning in output.warnings or []))
        self.assertTrue(any("Missing AGE tag" in warning for warning in output.warnings or []))
        self.assertEqual(set(output.dataset.df["group_level"]), {"TESTNAME"})

    def test_execute_work_runs_reference_interval_plugin(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "ri_input.tame"
            input_path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

        result = execute_work(dataset, "RI")
        self.assertEqual(result.outputs[0].name, "REFERENCE_INTERVAL")
        self.assertIn("ref_low", list(result.final_dataset.df.columns))
        self.assertIn("ref_high", list(result.final_dataset.df.columns))


class TransformTest(unittest.TestCase):
    def test_anonymize_dataset_hashes_ids_and_drops_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "anon.tame"
            path.write_text(SAMPLE_ANON_TAME, encoding="utf-8")
            dataset = read_tame(path)

        transformed, mappings = anonymize_dataset(dataset)
        self.assertIn("병원ID", mappings)
        self.assertIn("등록번호", mappings)
        self.assertNotIn("이름", list(transformed.df.columns))
        self.assertTrue(str(transformed.df.loc[0, "병원ID"]).startswith("anon_"))
        self.assertTrue(str(transformed.df.loc[0, "등록번호"]).startswith("anon_"))

    def test_write_mapping_tables_saves_manifest_and_csv_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "anon.tame"
            mapping_dir = Path(tmpdir) / "mappings"
            path.write_text(SAMPLE_ANON_TAME, encoding="utf-8")
            dataset = read_tame(path)

            _, mappings = anonymize_dataset(dataset)
            manifest = write_mapping_tables(mappings, mapping_dir)
            manifest_frame = pd.read_csv(mapping_dir / "manifest.csv", dtype={"rows": "int64"})
            hospital_mapping = pd.read_csv(mapping_dir / "병원ID.mapping.csv", dtype=str)

        self.assertEqual(set(manifest["column"]), {"병원ID", "등록번호"})
        self.assertEqual(set(manifest_frame["column"]), {"병원ID", "등록번호"})
        self.assertIn("anonymized", list(hospital_mapping.columns))
        self.assertTrue(str(hospital_mapping.loc[0, "anonymized"]).startswith("anon_"))

    def test_split_comparator_columns_creates_cmp_and_num_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "anon.tame"
            path.write_text(SAMPLE_ANON_TAME, encoding="utf-8")
            dataset = read_tame(path)

        transformed = split_comparator_columns(dataset, target_columns=["보고값"])
        self.assertIn("보고값__cmp", list(transformed.df.columns))
        self.assertIn("보고값__num", list(transformed.df.columns))
        self.assertEqual(transformed.df.loc[0, "보고값__cmp"], "LT")
        self.assertEqual(float(transformed.df.loc[0, "보고값__num"]), 30.0)
        self.assertEqual(transformed.df.loc[1, "보고값__cmp"], "EQ")

    def test_fix_num_comparator_values_deletes_or_strips_inequality_values(self) -> None:
        sample = textwrap.dedent(
            """
            <META></META>
            <DATA>
            [[RESULT::NUM]]보고값\t[[ID]]등록번호
            <3\tA
            4\tB
            >=5\tC
            bad\tD
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "num_fix.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        deleted, delete_table = fix_num_comparator_values(dataset, handling="delete")
        valued, value_table = fix_num_comparator_values(dataset, handling="value")

        self.assertEqual(deleted.df["등록번호"].tolist(), ["B", "D"])
        self.assertEqual(int(delete_table["affected_cells"].sum()), 2)
        self.assertEqual(valued.df["보고값"].tolist(), ["3", "4", "5", "bad"])
        self.assertEqual(int(value_table["affected_cells"].sum()), 2)

    def test_harmonize_comparator_thresholds_unifies_and_rewrites_band(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "harmonize.tame"
            path.write_text(SAMPLE_HARMONIZE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        transformed, preview = harmonize_comparator_thresholds(dataset)
        self.assertFalse(preview.empty)
        self.assertEqual(transformed.df.loc[0, "보고값"], "<30")
        self.assertEqual(transformed.df.loc[1, "보고값"], "<30")
        self.assertEqual(transformed.df.loc[2, "보고값"], "<30")
        self.assertEqual(transformed.df.loc[3, "보고값"], "18")

    def test_harmonize_comparator_thresholds_uses_item_tag_not_korean_column_name(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [SETTINGS]
            CRR = "HARMONIZE"
            </META>
            <DATA>
            [[ITEM]]test_name\t[[RESULT::<NUM>]]value
            AST\t<30
            AST\t<20
            AST\t22
            ALT\t22
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "harmonize_item_tag.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        transformed, preview = harmonize_comparator_thresholds(dataset)
        ast_rows = transformed.df.loc[transformed.df["test_name"] == "AST", "value"].tolist()
        alt_rows = transformed.df.loc[transformed.df["test_name"] == "ALT", "value"].tolist()

        self.assertFalse(preview.empty)
        self.assertEqual(ast_rows, ["<30", "<30", "<30"])
        self.assertEqual(alt_rows, ["22"])

    def test_sample_dataset_is_reproducible_with_seed(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(path)

        sampled_a = sample_dataset(dataset, rows=2, seed=17)
        sampled_b = sample_dataset(dataset, rows=2, seed=17)
        self.assertEqual(len(sampled_a.df), 2)
        self.assertTrue(sampled_a.df.equals(sampled_b.df))
        self.assertIsNotNone(sampled_a.first_column_with_tag("RESULT"))

    def test_sample_dataset_can_sample_per_group_by_tag(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(path)

        sampled = sample_dataset(dataset, rows=1, seed=5, by_tags=["TESTNAME"])
        self.assertEqual(len(sampled.df), 2)
        self.assertEqual(set(sampled.df["검사항목명"]), {"AST", "ALT"})

    def test_execute_work_runs_sample_step(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.tame"
            path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(path)

        meta = dict(dataset.meta)
        meta["WORKS"] = {"DEFAULT": ["SAMPLE"]}
        meta["SAMPLE"] = {"ROWS": 1, "BY_TAGS": ["TESTNAME"], "SEED": 3}
        result = execute_work(dataset.replace(meta=meta), "DEFAULT")
        self.assertEqual(result.outputs[0].name, "SAMPLE")
        self.assertEqual(len(result.final_dataset.df), 2)
        self.assertEqual(set(result.final_dataset.df["검사항목명"]), {"AST", "ALT"})


class XlsxRoundTripTest(unittest.TestCase):
    def test_write_and_read_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            tame_path = Path(tmpdir) / "sample.tame"
            xlsx_path = Path(tmpdir) / "sample.xlsx"
            roundtrip_path = Path(tmpdir) / "roundtrip.tame"
            tame_path.write_text(SAMPLE_TAME, encoding="utf-8")

            dataset = read_tame(tame_path)
            write_xlsx(xlsx_path, dataset)
            loaded = read_xlsx(xlsx_path)
            write_tame(roundtrip_path, loaded)

            self.assertEqual(list(dataset.df.columns), list(loaded.df.columns))
            self.assertEqual(len(loaded.df), 3)
            self.assertTrue(roundtrip_path.exists())

    def test_state_tokens_roundtrip_in_tame_xlsx_and_pandas_copy(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            tame_path = Path(tmpdir) / "states.tame"
            tame_roundtrip = Path(tmpdir) / "states_roundtrip.tame"
            xlsx_path = Path(tmpdir) / "states.xlsx"
            tame_path.write_text(SAMPLE_VALUE_STATES_TAME, encoding="utf-8")

            dataset = read_tame(tame_path)
            pandas_copy = pd.DataFrame(dataset.df.copy(deep=True), columns=dataset.df.columns, dtype=object)
            pandas_dataset = dataset.replace(df=pandas_copy)

            write_tame(tame_roundtrip, pandas_dataset)
            tame_loaded = read_tame(tame_roundtrip)

            write_xlsx(xlsx_path, pandas_dataset)
            xlsx_loaded = read_xlsx(xlsx_path)

        for loaded in (tame_loaded, xlsx_loaded):
            self.assertIsNone(loaded.df.loc[0, "비고"])
            self.assertIs(loaded.df.loc[1, "비고"], NULL)
            self.assertEqual(loaded.df.loc[2, "비고"], "")
            self.assertEqual(loaded.df.loc[3, "비고"], "   ")
            self.assertEqual(loaded.df.loc[4, "비고"], "<<NULL>>")
            self.assertEqual(loaded.df.loc[4, "메모"], "<<EMPTY>>")
            self.assertEqual(cell_state(loaded.df.loc[0, "비고"]), "ABSENT")
            self.assertEqual(cell_state(loaded.df.loc[1, "비고"]), "NULL")

    def test_read_xlsx_multisheet_flattens_and_write_xlsx_restores_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "multisheet.xlsx"
            roundtrip_path = Path(tmpdir) / "multisheet_roundtrip.xlsx"

            workbook = openpyxl.Workbook()
            ws1 = workbook.active
            ws1.title = "Chemistry"
            ws1.append(["[[SEX]]성별", "[[AGE]]나이", "[[RESULT::NUM]]보고값"])
            ws1.append(["M", "25", "30"])
            ws2 = workbook.create_sheet("Hematology")
            ws2.append(["[[SEX]]sex", "[[AGE]]ageYears", "[[RESULT::NUM]]resultValue"])
            ws2.append(["F", "41", "12"])
            workbook.save(xlsx_path)

            dataset = read_xlsx(xlsx_path)
            write_xlsx(roundtrip_path, dataset)
            reloaded = read_xlsx(roundtrip_path)
            roundtrip_workbook = openpyxl.load_workbook(roundtrip_path, data_only=True)

        self.assertTrue(any(column.has_tag("SHEET") for column in dataset.columns))
        self.assertIn("시트명", list(dataset.df.columns))
        self.assertIn("SEX", list(dataset.df.columns))
        self.assertIn("AGE", list(dataset.df.columns))
        self.assertIn("RESULT", list(dataset.df.columns))
        self.assertEqual(set(dataset.df["시트명"]), {"Chemistry", "Hematology"})
        self.assertIn("Chemistry", roundtrip_workbook.sheetnames)
        self.assertIn("Hematology", roundtrip_workbook.sheetnames)
        self.assertEqual(set(reloaded.df["시트명"]), {"Chemistry", "Hematology"})

    def test_read_xlsx_multisheet_source_column_name_is_configurable(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "multisheet.xlsx"
            workbook = openpyxl.Workbook()
            ws1 = workbook.active
            ws1.title = "A"
            ws1.append(["value"])
            ws1.append(["1"])
            ws2 = workbook.create_sheet("B")
            ws2.append(["value"])
            ws2.append(["2"])
            workbook.save(xlsx_path)

            dataset = read_xlsx(xlsx_path, source_column_name="sheet_name")

        self.assertIn("sheet_name", list(dataset.df.columns))
        self.assertEqual(set(dataset.df["sheet_name"]), {"A", "B"})
        self.assertEqual(dataset.meta["MULTISHEET"]["SOURCE_COLUMN_NAME"], "sheet_name")


class MergeTest(unittest.TestCase):
    def test_merge_warns_on_num_conflict_and_promotes(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            num_path = Path(tmpdir) / "num.tame"
            cnum_path = Path(tmpdir) / "cnum.tame"
            num_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")
            cnum_path.write_text(SAMPLE_CNUM_TAME, encoding="utf-8")
            num_dataset = read_tame(num_path)
            cnum_dataset = read_tame(cnum_path)

        result = merge_datasets([num_dataset, cnum_dataset], numeric_conflict="promote")
        result_column = result.dataset.first_column_with_tag("RESULT")
        self.assertIsNotNone(result_column)
        self.assertIn("<NUM>", result_column.tags)
        self.assertTrue(any("NUM_CONFLICT" in warning for warning in result.warnings))
        self.assertIn("source_dataset", list(result.dataset.df.columns))

    def test_merge_split_adds_comparator_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            num_path = Path(tmpdir) / "num.tame"
            cnum_path = Path(tmpdir) / "cnum.tame"
            num_path.write_text(SAMPLE_NUM_TAME, encoding="utf-8")
            cnum_path.write_text(SAMPLE_CNUM_TAME, encoding="utf-8")
            num_dataset = read_tame(num_path)
            cnum_dataset = read_tame(cnum_path)

        result = merge_datasets([num_dataset, cnum_dataset], numeric_conflict="split")
        self.assertIn("보고값__cmp", list(result.dataset.df.columns))
        self.assertIn("보고값__num", list(result.dataset.df.columns))


class ExportTest(unittest.TestCase):
    def test_export_csv_preserves_serialized_states_and_pandas_load(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "states.tame"
            output_path = Path(tmpdir) / "states.csv"
            input_path.write_text(SAMPLE_VALUE_STATES_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

            result = export_dataset(dataset, output_path)
            exported = pd.read_csv(output_path, dtype=str, keep_default_na=False)

        self.assertEqual(result.format, "csv")
        self.assertEqual(list(exported.columns), ["비고", "코드", "메모", "수치"])
        self.assertEqual(exported.loc[0, "비고"], "<<ABSENT>>")
        self.assertEqual(exported.loc[1, "비고"], "<<NULL>>")
        self.assertEqual(exported.loc[2, "비고"], "<<EMPTY>>")
        self.assertEqual(exported.loc[3, "비고"], "<<WS:3>>")
        self.assertEqual(exported.loc[4, "비고"], "\\<<NULL>>")
        self.assertEqual(exported.loc[4, "메모"], "\\<<EMPTY>>")

    def test_export_r_bundle_writes_r_loader_and_sidecars(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            bundle_path = Path(tmpdir) / "sample_bundle"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

            result = export_dataset(dataset, bundle_path, format="r_bundle")
            data_frame = pd.read_csv(bundle_path / "data.csv", dtype=str, keep_default_na=False)
            columns_frame = pd.read_csv(bundle_path / "columns.csv", dtype=str, keep_default_na=False)
            loader_text = (bundle_path / "load_tame_bundle.R").read_text(encoding="utf-8")
            readme_text = (bundle_path / "README_R.md").read_text(encoding="utf-8")
            manifest_text = (bundle_path / "bundle.json").read_text(encoding="utf-8")
            data_exists = (bundle_path / "data.csv").exists()
            columns_exists = (bundle_path / "columns.csv").exists()
            meta_exists = (bundle_path / "meta.toml").exists()

        self.assertEqual(result.format, "r_bundle")
        self.assertTrue(data_exists)
        self.assertTrue(columns_exists)
        self.assertTrue(meta_exists)
        self.assertIn("read_tame_bundle <- function", loader_text)
        self.assertIn("RESULT::<NUM>", "".join(columns_frame["tags"].tolist()))
        self.assertEqual(data_frame.loc[1, "보고값"], "<3")
        self.assertIn("\"data_format\": \"csv\"", manifest_text)
        self.assertIn("split-comparator", readme_text)

    def test_export_sql_writes_data_columns_and_sections(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            output_path = Path(tmpdir) / "sample_export.sql"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

            result = export_dataset(dataset, output_path, format="sql")
            sql_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(result.format, "sql")
        self.assertIn('CREATE TABLE "sample_export_data"', sql_text)
        self.assertIn('CREATE TABLE "sample_export_columns"', sql_text)
        self.assertIn('CREATE TABLE "sample_export_sections"', sql_text)
        self.assertIn("INSERT INTO \"sample_export_sections\" VALUES ('META'", sql_text)
        self.assertIn("보고값", sql_text)
        self.assertIn("<3", sql_text)

    def test_merge_harmonize_rewrites_threshold_band(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            p1 = Path(tmpdir) / "hospital_a.tame"
            p2 = Path(tmpdir) / "hospital_b.tame"
            p1.write_text(
                textwrap.dedent(
                    """
                    <META>
                    </META>
                    <DATA>
                    [[ITEM]]검사항목명\t[[RESULT::<NUM>]]보고값
                    AST\t<30
                    AST\t29
                    </DATA>
                    """
                ).strip(),
                encoding="utf-8",
            )
            p2.write_text(
                textwrap.dedent(
                    """
                    <META>
                    </META>
                    <DATA>
                    [[ITEM]]검사항목명\t[[RESULT::<NUM>]]보고값
                    AST\t<20
                    AST\t18
                    </DATA>
                    """
                ).strip(),
                encoding="utf-8",
            )
            d1 = read_tame(p1)
            d2 = read_tame(p2)

        result = merge_datasets([d1, d2], numeric_conflict="harmonize")
        ast_values = result.dataset.df.loc[result.dataset.df["검사항목명"] == "AST", "보고값"].tolist()
        self.assertIn("<30", ast_values)
        self.assertEqual(ast_values.count("<30"), 3)
        self.assertTrue(any("HARMONIZE_APPLIED" in warning for warning in result.warnings))

    def test_merge_uses_tag_name_when_column_names_differ(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            p1 = Path(tmpdir) / "a.tame"
            p2 = Path(tmpdir) / "b.tame"
            p1.write_text(
                textwrap.dedent(
                    """
                    <META></META>
                    <DATA>
                    [[SEX]]성별\t[[AGE]]나이\t[[RESULT::NUM]]보고값
                    M\t20\t10
                    </DATA>
                    """
                ).strip(),
                encoding="utf-8",
            )
            p2.write_text(
                textwrap.dedent(
                    """
                    <META></META>
                    <DATA>
                    [[SEX]]sex\t[[AGE]]ageYears\t[[RESULT::NUM]]resultValue
                    F\t30\t20
                    </DATA>
                    """
                ).strip(),
                encoding="utf-8",
            )
            d1 = read_tame(p1)
            d2 = read_tame(p2)

        result = merge_datasets([d1, d2], add_source_column=False)
        self.assertIn("SEX", list(result.dataset.df.columns))
        self.assertIn("AGE", list(result.dataset.df.columns))
        self.assertIn("RESULT", list(result.dataset.df.columns))
        self.assertTrue(any("tag_name 'SEX'" in warning for warning in result.warnings))


class SidecarTest(unittest.TestCase):
    def test_data_tame_auto_loads_sibling_meta_tame(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "sample.tame"
            meta_path = Path(tmpdir) / "sample.meta.tame"
            data_path = Path(tmpdir) / "sample.data.tame"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(input_path)

            write_meta_tame(meta_path, dataset)
            write_data_tame(data_path, dataset)
            meta_text = meta_path.read_text(encoding="utf-8")
            reloaded = read_tame(data_path)

        self.assertIn("[COLUMN.\"보고값\"]", meta_text)
        self.assertIn('TAGS = ["RESULT", "<NUM>"]', meta_text)
        self.assertEqual(reloaded.settings()["CRR"], "VALUE")
        self.assertIsNotNone(reloaded.first_column_with_tag("RESULT"))
        self.assertEqual(execute_work(reloaded, "DEFAULT").work_name, "DEFAULT")

    def test_xlsx_with_meta_tame_restores_tags_and_pipeline(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "template.tame"
            meta_path = Path(tmpdir) / "fresh.meta.tame"
            xlsx_path = Path(tmpdir) / "fresh.xlsx"
            input_path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(input_path)
            write_meta_tame(meta_path, dataset)

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["등록번호", "성별", "나이", "검사항목명", "보고값", "진료과"])
            sheet.append(["A1001", "F", "32", "AST", "25", "IM"])
            sheet.append(["A1002", "X", "bad", "ALT", "positive", "GS"])
            workbook.save(xlsx_path)

            loaded = read_xlsx(xlsx_path)
            result = execute_work(loaded, "DEFAULT")

        self.assertIsNotNone(loaded.first_column_with_tag("RESULT"))
        self.assertIsNotNone(loaded.first_column_with_tag("AGE"))
        self.assertEqual(result.outputs[0].name, "VALIDATE")
        self.assertEqual(result.outputs[1].name, "DESCRIBE")
        self.assertEqual(len(result.final_dataset.df), 1)

    def test_import_xlsx_into_tame_preserves_template_tags(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            template_path = Path(tmpdir) / "template.tame"
            output_path = Path(tmpdir) / "updated.tame"
            xlsx_path = Path(tmpdir) / "fresh.xlsx"
            template_path.write_text(SAMPLE_TAME, encoding="utf-8")
            template = read_tame(template_path)

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["등록번호", "성별", "나이", "검사항목명", "보고값", "진료과"])
            sheet.append(["B0001", "F", "44", "AST", "18", "IM"])
            sheet.append(["B0002", "M", "55", "ALT", "<12", "GS"])
            workbook.save(xlsx_path)

            imported = import_xlsx_into_tame(template, xlsx_path)
            write_tame(output_path, imported)
            written = output_path.read_text(encoding="utf-8")

        self.assertEqual(imported.columns, template.columns)
        self.assertEqual(imported.df.loc[0, "등록번호"], "B0001")
        self.assertEqual(imported.df.loc[1, "보고값"], "<12")
        self.assertEqual(imported.settings()["CRR"], "VALUE")
        self.assertEqual(imported.tagged_headers()[4], "[[RESULT::<NUM>]]보고값")
        self.assertIn("[[RESULT::<NUM>]]보고값", written)
        self.assertIn("B0002\tmale\t55\tALT\t<12\tGS", written)

    def test_import_xlsx_into_tame_rejects_header_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            template_path = Path(tmpdir) / "template.tame"
            xlsx_path = Path(tmpdir) / "fresh.xlsx"
            template_path.write_text(SAMPLE_TAME, encoding="utf-8")
            template = read_tame(template_path)

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            sheet.append(["등록번호", "성별", "나이", "검사항목명", "결과값", "진료과"])
            sheet.append(["B0001", "F", "44", "AST", "18", "IM"])
            workbook.save(xlsx_path)

            with self.assertRaisesRegex(ValueError, "XLSX import header validation failed"):
                import_xlsx_into_tame(template, xlsx_path)


class EmbeddedFunctionTest(unittest.TestCase):
    def test_execute_work_rejects_embedded_functions_without_opt_in(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "function.tame"
            path.write_text(SAMPLE_FUNCTION_TAME, encoding="utf-8")
            dataset = read_tame(path)

        with self.assertRaisesRegex(PermissionError, "allow_functions=True"):
            execute_work(dataset, "DEFAULT")

    def test_execute_work_runs_embedded_python_function(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "function.tame"
            path.write_text(SAMPLE_FUNCTION_TAME, encoding="utf-8")
            dataset = read_tame(path)

        result = execute_work(dataset, "DEFAULT", allow_functions=True)

        self.assertEqual(result.outputs[0].name, "TRIM_RESULT")
        self.assertEqual(result.final_dataset.df.loc[0, "보고값"], "alpha")
        self.assertEqual(result.final_dataset.df.loc[1, "보고값"], "beta")
        self.assertIsNotNone(result.final_dataset.first_column_with_tag("RESULT"))

    @unittest.skipIf(shutil.which("Rscript") is None, "Rscript is not available")
    def test_execute_work_runs_embedded_r_function(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "function_r.tame"
            path.write_text(SAMPLE_R_FUNCTION_TAME, encoding="utf-8")
            dataset = read_tame(path)

        result = execute_work(dataset, "DEFAULT", allow_functions=True)

        self.assertEqual(result.outputs[0].name, "TRIM_RESULT_R")
        self.assertEqual(result.outputs[0].message, "trimmed by R")
        self.assertEqual(result.final_dataset.df.loc[0, "보고값"], "gamma")
        self.assertEqual(result.final_dataset.df.loc[1, "보고값"], "delta")
        self.assertIsNotNone(result.final_dataset.first_column_with_tag("RESULT"))


class CommandPipelineTest(unittest.TestCase):
    def test_available_command_pipelines_reads_meta_section(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline.tame"
            path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(path)

        meta = dict(dataset.meta)
        meta["PIPELINES"] = {"DEFAULT": ["sample --rows 1 --by-tag TESTNAME --seed 11", "describe"]}
        names = available_command_pipelines(meta)
        self.assertEqual(names, ["DEFAULT"])

    def test_execute_command_pipeline_chains_sample_and_nested_work(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline.tame"
            path.write_text(SAMPLE_RI_TAME, encoding="utf-8")
            dataset = read_tame(path)

        meta = dict(dataset.meta)
        meta["PIPELINES"] = {"DEFAULT": ["sample --rows 1 --by-tag TESTNAME --seed 11", "run RI"]}
        result = execute_command_pipeline(dataset.replace(meta=meta), "DEFAULT")

        self.assertEqual(result.outputs[0].name, "sample --rows 1 --by-tag TESTNAME --seed 11")
        self.assertEqual(result.outputs[1].name, "run RI")
        self.assertEqual(len(result.final_dataset.df), 8)
        self.assertIn("ref_low", list(result.final_dataset.df.columns))
        self.assertEqual(set(result.outputs[1].dataset.df["test_name"]), {"AST", "ALT"})

    def test_execute_command_pipeline_supports_relative_export_paths(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline.tame"
            export_path = Path(tmpdir) / "exports" / "sample.sql"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

            meta = dict(dataset.meta)
            # 명령 문자열에는 OS 구분자 대신 정방향 슬래시를 쓴다(Windows 백슬래시는 shlex 에서 깨진다).
            meta["PIPELINES"] = {"DEFAULT": [f"export {export_path.relative_to(path.parent).as_posix()} --format sql"]}
            result = execute_command_pipeline(dataset.replace(meta=meta), "DEFAULT")
            sql_text = export_path.read_text(encoding="utf-8")

        self.assertEqual(result.outputs[0].name, "export exports/sample.sql --format sql")
        self.assertIn("CREATE TABLE", sql_text)

    def test_execute_command_pipeline_supports_evaluate_command(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline.tame"
            report_dir = Path(tmpdir) / "reports"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

            meta = dict(dataset.meta)
            meta["PIPELINES"] = {"DEFAULT": ["evaluate --tame-steps 2 --output-dir reports"]}
            result = execute_command_pipeline(dataset.replace(meta=meta), "DEFAULT")
            report_written = (report_dir / "dataset_quality.csv").exists()

        self.assertEqual(result.outputs[0].name, "evaluate --tame-steps 2 --output-dir reports")
        self.assertIn("dataset_quality", result.outputs[0].tables)
        self.assertTrue(report_written)

    def test_execute_command_pipeline_can_run_action_pipeline(self) -> None:
        sample = textwrap.dedent(
            """
            <META>
            [ACTIONS.KEEP_A]
            TYPE = "ROW.INCLUDE"
            COLUMN = "group"
            VALUE = "A"

            [ACTION_PIPELINES]
            CLEAN = ["KEEP_A"]

            [PIPELINES]
            DEFAULT = ["run-action-pipeline CLEAN", "describe"]
            </META>
            <DATA>
            group\tvalue
            A\t1
            B\t2
            </DATA>
            """
        ).strip()
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline_action.tame"
            path.write_text(sample, encoding="utf-8")
            dataset = read_tame(path)

        result = execute_command_pipeline(dataset, "DEFAULT")

        self.assertEqual(result.outputs[0].name, "run-action-pipeline CLEAN")
        self.assertEqual(result.outputs[0].dataset.df["group"].tolist(), ["A"])
        self.assertEqual(len(result.final_dataset.df), 1)

    def test_execute_command_pipeline_rejects_unknown_command(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        meta = dict(dataset.meta)
        meta["PIPELINES"] = {"DEFAULT": ["merge other.tame --output merged.tame"]}
        with self.assertRaisesRegex(CommandPipelineError, "Unsupported command"):
            execute_command_pipeline(dataset.replace(meta=meta), "DEFAULT")

    def test_execute_command_pipeline_rejects_undefined_pipeline_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pipeline.tame"
            path.write_text(SAMPLE_TAME, encoding="utf-8")
            dataset = read_tame(path)

        with self.assertRaisesRegex(CommandPipelineError, "Undefined command pipeline"):
            execute_command_pipeline(dataset, "DEFAULT")


class ImageTest(unittest.TestCase):
    def test_extract_and_embed_image_columns_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            tame_path = Path(tmpdir) / "images.tame"
            image_dir = Path(tmpdir) / "images"
            tame_path.write_text(SAMPLE_IMAGE_TAME, encoding="utf-8")
            dataset = read_tame(tame_path)

            extracted, manifest = extract_image_columns(dataset, image_dir)
            path_value = extracted.df.loc[0, "썸네일"]
            reembedded, embed_manifest = embed_image_columns(extracted, base_dir=Path.cwd())
            self.assertEqual(manifest.loc[0, "mime"], "image/png")
            self.assertTrue(Path(path_value).exists())
            self.assertTrue(any(column.has_tag("PATH") for column in extracted.columns if column.name == "썸네일"))
            self.assertEqual(reembedded.df.loc[0, "썸네일"], SAMPLE_IMAGE_DATA_URL)
            self.assertEqual(embed_manifest.loc[0, "mime"], "image/png")


class TutorialExamplesTest(unittest.TestCase):
    def test_capability_probe_20_examples_all_pass(self) -> None:
        # 저장된 20개 예제를 원본/컬럼명 변경 데이터 양쪽에서 전부 PASS하는지 확인.
        repo_root = Path(__file__).resolve().parents[2]
        runner = repo_root / "tutorial" / "14_capability_probe" / "examples" / "run_all.py"
        if not runner.exists():
            self.skipTest("examples/run_all.py not present")
        env = dict(os.environ)
        env["PYTHONPATH"] = str(repo_root / "tametools" / "src") + os.pathsep + env.get("PYTHONPATH", "")
        proc = subprocess.run(
            [sys.executable, str(runner)],
            capture_output=True, text=True, env=env, cwd=str(repo_root),
        )
        self.assertEqual(proc.returncode, 0, msg=proc.stdout + proc.stderr)
        self.assertIn("40/40 PASS", proc.stdout)


if __name__ == "__main__":
    unittest.main()
