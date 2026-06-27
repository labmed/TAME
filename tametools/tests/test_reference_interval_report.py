from __future__ import annotations

import importlib.util
import math
from pathlib import Path
from statistics import NormalDist
import sys
import tempfile
import textwrap
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tametools" / "src"))

from tametools.io import read_tame
from tametools.plugin_base import run_plugin


HAS_REPORT_DEPS = importlib.util.find_spec("docx") is not None and importlib.util.find_spec("matplotlib") is not None


def _read_tame_text(text: str):
    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "sample.tame"
        path.write_text(textwrap.dedent(text).strip(), encoding="utf-8")
        return read_tame(path)


def _clinical_rows(n: int = 130) -> str:
    rows = []
    for index in range(n):
        sex = "male" if index % 2 == 0 else "female"
        age = 20 + (index % 45)
        value = 20 + (index % 60) * 0.5
        rows.append(f"AST\t{sex}\t{age}\t{value:.1f}")
    return "\n".join(rows)


class ReferenceIntervalReportTests(unittest.TestCase):
    def test_reference_interval_returns_clinical_tables_charts_and_reliability(self) -> None:
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[SEX::CATEGORY::BY]]sex\t[[AGE]]age\t[[RESULT::NUM]]value
            {_clinical_rows()}
            </DATA>
            """
        )

        output = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"BOOTSTRAP_N": 40, "METHOD_COMPARISON_BOOTSTRAP_N": 20, "SEED": 7},
        )

        self.assertIsNotNone(output.dataset)
        self.assertIn("summary_by_test", output.tables)
        self.assertIn("reliability_summary", output.tables)
        self.assertIn("outliers", output.tables)
        self.assertIn("normality", output.tables)
        self.assertIn("method_comparison", output.tables)
        self.assertIn("group_difference_tests", output.tables)
        self.assertEqual({chart["name"] for chart in output.charts}, {"RI_REF_WIDTH", "RI_N_BY_GROUP", "RI_HIGH_LIMIT"})

        main = output.table
        ast_all = main.loc[(main["test_name"] == "AST") & (main["group_level"] == "TESTNAME")].iloc[0]
        self.assertEqual(int(ast_all["n"]), 130)
        self.assertEqual(ast_all["reliability"], "ok")
        self.assertEqual(ast_all["recommendation"], "candidate_for_verification")
        self.assertEqual(ast_all["outlier_method"], "NONE")
        self.assertTrue(ast_all["method"].startswith("nonparametric"))
        self.assertLessEqual(float(ast_all["ref_low_ci90_low"]), float(ast_all["ref_low_ci90_high"]))
        self.assertLessEqual(float(ast_all["ref_high_ci90_low"]), float(ast_all["ref_high_ci90_high"]))
        self.assertIn("RI_REF_WIDTH", output.dataset.meta["VISUALIZATIONS"])
        for row in output.table.itertuples(index=False):
            n = int(row.n)
            method = str(row.method)
            if n >= 120:
                self.assertTrue(method.startswith("nonparametric"))
            elif n >= 40:
                self.assertTrue(method.startswith("robust"))
            else:
                self.assertEqual(method, "not_calculated_insufficient_n")
        self.assertIn("adequate", set(output.table["ci_adequacy"]))
        self.assertEqual(
            set(output.tables["method_comparison"]["method_family"]),
            {"parametric", "nonparametric", "robust"},
        )
        self.assertIn("normality_decision", output.tables["normality"].columns)
        self.assertFalse(output.tables["group_difference_tests"].empty)

    def test_reference_interval_default_keeps_outliers_for_review(self) -> None:
        dataset = _read_tame_text(
            """
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            AST\t1
            AST\t2
            AST\t3
            AST\t4
            AST\t100
            </DATA>
            """
        )

        output = run_plugin(dataset, "REFERENCE_INTERVAL", dataset.meta, "REFERENCE_INTERVAL", {"BOOTSTRAP_N": 0})

        row = output.table.iloc[0]
        self.assertEqual(int(row["original_n"]), 5)
        self.assertEqual(int(row["n"]), 5)
        self.assertEqual(int(row["outliers_removed"]), 0)
        self.assertEqual(row["outlier_method"], "NONE")
        self.assertTrue(output.tables["outliers"].empty)

    def test_reference_interval_records_tukey_outliers(self) -> None:
        dataset = _read_tame_text(
            """
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            AST\t1
            AST\t2
            AST\t3
            AST\t4
            AST\t100
            </DATA>
            """
        )

        output = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"OUTLIER_METHOD": "TUKEY", "BOOTSTRAP_N": 10},
        )

        row = output.table.iloc[0]
        outliers = output.tables["outliers"]
        self.assertEqual(int(row["original_n"]), 5)
        self.assertEqual(int(row["n"]), 4)
        self.assertEqual(int(row["outliers_removed"]), 1)
        self.assertEqual(float(outliers.iloc[0]["value"]), 100.0)
        self.assertEqual(row["method"], "not_calculated_insufficient_n")
        self.assertTrue(math.isnan(float(row["ref_high"])))
        self.assertTrue(any("High outlier removal rate" in warning for warning in output.warnings))

    def test_reference_interval_log_tukey_is_less_aggressive_for_skewed_values(self) -> None:
        rows = "\n".join(f"FERR\t{math.exp(2.5 + index * 0.04):.2f}" for index in range(80))
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {rows}
            </DATA>
            """
        )

        raw = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"OUTLIER_METHOD": "TUKEY", "BOOTSTRAP_N": 0},
        )
        logged = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"OUTLIER_METHOD": "LOG_TUKEY", "BOOTSTRAP_N": 0},
        )

        self.assertGreater(int(raw.table.iloc[0]["outliers_removed"]), int(logged.table.iloc[0]["outliers_removed"]))
        self.assertEqual(logged.table.iloc[0]["outlier_method"], "LOG_TUKEY")
        self.assertEqual(int(logged.table.iloc[0]["outliers_removed"]), 0)

    def test_reference_interval_supports_parametric_methods(self) -> None:
        values = list(range(80, 121))
        rows = "\n".join(f"GLU\t{value}" for value in values)
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {rows}
            </DATA>
            """
        )

        output = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"METHOD": "PARAMETRIC", "BOOTSTRAP_N": 0},
        )

        row = output.table.iloc[0]
        z_low = NormalDist().inv_cdf(0.025)
        z_high = NormalDist().inv_cdf(0.975)
        mean = sum(values) / len(values)
        sd = math.sqrt(sum((value - mean) ** 2 for value in values) / (len(values) - 1))
        self.assertEqual(row["method"], "robust_median_mad_0.025_0.975")
        parametric = output.tables["method_comparison"].loc[
            output.tables["method_comparison"]["method_family"] == "parametric"
        ].iloc[0]
        self.assertAlmostEqual(float(parametric["ref_low"]), mean + z_low * sd, places=6)
        self.assertAlmostEqual(float(parametric["ref_high"]), mean + z_high * sd, places=6)

        log_rows = "\n".join(f"FERR\t{math.exp(3.0 + index * 0.03):.6f}" for index in range(50))
        log_dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {log_rows}
            </DATA>
            """
        )
        logged = run_plugin(
            log_dataset,
            "REFERENCE_INTERVAL",
            log_dataset.meta,
            "REFERENCE_INTERVAL",
            {"METHOD": "LOG_PARAMETRIC", "BOOTSTRAP_N": 0},
        )
        self.assertEqual(logged.table.iloc[0]["method"], "robust_median_mad_0.025_0.975")
        self.assertGreaterEqual(float(logged.table.iloc[0]["ref_low"]), 0.0)

    def test_reference_interval_supports_robust_method_for_review_groups(self) -> None:
        values = [float(value) for value in range(80, 137)] + [250.0, 280.0, 300.0]
        rows = "\n".join(f"GLU\t{value}" for value in values)
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {rows}
            </DATA>
            """
        )

        robust = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"METHOD": "ROBUST", "BOOTSTRAP_N": 0},
        )
        nonparametric = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"METHOD": "NONPARAMETRIC", "BOOTSTRAP_N": 0},
        )

        robust_row = robust.table.iloc[0]
        nonparametric_row = nonparametric.table.iloc[0]
        self.assertEqual(robust_row["reliability"], "review")
        self.assertEqual(robust_row["method"], "robust_median_mad_0.025_0.975")
        self.assertEqual(nonparametric_row["method"], "robust_median_mad_0.025_0.975")
        comparison = robust.tables["method_comparison"]
        robust_high = float(comparison.loc[comparison["method_family"] == "robust", "ref_high"].iloc[0])
        nonparametric_high = float(comparison.loc[comparison["method_family"] == "nonparametric", "ref_high"].iloc[0])
        self.assertLess(robust_high, nonparametric_high)

        ok_rows = "\n".join(f"GLU\t{value}" for value in range(1, 131))
        ok_dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {ok_rows}
            </DATA>
            """
        )
        ok_output = run_plugin(
            ok_dataset,
            "REFERENCE_INTERVAL",
            ok_dataset.meta,
            "REFERENCE_INTERVAL",
            {"METHOD": "ROBUST", "BOOTSTRAP_N": 0},
        )
        self.assertEqual(ok_output.table.iloc[0]["reliability"], "ok")
        self.assertEqual(ok_output.table.iloc[0]["method"], "nonparametric_quantile_0.025_0.975")

    def test_reference_interval_rank_ci_uses_expected_order_statistics(self) -> None:
        rows = "\n".join(f"RI\t{value}" for value in range(1, 121))
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {rows}
            </DATA>
            """
        )

        output = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"CI_METHOD": "RANK", "BOOTSTRAP_N": 0},
        )

        row = output.table.iloc[0]
        self.assertEqual(row["ci_method"], "rank_nonparametric_0.90")
        self.assertEqual(float(row["ref_low_ci90_low"]), 1.0)
        self.assertEqual(float(row["ref_low_ci90_high"]), 7.0)
        self.assertEqual(float(row["ref_high_ci90_low"]), 114.0)
        self.assertEqual(float(row["ref_high_ci90_high"]), 120.0)

    def test_reference_interval_rank_ci_warns_when_sample_is_too_small(self) -> None:
        rows = "\n".join(f"RI\t{value}" for value in range(1, 61))
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            {rows}
            </DATA>
            """
        )

        output = run_plugin(
            dataset,
            "REFERENCE_INTERVAL",
            dataset.meta,
            "REFERENCE_INTERVAL",
            {"CI_METHOD": "RANK", "BOOTSTRAP_N": 0},
        )

        row = output.table.iloc[0]
        self.assertEqual(row["method"], "robust_median_mad_0.025_0.975")
        self.assertEqual(row["ci_method"], "bootstrap_percentile_0.90_robust")
        self.assertLessEqual(float(row["ref_low_ci90_low"]), float(row["ref_low_ci90_high"]))
        self.assertLessEqual(float(row["ref_high_ci90_low"]), float(row["ref_high_ci90_high"]))

    def test_reference_interval_rank_ci_rejects_non_nonparametric_method(self) -> None:
        dataset = _read_tame_text(
            """
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[RESULT::NUM]]value
            RI\t1
            RI\t2
            </DATA>
            """
        )

        with self.assertRaisesRegex(ValueError, "CI_METHOD=RANK"):
            run_plugin(
                dataset,
                "REFERENCE_INTERVAL",
                dataset.meta,
                "REFERENCE_INTERVAL",
                {"METHOD": "PARAMETRIC", "CI_METHOD": "RANK"},
            )

    @unittest.skipUnless(HAS_REPORT_DEPS, "report extras are not installed")
    def test_reference_interval_writes_docx_report_with_chart(self) -> None:
        dataset = _read_tame_text(
            f"""
            <DATA>
            [[TESTNAME::CATEGORY]]test\t[[SEX::CATEGORY::BY]]sex\t[[AGE]]age\t[[RESULT::NUM]]value
            {_clinical_rows()}
            </DATA>
            """
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            report_path = Path(tmpdir) / "reference_interval_report.docx"
            output = run_plugin(
                dataset,
                "REFERENCE_INTERVAL",
                dataset.meta,
                "REFERENCE_INTERVAL",
                {"REPORT_PATH": str(report_path), "BOOTSTRAP_N": 20, "METHOD_COMPARISON_BOOTSTRAP_N": 20},
            )

            self.assertIn(str(report_path), output.files)
            self.assertTrue(report_path.exists())
            with zipfile.ZipFile(report_path) as archive:
                images = [name for name in archive.namelist() if name.startswith("word/media/") and name.endswith(".png")]
                document_xml = archive.read("word/document.xml").decode("utf-8")
                settings_xml = archive.read("word/settings.xml").decode("utf-8")
                footer_xml = "\n".join(
                    archive.read(name).decode("utf-8")
                    for name in archive.namelist()
                    if name.startswith("word/footer") and name.endswith(".xml")
                )
            self.assertGreater(len(images), 3)
            self.assertIn("TOC", document_xml)
            self.assertIn("w:tblHeader", document_xml)
            self.assertIn("DDEEDC", document_xml)
            self.assertIn("F4CCCC", document_xml)
            self.assertIn("updateFields", settings_xml)
            self.assertIn("PAGE", footer_xml)
            self.assertIn("NUMPAGES", footer_xml)

            from docx import Document  # noqa: PLC0415

            document = Document(report_path)
            table_widths = [len(table.columns) for table in document.tables]
            self.assertTrue(table_widths)
            self.assertLessEqual(max(table_widths), 10)
            headings = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertNotIn("Showing first", headings)
            self.assertIn("Table of Contents", headings)
            self.assertIn("Executive Summary", headings)
            self.assertIn("Normality Assessment", headings)
            self.assertIn("Candidate Reference Intervals", headings)
            self.assertIn("Review-Grade Reference Intervals", headings)
            self.assertIn("Do Not Use / Additional Data Required", headings)
            self.assertIn("Method Comparison", headings)
            self.assertIn("Reference Interval Difference Tests", headings)
            self.assertNotIn("Legacy Harris-Boyd Partition Tests", headings)
            self.assertIn("Group Distributions", headings)
            table_headers = [[cell.text for cell in table.rows[0].cells] for table in document.tables if table.rows]
            candidate_header = next(header for header in table_headers if "Ci Adequacy" in header and "Method" in header)
            self.assertIn("Ref Low Ci", candidate_header)
            self.assertIn("Ref High Ci", candidate_header)
            do_not_use_header = next(header for header in table_headers if "N Range" in header and "Partitions" in header)
            self.assertNotIn("Recommendation", do_not_use_header)
            candidate_table = next(
                table for table in document.tables if table.rows and "Ci Adequacy" in [cell.text for cell in table.rows[0].cells]
            )
            candidate_text = "\n".join(cell.text for row in candidate_table.rows[1:] for cell in row.cells)
            self.assertIn("Nonparametric", candidate_text)
            self.assertNotIn("nonparametric_quantile", candidate_text)
            normality_table = next(
                table for table in document.tables if table.rows and "Skewness" in [cell.text for cell in table.rows[0].cells]
            )
            normality_levels = [row.cells[1].text for row in normality_table.rows[1:]]
            normality_rank = {"All": 0, "Sex": 1, "Age": 2, "Sex + Age": 3}
            self.assertEqual(normality_levels[0], "All")
            self.assertEqual(
                [normality_rank[level] for level in normality_levels],
                sorted(normality_rank[level] for level in normality_levels),
            )
